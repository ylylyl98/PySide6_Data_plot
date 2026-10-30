"""Pure MCD numerical transactions plus real native Quick controls, isolated."""
import os
from pathlib import Path
import subprocess
import sys
from tempfile import TemporaryDirectory
import unittest


class QuickMcdTests(unittest.TestCase):
    def test_native_mcd(self):
        with TemporaryDirectory() as temporary:
            env=dict(os.environ,QT_QPA_PLATFORM="offscreen",QT_QUICK_BACKEND="software",QT_SCALE_FACTOR="1",
                MPLBACKEND="agg",MPLCONFIGDIR=temporary,LOCALAPPDATA=temporary,XDG_CONFIG_HOME=temporary,
                XDG_CACHE_HOME=temporary,QML_DISK_CACHE_PATH=temporary)
            result=subprocess.run([sys.executable,str(Path(__file__).resolve()),"--isolated"],
                cwd=Path(__file__).resolve().parents[1],env=env,capture_output=True,text=True,timeout=240)
        self.assertEqual(result.returncode,0,result.stdout+result.stderr)


def _isolated():
    sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
    import csv
    from dataclasses import replace
    import json
    from threading import Event
    import time
    import numpy as np
    from PySide6.QtCore import QObject,QSettings,QUrl
    from PySide6.QtGui import QGuiApplication
    from PySide6.QtQml import QQmlComponent,QQmlEngine
    from PySide6.QtQuick import QQuickItem,QQuickWindow
    from core.mcd import McdSettings,process_mcd,pair_window_trace_by_branch,suggest_mcd_background_ranges
    from core.mcd_analysis import fit_mcd_slopes
    from ui_qt.quick.plot_surface import register_plot_surface
    from ui_qt.quick.mcd_bridge import McdBridge
    from ui_qt.quick.mcd_workflow import McdRequest,load_mcd,window_snapshot,fingerprints,input_paths,export_snapshot,export_mcd

    app=QGuiApplication([]);register_plot_surface()
    from ui_qt.typography import install_application_font
    install_application_font(app)
    colors=dict(surface="#f3f5f8",control="#ffffff",text="#20242b",muted="#606978",border="#c6cbd3",accent="#0067c0",hover="#edf1f6",pressed="#e0e7ef",error="#b42318",disabled="#939aa5",selectedText="#ffffff",inset="#e9edf2",divider="#dce1e8",borderHover="#aeb7c4",borderDisabled="#dce1e8",controlDisabled="#edf0f4",focus="#20242b",accentText="#0067c0",selected="#e6f0fa",selectedBorder="#b4d6fa",errorSurface="#fdf1f0")

    def settle(milliseconds=15):
        deadline=time.monotonic()+milliseconds/1000
        while time.monotonic()<deadline:app.processEvents();time.sleep(.001)
        app.processEvents()

    class Cases(unittest.TestCase):
        def setUp(self):
            self.temporary=TemporaryDirectory();self.addCleanup(self.temporary.cleanup)
            self.root=Path(self.temporary.name).resolve();self.energy=np.linspace(1.5,1.9,201)
            self.source=self.root/"mcd"/"synthetic_MCD.csv";self.source.parent.mkdir()
            self.write_source(self.source)
            self.settings=QSettings(str(self.root/"test.ini"),QSettings.Format.IniFormat)
            self.bridge=McdBridge(settings=self.settings);self.addCleanup(self.bridge.close)

        def write_source(self,path,scale=1.,fields=None):
            fields=np.r_[np.linspace(-2,2,41),np.linspace(1.9,-2,40)] if fields is None else fields
            with path.open("w",newline="") as stream:
                writer=csv.writer(stream);writer.writerow(["B_T","angle_deg",*list(1239.841984/self.energy)])
                for field in fields:
                    for angle,sign,gain in ((90,1,1.),(0,-1,1.08)):
                        raw=scale*gain*(10000+1800*np.exp(-((self.energy-1.7-sign*.001*field)/.009)**2)+300*np.exp(-((self.energy-1.78)/.014)**2))
                        writer.writerow([field,angle,*raw])

        def wait(self,kind="all"):
            deadline=time.monotonic()+45
            while time.monotonic()<deadline:
                settle()
                state=self.bridge.snapshot
                pending=any(timer.isActive() for timer in (self.bridge._load_timer,self.bridge._window_timer,self.bridge._feature_timer))
                if not pending and not state["busy"] and not state["catalogBusy"] and not state["analysisBusy"] and (not state["inspectionVisible"] or state["windowReady"]):break
            self.assertFalse(self.bridge.snapshot["busy"],self.bridge.snapshot["error"])
            self.assertFalse(self.bridge.snapshot["error"],self.bridge.snapshot["error"])
            settle(40)

        def load(self):
            self.bridge.setExperimentFolder(str(self.root));self.bridge.openFiles("sources",[QUrl.fromLocalFile(str(self.source))]);self.wait()

        def test_service_matches_exact_core_pairing_correction_and_dark_references(self):
            dark=self.root/"dark.csv"
            with dark.open("w",newline="") as stream:
                writer=csv.writer(stream);writer.writerow(["B_T","angle_deg",*list(1239.841984/self.energy)]);writer.writerow([0,0,*np.full(len(self.energy),50.)])
            settings=McdSettings(correction_mode="pair_affine",pair_b_alignment="interpolate",dark_pos_file=str(dark),dark_neg_file=str(dark))
            request=McdRequest(str(self.root),str(self.source),settings)
            products=load_mcd(request);expected=process_mcd(str(self.source),settings)
            for key in ("pair_raw_pos","pair_raw_neg","pair_corrected_pos","pair_corrected_neg","pair_mcd_raw","pair_mcd_corrected","pair_b","pair_labels"):
                np.testing.assert_array_equal(getattr(products.result,key),getattr(expected,key))
            self.assertEqual(products.angles,(0.,90.))
            self.assertEqual(len(products.signatures),2)
            self.assertTrue(products.candidates)
            self.assertEqual(set(products.result._quick_mcd_auto_limits),set(products.result.maps))

        def test_native_state_loads_defaults_and_exposes_every_scientific_control(self):
            self.load();state=self.bridge.snapshot
            self.assertTrue(state["inspectionVisible"]);self.assertTrue(state["canExport"])
            self.assertEqual(state["values"]["correction_mode"],"pair_spectral")
            self.assertEqual(state["values"]["spectral_order"],2)
            self.assertTrue(state["values"]["showMean"]);self.assertFalse(state["values"]["showRaw"])
            keys={field["key"] for group in state["parameterGroups"] for field in group["fields"]}
            for key in ("automaticAngles","pos_angle","neg_angle","gain_mode","gain_smooth_window","pair_b_alignment","zero_window_t","max_delta_b","background_ranges_ev","featureMethod","referenceMode","channelMapping","slopeHighPositiveEnd","retain-window","export-retained","pngWidthScale"):
                self.assertIn(key,keys)

        def test_branch_window_and_slopes_match_core_all_metrics(self):
            self.load();products=self.bridge._products
            for metric in ("mean","field_signed_absolute_mean","absolute_mean","integral"):
                snapshot,traces=window_snapshot(products,center=1.7,width=20,metric=metric,ranges=self.bridge._range_settings(),generation=7)
                expected=pair_window_trace_by_branch(products.result,1.7,20)
                for branch in expected:
                    np.testing.assert_array_equal(traces[branch]["corrected_"+metric][1],expected[branch]["corrected_"+metric][1])
                    mask=products.result.pair_labels==branch
                    np.testing.assert_array_equal(snapshot["trace_values"][mask],expected[branch]["corrected_"+metric][1])
                self.assertEqual(len(snapshot["slopes"]["fits"]),6)

        def test_protection_review_preserves_manual_regions_and_never_auto_selects_model(self):
            self.load();self.bridge.invoke("protection")
            deadline=time.monotonic()+20
            while self.bridge._review_result is None and time.monotonic()<deadline:settle()
            self.assertIsNotNone(self.bridge._review_result)
            self.bridge.addProtection(1.76,1.8)
            self.assertIsNone(self.bridge._review_result)
            deadline=time.monotonic()+20
            while (self.bridge._review_result is None or self.bridge._review_timer.isActive() or not self.bridge._review_result.manual_protected_ranges) and time.monotonic()<deadline:settle()
            value=self.bridge._review_result
            expected=suggest_mcd_background_ranges(self.bridge._products.result,protected_ranges_ev=value.requested_protected_ranges+((1.76,1.8),),auto_detect_features=False,use_all_unprotected_bands=True)
            self.assertEqual(value.ranges,expected.ranges)
            self.bridge.invoke("protection-apply");self.wait()
            self.assertEqual(self.bridge._products.request.settings.manual_protected_ranges_ev,((1.76,1.8),))
            self.assertEqual(self.bridge._products.request.settings.spectral_order,2)
            self.assertEqual(self.bridge._products.request.settings.background_selection,"suggested")

        def test_selected_feature_tracks_measured_channels_manual_reference_and_mapping(self):
            self.load();self.bridge.setParameter("candidateFilter","Spectrum")
            candidate=next(item for item in self.bridge._candidates() if item["source"]=="raw pos" and item["kind"]=="peak" and abs(item["energy_ev"]-1.7)<.01)
            self.bridge.setParameter("width",20.);self.bridge.setParameter("featureLow",1.68);self.bridge.setParameter("featureHigh",1.72)
            self.bridge.setParameter("referenceMode","manual");self.bridge.setParameter("manualReference",1.695)
            self.bridge.setParameter("candidate",candidate["id"]);self.wait()
            payload=self.bridge.snapshot["featurePayload"]
            self.assertEqual(payload["status"],"ok",payload.get("reason"))
            self.assertTrue(payload["tracks"]);self.assertEqual(payload["mapping"]["status"],"unknown");self.assertFalse(payload["splitting"])
            for track in payload["tracks"]:
                self.assertEqual(track["reference_method"],"manual")
                measured=getattr(self.bridge._products.result,"pair_b_neg" if track["channel"].endswith("neg") else "pair_b_pos")
                for point in track["points"]:
                    self.assertIn(point["field_t"],measured)
                    if point.get("energy_ev") is not None:self.assertAlmostEqual(point["delta_energy_ev"],point["energy_ev"]-1.695)
            self.bridge.setParameter("channelMapping","fixed_pos");self.wait()
            fixed=self.bridge.snapshot["featurePayload"];self.assertEqual(fixed["mapping"]["status"],"fixed")
            self.assertTrue(fixed["splitting"],fixed.get("splitting_reason"))
            self.bridge.invoke("retain-feature");self.assertFalse(self.bridge.snapshot["error"])
            item=self.bridge._retention.items("feature")[0]
            self.assertTrue(item.computed);self.assertTrue(item.links);self.assertTrue(item.splitting)
            self.bridge.setParameter("featureMethod","Second derivative");self.wait()
            self.assertEqual(self.bridge.snapshot["featurePayload"]["method"],"Second derivative")
            self.assertTrue(self.bridge.snapshot["featurePayload"]["tracks"])

        def test_local_mixed_fit_runs_actual_core_model_on_small_acquisition(self):
            from core.mcd_local_fit import mixed_model
            with self.source.open("w",newline="") as stream:
                writer=csv.writer(stream);writer.writerow(["B_T","angle_deg",*list(1239.841984/self.energy)])
                for field in (-.5,.5):
                    for angle,sign in ((90,1),(0,-1)):
                        raw=mixed_model((self.energy-1.7)*1000.,np.array([10000.,.1,1800.,200.,sign*.5*field,4.]))
                        writer.writerow([field,angle,*raw])
            self.load()
            self.bridge.setParameter("candidateFilter","Spectrum")
            candidate=next(item for item in self.bridge._candidates() if item["source"]=="raw pos" and item["kind"]=="peak" and abs(item["energy_ev"]-1.7)<.01)
            self.bridge.setParameter("featureLow",1.68);self.bridge.setParameter("featureHigh",1.72)
            self.bridge.setParameter("candidate",candidate["id"]);self.bridge.setParameter("featureMethod","Local mixed fit");self.wait()
            local=self.bridge.snapshot["featurePayload"]
            self.assertEqual(local.get("method"),"Local mixed fit",self.bridge.snapshot["status"])
            self.assertTrue(local["tracks"],local.get("reason"))
            self.assertTrue(all(row["tracking_method"]=="Local mixed fit" for row in local["analysis_results"].values()))

        def test_source_catalog_is_unique_portable_and_remembers_injected_filter(self):
            self.load();catalog=self.bridge.snapshot["catalog"]
            self.assertEqual(len(catalog),1);self.assertEqual(catalog[0]["name"],"mcd/synthetic_MCD.csv")
            self.bridge.setParameter("historyFilter","new")
            self.assertEqual(self.settings.value("mcd/source_filter"),"unprocessed")
            other=McdBridge(settings=self.settings)
            try:self.assertEqual(other.snapshot["values"]["historyFilter"],"new")
            finally:other.close()

        def test_cancelled_worker_cannot_replace_completed_source_and_manual_links_clear(self):
            from unittest.mock import patch
            from core.load_cancellation import LoadCancelled
            self.load();old=self.bridge._products;entered=Event();release=Event()
            def slow(request,cancellation=None,previous=None):
                entered.set();release.wait(5)
                if cancellation:cancellation.check()
                return old
            with patch("ui_qt.quick.mcd_workflow.load_mcd",slow):
                self.bridge.invoke("reload")
                deadline=time.monotonic()+5
                while not entered.is_set() and time.monotonic()<deadline:settle()
                self.assertTrue(entered.is_set());self.assertTrue(self.bridge.snapshot["canCancel"])
                self.bridge.invoke("cancel");release.set();settle(100)
                self.assertIs(self.bridge._products,old);self.assertFalse(self.bridge.snapshot["busy"])
            self.bridge._manual_links=[dict(mcd_id="old",spectrum_id="old")]
            self.bridge.openFiles("source",[str(self.source)]);self.wait();self.assertFalse(self.bridge.snapshot["manualLinks"])

        def test_retention_is_immutable_independent_checked_and_stale_guarded(self):
            self.load();self.bridge.invoke("retain-window")
            self.assertFalse(self.bridge.snapshot["error"],self.bridge.snapshot["error"])
            item=self.bridge._retention.items("window")[0]
            before=item.trace_values.copy();self.assertFalse(item.trace_values.flags.writeable)
            self.bridge.setParameter("center",1.78);self.wait()
            np.testing.assert_array_equal(item.trace_values,before)
            self.bridge.invoke("retain-window");self.assertEqual(len(self.bridge._retention.items("window")),2)
            self.bridge.includeRetained("window",item.id,True)
            selected=self.bridge._retention.get_included(source_generation=self.bridge._generation)
            self.assertEqual(len(selected["window"]),1)
            self.bridge.setParameter("correction_mode","global");self.wait()
            with self.assertRaises(ValueError):self.bridge._retention.get_included(source_generation=self.bridge._generation)

        def test_superseded_window_cannot_publish_during_debounce(self):
            from unittest.mock import patch
            self.load();old=(self.bridge._window,self.bridge._traces);entered=Event();release=Event()
            def slow(*args,**kwargs):
                entered.set();release.wait(5);return old
            with patch("ui_qt.quick.mcd_workflow.window_snapshot",slow):
                self.bridge._calculate_window()
                deadline=time.monotonic()+5
                while not entered.is_set() and time.monotonic()<deadline:settle()
                self.assertTrue(entered.is_set())
                self.bridge.setParameter("center",1.78);self.bridge._window_timer.stop()
                release.set();settle(120)
                self.assertIsNone(self.bridge._window)
                self.assertFalse(self.bridge.snapshot["canExport"])
            self.bridge._calculate_window();self.wait()
            self.assertEqual(self.bridge._window["center_ev"],1.78)

        def test_reload_clears_superseded_feature_busy_state(self):
            from unittest.mock import patch
            self.load();entered=Event();release=Event()
            def slow(*args,**kwargs):
                entered.set();release.wait(5);return {"status":"ok","tracks":[]}
            with patch("ui_qt.quick.mcd_workflow.analyze_feature",slow):
                self.bridge.setParameter("candidateFilter","Spectrum");self.bridge._feature_timer.stop()
                self.bridge._analyze_feature()
                deadline=time.monotonic()+5
                while not entered.is_set() and time.monotonic()<deadline:settle()
                self.assertTrue(entered.is_set());self.assertTrue(self.bridge.snapshot["analysisBusy"])
                self.bridge.invoke("reload")
                self.assertFalse(self.bridge.snapshot["analysisBusy"])
                release.set();self.wait()
            self.assertFalse(self.bridge.snapshot["analysisBusy"])

        def test_export_writes_complete_core_package_and_tracks_exact_source(self):
            self.load();self.bridge.invoke("export");self.wait()
            paths=self.bridge.snapshot["resultPaths"]
            self.assertTrue(any(str(path).endswith(".xlsx") for path in paths))
            self.assertTrue(any(str(path).endswith(".dat") for path in paths))
            metadata=json.loads(next(Path(path) for path in paths if str(path).endswith(".json")).read_text())
            self.assertEqual(metadata["workflow"],"MCD")
            self.assertEqual(metadata["source_path"],str(self.source))
            self.assertEqual(metadata["source_relative_path"],"mcd/synthetic_MCD.csv")
            self.assertTrue(metadata["slopes"]);self.assertTrue(metadata["windows"])
            from openpyxl import load_workbook
            workbook=load_workbook(next(path for path in paths if str(path).endswith(".xlsx")),read_only=True)
            self.assertIn("MCD",workbook.sheetnames);self.assertIn("Slopes",workbook.sheetnames);workbook.close()

        def test_custom_export_roots_restore_legacy_history_and_processed_catalog_after_restart(self):
            import hashlib
            from ui_qt.quick.mcd_workflow import history_settings_key
            with TemporaryDirectory() as external:
                external = str(Path(external).resolve())
                self.load();self.bridge.setParameter("center",1.731);self.bridge.setParameter("width",9.);self.wait()
                self.bridge.openFiles("output-folder",[external]);self.bridge.invoke("export");self.wait()
                expected_key="mcd/history_roots/"+hashlib.sha256(str(self.root.resolve()).casefold().encode("utf-8")).hexdigest()
                self.assertEqual(history_settings_key(self.root),expected_key)
                self.assertIn(external,self.settings.value(expected_key,[]))
                deadline=time.monotonic()+10
                while not self.bridge._products.history and time.monotonic()<deadline:settle()
                self.assertTrue(any(abs(row["center_ev"]-1.731)<1e-9 for row in self.bridge._products.history))
                self.assertEqual(self.bridge.snapshot["catalog"][0]["status"],"processed")
                self.bridge.close()
                # Reopen the on-disk INI, rather than sharing in-memory bridge state.
                self.settings=QSettings(self.settings.fileName(),QSettings.IniFormat)
                self.bridge=McdBridge(settings=self.settings);self.addCleanup(self.bridge.close)
                self.load()
                self.assertEqual(self.bridge.snapshot["catalog"][0]["status"],"processed")
                self.assertTrue(any(abs(row["center_ev"]-1.731)<1e-9 and row["width_mev"]==9. for row in self.bridge._products.history))
                self.assertTrue(any(item.get("history") and abs(item["energy_ev"]-1.731)<1e-9 for item in self.bridge._products.candidates))

        def test_input_fingerprint_rejects_same_size_restored_mtime_change(self):
            self.load();products=self.bridge._products
            stat=self.source.stat();data=self.source.read_text();self.source.write_text(data.replace("10000","10001",1))
            os.utime(self.source,ns=(stat.st_atime_ns,stat.st_mtime_ns))
            if self.source.read_text()==data:
                self.source.write_text(data.replace("1.08","1.09",1))
            self.assertNotEqual(fingerprints(input_paths(products.request)),products.signatures)
            self.bridge.invoke("export")
            deadline=time.monotonic()+10
            while self.bridge.snapshot["busy"] and time.monotonic()<deadline:settle()
            self.assertIn("changed",self.bridge.snapshot["error"])

        def test_loaded_cursor_and_window_reuse_map_and_publish_partial_frames(self):
            self.load();settle(60)
            def assert_same_layout_pixels():
                partial=np.asarray(self.bridge._canvas.buffer_rgba()).copy()
                layout=self.bridge._canvas.figure.get_layout_engine()
                self.bridge._canvas.figure.set_layout_engine(None)
                self.bridge._canvas.draw()
                self.bridge._canvas.figure.set_layout_engine(layout)
                full=np.asarray(self.bridge._canvas.buffer_rgba()).copy()
                np.testing.assert_array_equal(partial,full,err_msg="blit differs from full rendering at identical axes positions")
            axes=dict(self.bridge._axes);mesh=axes["map"].collections[0]
            draws=self.bridge._canvas.draw_count;blits=self.bridge._canvas.blit_count
            self.bridge.setParameter("bCursor",1.)
            settle(40)
            self.assertEqual(self.bridge._axes,axes)
            self.assertIs(self.bridge._axes["map"].collections[0],mesh)
            self.assertEqual(self.bridge._canvas.draw_count,draws)
            self.assertGreater(self.bridge._canvas.blit_count,blits)
            from core.mcd_peak_shift import spectrum_energy_order
            index=self.bridge._b_index();order=spectrum_energy_order(self.bridge._products.result)
            np.testing.assert_array_equal(axes["spectra"].lines[0].get_ydata(),self.bridge._products.result.pair_raw_pos[index,order])
            assert_same_layout_pixels()
            self.bridge.setParameter("center",1.78);self.wait()
            self.assertIs(self.bridge._axes["map"].collections[0],mesh)
            self.assertEqual(self.bridge._window["center_ev"],1.78)
            self.bridge.setParameter("candidateFilter","Spectrum");self.wait()
            self.assertIs(self.bridge._axes["map"].collections[0],mesh)
            assert_same_layout_pixels()
            self.bridge.setParameter("xmin",1.55);settle(60)
            self.assertIsNot(self.bridge._axes["map"].collections[0],mesh)

        def test_cursor_spectra_y_follow_current_energy_window_without_marker_bounds(self):
            header=self.source.read_text().splitlines()[0]
            table=np.loadtxt(self.source,delimiter=",",skiprows=1)
            table[:,2:]*=np.where(abs(table[:,0])<.001,100.,1.)[:,None]
            np.savetxt(self.source,table,delimiter=",",header=header,comments="")
            self.load();settle(60)
            axis=self.bridge._axes["spectra"];full=axis.get_ylim()
            axis.set_xlim(1.5,1.56)
            narrowed=axis.get_ylim()
            self.assertLess(narrowed[1]-narrowed[0],(full[1]-full[0])*.4)
            self.bridge._navigation.push_current()
            axis.axvline(1.53,ymin=-1000,ymax=1000)
            self.bridge.setParameter("bCursor",1.);settle(40)
            low,high=axis.get_ylim()
            self.assertGreater(low,9000.);self.assertLess(high,11000.)
            self.assertLess(high,narrowed[1]*.1)
            axis.set_ylim(4.,5.)
            self.bridge._navigation.push_current()
            self.bridge.setParameter("bCursor",-1.);settle(40)
            self.assertEqual(axis.get_ylim(),(4.,5.))
            self.bridge.invoke("home");settle(40)
            self.assertTrue(axis.get_autoscaley_on())
            self.assertGreater(axis.get_ylim()[0],9000.);self.assertLess(axis.get_ylim()[1],11000.)

        def test_color_edits_retain_mesh_zoom_science_and_exact_static_labels(self):
            from matplotlib.colors import Normalize,TwoSlopeNorm
            self.load();settle(60)
            canvas=self.bridge._canvas;axes=dict(self.bridge._axes)
            render=self.bridge._heatmap_render;mesh=render.primary
            cube_values=self.bridge._heatmap_cube.Z.copy()
            def exact_pixels():
                partial=np.asarray(canvas.buffer_rgba()).copy();layout=canvas.figure.get_layout_engine()
                canvas.figure.set_layout_engine(None);canvas.draw();canvas.figure.set_layout_engine(layout)
                np.testing.assert_array_equal(partial,np.asarray(canvas.buffer_rgba()),err_msg="recolor/cursor erased static labels or retained old colors")
            self.bridge.setParameter("bCursor",.5);settle(30)
            self.assertIsNotNone(self.bridge._cursor_blitter)
            axes["map"].set_xlim(1.6,1.8);axes["map"].set_ylim(-1.,1.)
            axes["spectra"].set_xlim(1.62,1.76);canvas.draw()
            limits={key:(axis.get_xlim(),axis.get_ylim()) for key,axis in axes.items()}
            draws=canvas.draw_count;lines={key:tuple(axis.lines) for key,axis in axes.items()}
            self.bridge.setParameter("cmap","viridis");settle(30)
            self.assertEqual(self.bridge._axes,axes);self.assertIs(self.bridge._heatmap_render,render)
            self.assertEqual(canvas.draw_count,draws)
            self.assertIsNone(self.bridge._cursor_blitter)
            for key,axis in axes.items():
                self.assertEqual((axis.get_xlim(),axis.get_ylim()),limits[key]);self.assertEqual(tuple(axis.lines),lines[key])
            self.assertEqual(mesh.cmap.name,"viridis");exact_pixels()
            self.bridge.setParameter("centerZero",False);settle(30)
            self.assertIs(type(mesh.norm),Normalize);exact_pixels()
            self.bridge.setParameter("vmin",-.02);self.bridge.setParameter("vmax",.05)
            self.bridge.setParameter("centerZero",True);settle(30)
            self.assertIsInstance(mesh.norm,TwoSlopeNorm);self.assertEqual(mesh.norm.vcenter,0.)
            self.assertEqual(mesh.get_clim(),(-.02,.05));exact_pixels()
            self.bridge.setParameter("vmax",.2);settle(30)
            np.testing.assert_array_equal(mesh.get_array(),cube_values)
            np.testing.assert_array_equal(self.bridge._heatmap_cube.Z,cube_values)
            self.assertIs(self.bridge._heatmap_render.primary,mesh);exact_pixels()
            self.bridge.setParameter("bCursor",-.5);settle(30)
            self.assertEqual(mesh.cmap.name,"viridis");exact_pixels()
            from unittest.mock import patch
            draws=canvas.draw_count
            with patch.object(self.bridge._plot_blitter,"draw_colors",return_value=False):
                self.bridge.setParameter("cmap","plasma");settle(30)
            self.assertGreater(canvas.draw_count,draws)
            self.assertEqual(self.bridge._axes,axes);self.assertIs(self.bridge._heatmap_render.primary,mesh)
            self.assertEqual(axes["map"].get_xlim(),limits["map"][0]);exact_pixels()

        def test_export_snapshot_preparation_runs_on_worker(self):
            import threading
            from unittest.mock import patch
            self.load();threads=[]
            from ui_qt.quick.mcd_workflow import export_snapshot as original
            def capture(*args,**kwargs):
                threads.append(threading.get_ident())
                return original(*args,**kwargs)
            def stop(*args,**kwargs):raise ValueError("isolated export preparation complete")
            with patch("ui_qt.quick.mcd_workflow.export_snapshot",capture),patch("ui_qt.quick.mcd_workflow.export_mcd",stop):
                self.bridge.invoke("export")
                deadline=time.monotonic()+20
                while self.bridge._busy and time.monotonic()<deadline:settle()
            self.assertTrue(threads)
            self.assertNotEqual(threads[0],threading.get_ident())
            self.assertIn("isolated export preparation complete",self.bridge.snapshot["error"])
            self.assertFalse(self.bridge._exporting)

        def test_real_qml_wide_compact_and_review_have_no_widget_dependency(self):
            self.load();engine=QQmlEngine();warnings=[]
            engine.warnings.connect(lambda messages:warnings.extend(message.toString() for message in messages))
            path=Path(__file__).resolve().parents[1]/"ui_qt"/"quick"/"qml"/"McdWorkspace.qml"
            component=QQmlComponent(engine,QUrl.fromLocalFile(str(path)))
            item=component.createWithInitialProperties(dict(bridge=self.bridge,colors=colors))
            self.assertIsNotNone(item,"\n".join(error.toString() for error in component.errors()))
            pane_component=QQmlComponent(engine,QUrl.fromLocalFile(str(path.with_name("McdParameters.qml"))))
            pane=pane_component.createWithInitialProperties(dict(bridge=self.bridge,colors=colors))
            self.assertIsNotNone(pane,"\n".join(error.toString() for error in pane_component.errors()))
            window=QQuickWindow();window.resize(1440,780);pane.setParentItem(window.contentItem());pane.setWidth(240);pane.setHeight(780)
            item.setParentItem(window.contentItem());item.setX(240);item.setWidth(1200);item.setHeight(780);window.show()
            settle(80)
            recommendations=item.findChild(QObject,"mcd_peak_center_recommendations")
            self.assertIsNotNone(recommendations)
            self.assertTrue(recommendations.property("visible"))
            centers=self.bridge.snapshot["recommendedCenters"]
            self.assertTrue(centers)
            button=next((child for child in recommendations.childItems()
                         if child.objectName()=="mcd_recommended_center_"+centers[-1]["id"]),None)
            self.assertIsNotNone(button)
            self.bridge.setParameter("candidateFilter","Spectrum")
            self.assertEqual([row["id"] for row in self.bridge.snapshot["recommendedCenters"]],[row["id"] for row in centers])
            button.clicked.emit();self.wait()
            self.assertEqual(self.bridge._values["center"],centers[-1]["energy_ev"])
            self.assertEqual(self.bridge._values["candidateFilter"],"Spectrum")
            surface=item.findChild(QObject,"mcd_plot_surface");self.assertIsNotNone(surface)
            def plot_ready(axes):
                canvas=self.bridge._canvas
                size=(max(1,round(surface.width()*window.devicePixelRatio())),max(1,round(surface.height()*window.devicePixelRatio())))
                return canvas is not None and not surface._size_pending and not surface._size_timer.isActive() and canvas.get_width_height(physical=True)==size and not self.bridge._plot_timer.isActive() and not canvas.backend._timer.isActive() and set(self.bridge._axes)==set(axes)
            def wait_plot(axes):
                # Viewport polish and the queued Agg render must both finish.
                deadline=time.monotonic()+10
                while not plot_ready(axes) and time.monotonic()<deadline:settle()
                self.assertTrue(plot_ready(axes),f"Plot did not settle: expected {axes}, actual {tuple(self.bridge._axes)}")
            wait_plot({"map","spectra","mcd","trace","feature"})
            self.assertFalse(window.grabWindow().isNull())
            capture=os.environ.get("DPTK_MCD_CAPTURE")
            if capture:window.grabWindow().save(capture+"-wide.png")
            self.bridge.invoke("protection")
            deadline=time.monotonic()+20
            while self.bridge._review_result is None and time.monotonic()<deadline:settle()
            self.assertIsNotNone(self.bridge._review_result);settle(180)
            if capture:window.grabWindow().save(capture+"-protection.png")
            self.bridge.invoke("close-sheet");pane.setWidth(220);pane.setHeight(600);item.setX(220);item.setWidth(500);item.setHeight(600);window.resize(720,600);wait_plot({"map","spectra"})
            self.assertTrue(self.bridge.snapshot["compactInspection"])
            self.assertEqual(set(self.bridge._axes),{"map","spectra"})
            self.assertEqual(self.bridge._values["heatmapFraction"],.7)
            headers=[axis for axis in self.bridge._canvas.figure.axes if getattr(axis,"_dptk_inspection_header",False)]
            self.assertEqual(len(headers),1)
            self.assertTrue(all(child._colorbar.orientation=="horizontal" for child in headers[0].child_axes))
            self.assertEqual(self.bridge._axes["spectra"].title.get_position(),(.02,.97))
            if capture:window.grabWindow().save(capture+"-compact-map.png")
            previous=self.bridge._values["heatmapFraction"]
            self.bridge.setParameter("dividerY",self.bridge.snapshot["dividerY"]-.08);wait_plot({"map","spectra"})
            self.assertLess(self.bridge._values["heatmapFraction"],previous)
            self.bridge.invoke("toggle-spectrum");wait_plot({"map"});self.assertEqual(set(self.bridge._axes),{"map"})
            self.bridge.invoke("toggle-spectrum");self.bridge.setParameter("heatmapFraction",.7);wait_plot({"map","spectra"})
            self.assertEqual(set(self.bridge._axes),{"map","spectra"})
            for panel in ("map","spectra","mcd","trace","feature"):
                self.bridge.setParameter("panel",panel);wait_plot({"map","spectra"} if panel=="map" else {panel})
            self.assertFalse(warnings,"\n".join(warnings))
            self.assertNotIn("PySide6.QtWidgets",sys.modules);self.assertNotIn("matplotlib.backends.backend_qtagg",sys.modules)
            if capture:window.grabWindow().save(capture+"-compact.png")
            window.close();item.setParentItem(None);pane.setParentItem(None);item.deleteLater();pane.deleteLater();settle()

    suite=unittest.TestSuite([Cases("test_real_qml_wide_compact_and_review_have_no_widget_dependency")]) if "--qml-only" in sys.argv else unittest.defaultTestLoader.loadTestsFromTestCase(Cases)
    result=unittest.TextTestRunner(verbosity=2).run(suite)
    return 0 if result.wasSuccessful() else 1


if __name__=="__main__":
    if "--isolated" in sys.argv:raise SystemExit(_isolated())
    unittest.main()
