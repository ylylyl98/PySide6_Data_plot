import unittest
import tempfile
from pathlib import Path
from dataclasses import replace
import numpy as np

from core.loader import DataCube
from core import drr_analysis_workspace as workspace


class WorkspaceTests(unittest.TestCase):
    def dataset(self, name='one'):
        x = np.linspace(1, 1.1, 101)
        peak = np.exp(-((x-1.05)/.004)**2)
        cube = DataCube(x, np.array([0., 1., 2.]), np.array([peak, peak*0, peak]), 'Gate (V)', name, 'DR/R')
        return workspace.create_dataset(cube, name, {'source': name})

    def test_snapshot_identity_and_full_ranges(self):
        dataset = self.dataset()
        original = dataset.cube
        copy = workspace.create_dataset(original, 'renamed', {'source': 'one'})
        self.assertEqual(dataset.key, copy.key)
        self.assertFalse(np.shares_memory(original.Z, copy.cube.Z))
        self.assertFalse(copy.cube.Z.flags.writeable)
        self.assertEqual((copy.settings.x_min, copy.settings.x_max, copy.settings.y_min, copy.settings.y_max), (1., 1.1, 0., 2.))
        changed = replace(original, Z=original.Z+1)
        self.assertNotEqual(dataset.key, workspace.create_dataset(changed, 'one', {'source': 'one'}).key)

    def test_settings_invalidation_and_batch_preserves_seeds(self):
        source, target = self.dataset(), self.dataset('two')
        target.settings = replace(target.settings, seed_energy=1.08, seed_y=2.)
        target.result = {'old': True}
        self.assertFalse(workspace.apply_settings(target, target.settings))
        self.assertEqual(target.revision, 0)
        source.settings = replace(source.settings, prominence=.2, seed_energy=1.03, seed_y=1.)
        self.assertEqual(workspace.apply_common_settings(source, [source, target]), 1)
        self.assertIsNone(target.result)
        self.assertEqual(target.revision, 1)
        self.assertEqual((target.settings.seed_energy, target.settings.seed_y), (1.08, 2.))
        self.assertEqual(target.settings.prominence, .2)

    def test_analysis_is_unassigned_and_has_dataset_provenance(self):
        dataset = self.dataset()
        dataset.settings = replace(dataset.settings, source='raw', polarity='peaks')
        result = workspace.analyze_dataset(dataset)
        self.assertIsNone(dataset.result)
        self.assertEqual(result['dataset_id'], dataset.key)
        self.assertEqual(result['provenance'], {'source': 'one'})
        self.assertEqual([p['row_index'] for p in result['products']['raw']['points']], [0, 2])

    def test_loading_explicit_background_recipe_without_writes(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root / 'sample_TG-1.087BG=18.csv'
            path.write_text('Vbg,Vtg,700,710,720,730,740\n0,18,2,4,6,4,2\n1,19.087,4,8,12,8,4\n')
            bg = root / 'background.csv'
            bg.write_text('Vbg,Vtg,700,710,720,730,740\n0,0,1,2,3,2,1\n0,0,3,6,9,6,3\n')
            before = {p.name: p.read_bytes() for p in root.iterdir()}
            external = workspace.load_dataset(path, [bg])
            first = workspace.load_dataset(path, [], background_mode='self_first')
            last = workspace.load_dataset(path, [], background_mode='self_last')
            self.assertEqual(external.provenance['background_which'], 'all')
            np.testing.assert_allclose(external.cube.Z, [[0]*5, [1]*5])
            np.testing.assert_allclose(first.cube.Z, [[0]*5, [1]*5])
            np.testing.assert_allclose(last.cube.Z, [[-.5]*5, [0]*5])
            self.assertNotEqual(external.cube.gate_label, 'BG (V)')
            self.assertEqual(before, {p.name: p.read_bytes() for p in root.iterdir()})
            with self.assertRaises(ValueError):
                workspace.load_dataset(path, [])
            bg.write_text('Vbg,Vtg,700,710,720,730,741\n0,0,1,2,3,2,1\n')
            with self.assertRaises(ValueError):
                workspace.load_dataset(path, [bg])

    def test_summary_numeric_tables_keep_missing_rows_and_reject_stale(self):
        from openpyxl import load_workbook
        datasets = [self.dataset(), self.dataset('two')]
        for dataset in datasets:
            dataset.settings = replace(dataset.settings, source='raw', polarity='peaks')
            dataset.result = workspace.analyze_dataset(dataset)
        with tempfile.TemporaryDirectory() as tmp:
            path = workspace.export_summary(tmp, datasets)
            book = load_workbook(path)
            try:
                info = list(book['Peak_Info'].values)
                self.assertEqual(info[0][:3], ('dataset_id', 'dataset_name', 'product'))
                self.assertEqual({row[0] for row in info[1:]}, {d.key for d in datasets})
                self.assertIsInstance(info[1][5], (int, float))
                wide = list(book['D001_DRR_Peaks'].values)
                self.assertEqual(wide[2], (1., None, None))
                self.assertEqual(len(wide), 4)
                self.assertIn('Parameters', book.sheetnames)
            finally:
                book.close()
            outputs = workspace.export_dataset(tmp, datasets[0])
            self.assertEqual(set(outputs), {'xlsx', 'json', 'png'})
            self.assertTrue(all(p.is_file() for p in outputs.values()))
            datasets[0].revision += 1
            with self.assertRaises(ValueError):
                workspace.export_summary(tmp, datasets)
            with self.assertRaises(ValueError):
                workspace.export_dataset(tmp, datasets[0])
