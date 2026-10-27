import csv
from dataclasses import replace
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import numpy as np
from openpyxl import load_workbook

from core.export import SHG_RESULT_FIELDS, export_shg_results
from core.load_cancellation import LoadCancelled
from core.shg import ShgSettings, process_shg_sweep
from core.shg_fit import ShgFitSettings, evaluate_shg_angular_model, fit_shg_angular_result, fit_shg_twist_comparison
from core.shg_history import scan_shg_history
from core.shg_workbook import export_shg_workbook
from tests.test_shg import _synthetic_data
from tests.test_shg_fit import _angular_sweep


def read_sheets(path):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        output = {}
        for sheet in book:
            rows = list(sheet.values)
            width = len(rows[0])
            # Streaming XLSX omits trailing blank cells; recover the rectangular
            # table implied by its header when comparing independent XY blocks.
            output[sheet.title] = [tuple(row) + (None,) * (width - len(row)) for row in rows]
        return output
    finally:
        book.close()


class ShgWorkbookTests(unittest.TestCase):
    def test_overview_identifies_same_named_sources_and_backgrounds_without_json(self):
        data, _ = _synthetic_data()
        settings = ShgSettings(background_method="external")
        with tempfile.TemporaryDirectory() as folder:
            results = []
            for name in ("reference", "sample"):
                source = replace(data, source_file=str(Path(folder) / name / "measurement.csv"))
                background = replace(data, source_file=str(Path(folder) / name / "background.csv"),
                                     spectra=data.spectra * .1)
                results.append(process_shg_sweep(source, settings, background=background))
            paths = export_shg_workbook(folder, results=results, include_csv=True)
            metadata = json.loads(paths["settings"].read_text(encoding="utf-8"))
            sheets = read_sheets(paths["workbook"])
            self.assertEqual(next(iter(sheets)), "Overview")
            rows = {(row[0], row[1]): row for row in sheets["Overview"][1:]}
            for role, label in (("reference", "Reference A"), ("sample", "Sample B")):
                row = rows["Source files", label]
                self.assertEqual(row[2], "measurement.csv")
                self.assertEqual(row[3], str((Path(folder) / role / "measurement.csv").resolve()))
                background_row = rows["Background files", label]
                self.assertEqual(background_row[2], "background.csv")
                self.assertEqual(background_row[3], str((Path(folder) / role / "background.csv").resolve()))
                self.assertEqual(rows["Processing", label + " background"][2], "external")
                self.assertIn(role.title() + "_Raw", rows["Data mapping", label][2])
            self.assertEqual(rows["Analysis", "Exported (UTC)"][2], metadata["created_utc"])
            self.assertEqual({row[1] for row in sheets["Overview"][1:] if row[0] == "Worksheet guide"},
                             set(sheets))
            with paths["csv_overview"].open(newline="", encoding="utf-8") as handle:
                csv_rows = list(csv.reader(handle))
            self.assertEqual(csv_rows[0], list(sheets["Overview"][0]))
            self.assertEqual(csv_rows[2][2:], list(sheets["Overview"][2][2:]))
            # A copy of the workbook alone still carries both absolute identities.
            copied = Path(folder) / "shared.xlsx"
            copied.write_bytes(paths["workbook"].read_bytes())
            self.assertEqual(read_sheets(copied)["Overview"], sheets["Overview"])

    def test_single_overview_is_active_and_long_paths_wrap(self):
        data, _ = _synthetic_data()
        with tempfile.TemporaryDirectory() as folder:
            source = "nested/" + "long_measurement_" * 15 + ".csv"
            data = replace(data, source_file=source)
            result = process_shg_sweep(data, ShgSettings(background_method="local_quadratic"))
            paths = export_shg_workbook(folder, results=(result,))
            self.assertEqual({path.name for path in paths.values()}, {"shg.xlsx", "analysis_settings.json"})
            book = load_workbook(paths["workbook"], data_only=True)
            try:
                self.assertEqual(book.active.title, "Overview")
                sheet = book["Overview"]
                self.assertEqual(sheet["B3"].value, "Measurement")
                self.assertEqual(sheet["C3"].value, Path(source).name)
                self.assertEqual(sheet["D3"].value, str((Path(folder) / source).resolve()))
                self.assertTrue(sheet["D3"].alignment.wrap_text)
                self.assertGreater(sheet.row_dimensions[3].height, 30)
                rows = {(row[0], row[1]): row for row in sheet.values if row[0] != "Section"}
                self.assertEqual(rows["Processing", "Measurement background"][2], "local_quadratic")
                self.assertNotIn(("Source files", "Sample B"), rows)
            finally:
                book.close()

    def test_transposed_all_angle_spectra_keep_duplicate_angles_and_source_rows(self):
        data, _ = _synthetic_data()
        data = replace(data, measured_angle_deg=np.array([360., 0., 5.]),
                       target_angle_deg=np.array([360., 0., 5.]))
        result = process_shg_sweep(data, ShgSettings(background_method="local_quadratic", angle_wrap_deg=360.))
        with tempfile.TemporaryDirectory() as folder:
            paths = export_shg_workbook(folder, results=(result,))
            self.assertEqual({path.name for path in paths.values()}, {"shg.xlsx", "analysis_settings.json"})
            self.assertEqual(len(list(paths["workbook"].parent.iterdir())), 2)
            sheets = read_sheets(paths["workbook"])
            payload = json.loads(paths["settings"].read_text(encoding="utf-8"))
        raw = sheets["Spectra_Raw"]
        self.assertEqual(raw[0], ("Wavelength_nm", "0 deg [row 2]", "0 deg [row 3]", "5 deg [row 4]"))
        np.testing.assert_allclose(np.asarray(raw[1:], float)[:, 0], data.wavelength_nm)
        np.testing.assert_allclose(np.asarray(raw[1:], float)[:, 1:], data.spectra.T)
        # Failed motion rows remain in spectra and diagnostics, but not Plot.
        self.assertEqual(len(sheets["Plot"]), 3)
        self.assertEqual(len(sheets["Diagnostics"]), 4)
        columns = payload["sheets"]["Spectra_Raw"]["columns"]
        self.assertEqual([item["source_row"] for item in columns], [2, 3, 4])
        self.assertEqual([item["raw_measured_angle_deg"] for item in columns], [360., 0., 5.])
        self.assertFalse(columns[2]["included"])
        self.assertEqual(payload["sheets"]["Spectra_Raw"]["intensity_units"], "counts")

    def test_corrected_window_and_optional_full_cleaned_spectra_are_exact(self):
        data, _ = _synthetic_data()
        result = process_shg_sweep(data, ShgSettings(background_method="local_quadratic"))
        with tempfile.TemporaryDirectory() as folder:
            paths = export_shg_workbook(folder, results=(result,), include_cleaned=True, include_csv=True)
            sheets = read_sheets(paths["workbook"])
            order = np.argsort(result.measured_angle_deg, kind="stable")
            mask = ((data.wavelength_nm >= result.settings.left_min_nm)
                    & (data.wavelength_nm <= result.settings.right_max_nm))
            corrected = np.asarray(sheets["Spectra_Corrected"][1:], float)
            np.testing.assert_allclose(corrected[:, 0], data.wavelength_nm[mask])
            np.testing.assert_allclose(corrected[:, 1:], result.corrected[order][:, mask].T)
            cleaned = np.asarray(sheets["Spectra_Cleaned"][1:], float)
            np.testing.assert_allclose(cleaned[:, 1:], result.cleaned_spectra[order].T)
            self.assertEqual(len(paths), 2 + len(sheets))
            for name, rows in sheets.items():
                with paths["csv_" + name.casefold()].open(newline="", encoding="utf-8") as handle:
                    csv_rows = list(csv.reader(handle))
                self.assertEqual(csv_rows[0], list(rows[0]))
                self.assertEqual(len(csv_rows), len(rows))
            with paths["csv_spectra_corrected"].open(newline="", encoding="utf-8") as handle:
                csv_values = np.asarray(list(csv.reader(handle))[1:], float)
            np.testing.assert_allclose(csv_values, corrected)
            none = process_shg_sweep(data, ShgSettings(background_method="none"))
            none_paths = export_shg_workbook(folder, results=(none,), processed_name="none")
            none_sheets = read_sheets(none_paths["workbook"])
            self.assertEqual(len(none_sheets["Spectra_Corrected"]), len(data.wavelength_nm) + 1)

    def test_compare_keeps_independent_angle_and_wavelength_axes(self):
        _first_data, first = _angular_sweep(8.)
        second_data, _ = _angular_sweep(14., angles=np.linspace(.5, 179.5, 31))
        second_data = replace(second_data, wavelength_nm=second_data.wavelength_nm[::2],
                              spectra=second_data.spectra[:, ::2])
        second = process_shg_sweep(second_data, first.settings)
        twist = fit_shg_twist_comparison(first, second)
        with tempfile.TemporaryDirectory() as folder:
            paths = export_shg_workbook(folder, results=(first, second), twist=twist)
            sheets = read_sheets(paths["workbook"])
            payload = json.loads(paths["settings"].read_text(encoding="utf-8"))
        plot = sheets["Plot"]
        self.assertEqual(plot[0], ("Reference_Angle_deg", "Reference_Intensity_counts_nm", "Reference_Error_counts_nm",
                                   "Sample_Angle_deg", "Sample_Intensity_counts_nm", "Sample_Error_counts_nm"))
        self.assertEqual(plot[1][0], 0.)
        self.assertEqual(plot[1][3], .5)
        self.assertEqual(plot[-1][3:], (None, None, None))
        self.assertEqual(len(sheets["Reference_Raw"]), 322)
        self.assertEqual(len(sheets["Sample_Raw"]), 162)
        np.testing.assert_allclose(np.asarray(sheets["Sample_Raw"][1:], float)[:, 1:], second_data.spectra.T)
        self.assertAlmostEqual(payload["twist_fit"]["absolute_twist_angle_deg"], 4., delta=.05)
        self.assertEqual(set(payload["angular_fits"]), {"reference", "sample"})

    def test_fit_samples_its_actual_fit_mask_including_explicitly_allowed_excluded_rows(self):
        data, result = _angular_sweep(12.)
        included = result.included.copy()
        included[0] = False
        result = replace(result, included=included)
        settings = ShgFitSettings(include_excluded_rows=True, angle_max_deg=120.)
        fit = fit_shg_angular_result(result, settings)
        self.assertTrue(fit.fit_mask[0])
        with tempfile.TemporaryDirectory() as folder:
            paths = export_shg_workbook(folder, results=(result,), fits=(fit,), fit_settings=settings)
            sheets = read_sheets(paths["workbook"])
        self.assertEqual(sheets["Plot"][1][0], 5.)
        curve = np.asarray(sheets["Fit"][1:], float)
        self.assertEqual(curve.shape, (721, 2))
        self.assertEqual(curve[0, 0], 0.)
        self.assertEqual(curve[-1, 0], 120.)
        np.testing.assert_allclose(curve[:, 1], evaluate_shg_angular_model(curve[:, 0], fit.i0, fit.amplitude, fit.x_center_deg))

    def test_diagnostics_preserve_legacy_csv_fields_and_settings_keep_history_identity(self):
        data, _ = _synthetic_data()
        result = process_shg_sweep(data, ShgSettings(background_method="local_quadratic"))
        with tempfile.TemporaryDirectory() as folder:
            (Path(folder) / data.source_file).write_text("source identity", encoding="utf-8")
            paths = export_shg_workbook(folder, results=(result,), processed_name="Processed Data/SHG/test")
            old = export_shg_results(folder, data=data, result=result, settings=result.settings, processed_name="legacy")
            with old["csv"].open(newline="", encoding="utf-8") as handle:
                csv_rows = list(csv.DictReader(handle))
            sheets = read_sheets(paths["workbook"])
            diagnostic = [dict(zip(sheets["Diagnostics"][0], row)) for row in sheets["Diagnostics"][1:]]
            payload = json.loads(paths["settings"].read_text(encoding="utf-8"))
            history = scan_shg_history(folder, [data.source_file])
            self.assertIn(data.source_file, history.processed_at)
            self.assertEqual(payload["inputs"][0]["source_relative_path"], data.source_file)
            self.assertEqual(payload["inputs"][0]["source_path"], str((Path(folder) / data.source_file).resolve()))
        self.assertEqual(set(SHG_RESULT_FIELDS), set(csv_rows[0]))
        for expected, actual in zip(csv_rows, diagnostic):
            for key in SHG_RESULT_FIELDS:
                if expected[key] == "":
                    self.assertIsNone(actual[key])
                elif isinstance(actual[key], (int, float)):
                    self.assertAlmostEqual(float(expected[key]), actual[key], delta=max(1e-10, abs(actual[key]) * 1e-11))
                else:
                    self.assertEqual(expected[key], actual[key])
        self.assertEqual(payload["processing_settings"], result.settings.to_dict())
        self.assertIn("fit_settings", payload)
        self.assertEqual({item["filename"] for item in payload["output_manifest"]}, {"shg.xlsx", "analysis_settings.json"})

    def test_disabled_and_partial_fits_are_explicit_without_fabricated_results(self):
        _, first = _angular_sweep(8.)
        _, second = _angular_sweep(14., angles=np.array([0., 5., 10.]))
        fit = fit_shg_angular_result(first)
        with tempfile.TemporaryDirectory() as folder:
            paths = export_shg_workbook(folder, results=(first, second), fits=(fit, None),
                                       fit_settings=ShgFitSettings(), fit_message="Sample: insufficient fit points")
            sheets = read_sheets(paths["workbook"])
            summary = {(row[0], row[1]): row for row in sheets["Summary"][1:]}
            self.assertEqual(summary["reference", "fit_status"][2], "success")
            self.assertEqual(summary["sample", "fit_status"][2], "unavailable")
            self.assertIsNone(summary["sample", "amplitude"][2])
            self.assertIsNone(summary["comparison", "absolute_twist_angle_deg"][2])
            self.assertEqual(len(sheets["Fit"][0]), 2)
            disabled = export_shg_workbook(folder, results=(first,), fit_settings=ShgFitSettings(enabled=False),
                                          processed_name="disabled")
            disabled_sheets = read_sheets(disabled["workbook"])
            self.assertEqual(len(disabled_sheets["Fit"]), 1)
            self.assertEqual(disabled_sheets["Summary"][1][2], "disabled")

    def test_nonfinite_values_are_blank_and_source_strings_are_not_formulas(self):
        data, _ = _synthetic_data()
        result = process_shg_sweep(data, ShgSettings(background_method="local_quadratic"))
        raw = data.spectra.copy()
        raw[0, 0] = np.nan
        data = replace(data, source_file="=SUM(1,2)\x01.csv", spectra=raw)
        result = replace(result, data=data)
        with tempfile.TemporaryDirectory() as folder:
            paths = export_shg_workbook(folder, results=(result,))
            book = load_workbook(paths["workbook"], read_only=True, data_only=False)
            try:
                rows = list(book["Diagnostics"].rows)
                source_index = [cell.value for cell in rows[0]].index("source_csv")
                self.assertEqual(rows[1][source_index].data_type, "s")
                self.assertEqual(rows[1][source_index].value, "=SUM(1,2)\\u0001.csv")
                # The first acquisition is C because the negative angle sorts first.
                raw_rows = list(book["Spectra_Raw"].values)
                self.assertIsNone(raw_rows[1][2])
            finally:
                book.close()

    def test_failure_and_cancellation_remove_only_new_files_and_never_overwrite(self):
        data, _ = _synthetic_data()
        result = process_shg_sweep(data, ShgSettings(background_method="local_quadratic"))
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "Processed Data"
            output.mkdir()
            marker = output / "keep.txt"
            marker.write_text("keep", encoding="utf-8")
            with patch("openpyxl.Workbook.save", side_effect=OSError("disk full")):
                with self.assertRaisesRegex(OSError, "disk full"):
                    export_shg_workbook(folder, results=(result,), include_csv=True)
            self.assertEqual(list(output.iterdir()), [marker])
            checks = 0

            def cancel_midway():
                nonlocal checks
                checks += 1
                if checks > 20:
                    raise LoadCancelled()

            with patch("core.shg_workbook.check_load_cancelled", side_effect=cancel_midway):
                with self.assertRaises(LoadCancelled):
                    export_shg_workbook(folder, results=(result,), include_csv=True)
            self.assertEqual(list(output.iterdir()), [marker])
            paths = export_shg_workbook(folder, results=(result,))
            originals = {key: path.read_bytes() for key, path in paths.items()}
            with self.assertRaises(FileExistsError):
                export_shg_workbook(folder, results=(result,))
            self.assertEqual({key: path.read_bytes() for key, path in paths.items()}, originals)
            self.assertEqual(marker.read_text(encoding="utf-8"), "keep")

    def test_xlsx_limits_reject_before_writing_instead_of_truncating(self):
        data, _ = _synthetic_data()
        result = process_shg_sweep(data, ShgSettings(background_method="local_quadratic"))
        with tempfile.TemporaryDirectory() as folder, patch("core.shg_workbook._MAX_COLUMNS", 3):
            with self.assertRaisesRegex(ValueError, "exceeds the XLSX limit"):
                export_shg_workbook(folder, results=(result,))
            self.assertEqual(list(Path(folder).iterdir()), [])


if __name__ == "__main__":
    unittest.main()
