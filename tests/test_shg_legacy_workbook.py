"""Exercise the legacy worker's SHG package transaction with real workbooks."""
from dataclasses import replace
import json
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from core.application_models import ExportOptions, LoadedState
from core.shg_fit import ShgFitSettings, fit_shg_angular_result
from core.shg_history import scan_shg_history
from tests.test_shg_fit import _angular_sweep
from ui_qt.main_window import MainWindow


class ShgLegacyWorkbookTests(unittest.TestCase):
    def _state(self, folder, *, compare=False):
        first_data, first = _angular_sweep(12.)
        second_data, second = _angular_sweep(19.)
        for data in (first_data, second_data):
            (Path(folder) / data.source_file).write_text("source retained\n", encoding="utf-8")
        fit_settings = ShgFitSettings()
        return LoadedState(
            mode="SHG Processing", folder=folder,
            selected_files=[first_data.source_file, *([second_data.source_file] if compare else [])],
            shg_data=first_data, shg_result=first, shg_fit=fit_shg_angular_result(first, fit_settings),
            shg_data_b=second_data if compare else None, shg_result_b=second if compare else None,
            shg_fit_b=None, shg_compare=compare, shg_settings=first.settings,
            shg_fit_settings=fit_settings,
        )

    def _export(self, loaded):
        # The loaded processing snapshot remains authoritative even if a UI
        # options object is stale; no scientific work should occur while saving.
        options = ExportOptions(mode="SHG Processing", params=None,
            shg_settings=replace(loaded.shg_settings, peak_center_nm=516.),
            shg_fit_settings=ShgFitSettings(enabled=False))
        signal = SimpleNamespace(emit=lambda *_args: None)
        return MainWindow._export_task(None, loaded, options, progress=signal, log=signal)

    def test_default_package_uses_processed_snapshot_and_retains_sources(self):
        with tempfile.TemporaryDirectory() as folder:
            loaded = self._state(folder)
            result = self._export(loaded)
            package = Path(result["out_folder"])
            self.assertEqual({path.name for path in package.iterdir()}, {"shg.xlsx", "analysis_settings.json"})
            self.assertIn("_515nm_", package.name)
            metadata = json.loads((package / "analysis_settings.json").read_text(encoding="utf-8"))
            self.assertEqual(metadata["processing_settings"]["peak_center_nm"], 515.)
            self.assertTrue(metadata["fit_settings"]["enabled"])
            with (package / "shg.xlsx").open("rb") as stream:
                book = load_workbook(stream, read_only=True, data_only=True)
                try:
                    self.assertIn("Spectra_Raw", book.sheetnames)
                    self.assertIn("Spectra_Corrected", book.sheetnames)
                    rows = book["Spectra_Raw"].iter_rows(values_only=True)
                    next(rows)
                    first_row = next(rows)
                    self.assertEqual(first_row[0], loaded.shg_data.wavelength_nm[0])
                    self.assertAlmostEqual(first_row[1], loaded.shg_data.spectra[0, 0])
                finally:
                    book.close()
            self.assertEqual(result["source_files"], loaded.selected_files)
            self.assertEqual((Path(folder) / loaded.selected_files[0]).read_text(), "source retained\n")
            history = scan_shg_history(folder, loaded.selected_files)
            self.assertIn(loaded.selected_files[0], history.processed_at)

    def test_partial_comparison_keeps_both_sources_and_the_valid_fit(self):
        with tempfile.TemporaryDirectory() as folder:
            loaded = self._state(folder, compare=True)
            result = self._export(loaded)
            package = Path(result["out_folder"])
            self.assertEqual(len(list(package.iterdir())), 2)
            with (package / "shg.xlsx").open("rb") as stream:
                book = load_workbook(stream, read_only=True, data_only=True)
                try:
                    self.assertIn("Reference_Raw", book.sheetnames)
                    self.assertIn("Sample_Raw", book.sheetnames)
                    self.assertIn("Fit", book.sheetnames)
                finally:
                    book.close()
            history = scan_shg_history(folder, loaded.selected_files)
            self.assertEqual(set(history.processed_at), set(loaded.selected_files))
            self.assertIn("reference", history.roles[loaded.selected_files[0]])
            self.assertIn("sample", history.roles[loaded.selected_files[1]])

    def test_failure_cleans_staging_and_does_not_publish_history(self):
        def fail(_folder, *, processed_name, **_kwargs):
            staged = Path(processed_name)
            staged.mkdir(parents=True)
            (staged / "shg.xlsx").write_bytes(b"incomplete")
            raise OSError("simulated write failure")

        with tempfile.TemporaryDirectory() as folder:
            loaded = self._state(folder)
            with patch("core.shg_workbook.export_shg_workbook", side_effect=fail):
                with self.assertRaisesRegex(OSError, "simulated write failure"):
                    self._export(loaded)
            self.assertEqual(list((Path(folder) / "Processed Data" / "SHG").iterdir()), [])
            self.assertEqual(list(Path(folder).glob(".shg-export-*")), [])
            self.assertEqual(scan_shg_history(folder, loaded.selected_files).processed_at, {})


if __name__ == "__main__":
    unittest.main()
