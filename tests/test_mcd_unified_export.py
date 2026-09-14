from __future__ import annotations

import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from core.mcd_unified_export import (
    _mcd_data,
    _selected_map_data,
    _plot_pngs,
    build_mcd_export_snapshot,
    export_mcd_analysis,
    freeze_mcd_export_snapshot,
)
from core.loader import load_dat
from core.mcd import McdSettings, discover_mcd_processing_status, process_mcd
from core.mcd_analysis import fit_mcd_slopes
from core.mcd_peak_shift import PeakCandidate, PeakPoint, PeakShiftResult, PeakTrack


def _snapshot() -> dict:
    energy = np.asarray([1.0, 1.5, 2.0], float)
    fields = np.asarray([-2.0, -1.0, 0.0, 1.0, 2.0], float)
    branches = np.asarray(
        ["B decreasing", "B decreasing", "B increasing", "B increasing", "B increasing"],
        dtype=str,
    )
    values = np.asarray(
        [[-2.0, -1.0, 0.0], [-1.0, -0.5, 0.0], [0.0, 0.0, 0.0], [1.0, 0.5, 0.0], [2.0, 1.0, 0.0]],
        float,
    )
    return {
        "source_descriptor": {
            "name": "sample_P3_T1.67K_D-0.2V_F0.1V_Bsweep_repeat02.csv",
            "filename": "sample_P3_T1.67K_D-0.2V_F0.1V_Bsweep_repeat02.csv",
            "path": "/experiment/raw/sample_P3_T1.67K_D-0.2V_F0.1V_Bsweep_repeat02.csv",
            "relative_path": "raw/sample_P3_T1.67K_D-0.2V_F0.1V_Bsweep_repeat02.csv",
            "repeat": "repeat02",
        },
        "mcd": {
            "energy_ev": energy,
            "fields_t": fields,
            "branches": branches,
            "values": values,
            "metric": "corrected signed mean",
            "channel": "MCD",
        },
        "spectra": {
            "K_raw": {
                "energy_ev": energy,
                "fields_t": np.asarray([-2.0, 0.0, 2.0]),
                "values": np.asarray([[-1.0, 1.0, 2.0], [-1.0, 1.1, 2.1], [-1.0, 1.2, 2.2]]),
                "unit": "a.u.",
                "branch": "B increasing",
            },
            "Kp_raw": {
                "energy_ev": energy,
                "fields_t": np.asarray([-1.0, 1.0]),
                "values": np.asarray([[-2.0, 0.0, 3.0], [-2.1, 0.1, 3.1]]),
                "unit": "a.u.",
                "branch": "B decreasing",
            },
        },
        "features": [
            {"feature_id": "peak-1", "feature_kind": "peak", "branch": "B increasing", "energy_ev": 1.5, "delta_energy_ev": 0.01, "field_t": 1.0, "status": "tracked"},
            {"feature_id": "dip-1", "feature_kind": "dip", "branch": "B decreasing", "energy_ev": 1.4, "delta_energy_ev": -0.01, "field_t": -1.0, "status": "tracked"},
        ],
        "feature_points": [
            {"feature_id": "peak-1", "feature_kind": "peak", "branch": "B increasing", "energy_ev": 1.5, "delta_energy_ev": 0.01, "field_t": 1.0, "status": "tracked"},
            {"feature_id": "dip-1", "feature_kind": "dip", "branch": "B decreasing", "energy_ev": 1.4, "delta_energy_ev": -0.01, "field_t": -1.0, "status": "tracked"},
        ],
        "windows": [
            {"window_id": "w1", "center_ev": 1.5, "width_mev": 20.0, "metric": "mean", "source": "Combo"},
            {"window_id": "w2", "center_ev": 1.0, "width_mev": 20.0, "metric": "mean", "source": "Combo"},
        ],
        "slopes": [
            {"branch": "B increasing", "region": "low", "slope": 0.10, "slope_se": 0.01, "difference": 0.03, "difference_se": 0.004, "n": 5},
            {"branch": "B decreasing", "region": "low", "slope": -0.08, "slope_se": 0.02, "difference": 0.03, "difference_se": 0.004, "n": 5},
        ],
        "splitting": [{"branch": "B increasing", "points": [{"field_t": 1.0, "splitting_ev": 0.02, "status": "tracked"}]}],
        "links": [{"link_id": "link-1", "source": "peak-1", "target": "dip-1", "score": 0.8, "state": "ambiguous", "reasons": ["support overlap"]}],
        "settings": {"feature_method": "Raw spectrum", "mcd_center_ev": 1.5, "mcd_width_mev": 20.0, "fit_requested": False},
        "provenance": {"source_repeat": "repeat02", "processing_revision": 7},
        "plot_state": {"show_raw": True, "show_corrected": True, "visible_features": ["peak-1", "dip-1"]},
    }


class UnifiedMcdExportTests(unittest.TestCase):
    def test_fingerprint_hashes_array_buffers_without_python_list_expansion(self):
        from core.mcd_unified_export import _mcd_export_fingerprint
        class BufferOnly(np.ndarray):
            def tolist(self):
                raise AssertionError('large arrays must not expand into Python lists')
        values = np.arange(2000., dtype=float).reshape(100, 20).view(BufferOnly)
        snapshot = {'mcd': {'values': values}}
        first = _mcd_export_fingerprint(snapshot, [])
        self.assertEqual(first, _mcd_export_fingerprint({'mcd': {'values': np.asfortranarray(values)}}, []))
        values[0, 0] = 42.
        self.assertNotEqual(first, _mcd_export_fingerprint(snapshot, []))

    def test_fingerprint_supports_dataclass_results(self):
        from dataclasses import dataclass
        from core.mcd_unified_export import _mcd_export_fingerprint
        @dataclass
        class Result:
            values: np.ndarray
        result = Result(np.array([1., 2.]))
        first = _mcd_export_fingerprint({'result': result}, [])
        self.assertEqual(first, _mcd_export_fingerprint({'result': Result(result.values.copy())}, []))
        result.values[1] = 3.
        self.assertNotEqual(first, _mcd_export_fingerprint({'result': result}, []))

    def test_same_center_overwrites_and_different_centers_remain(self):
        snapshot = _snapshot()
        snapshot['windows'] = [{'window_id': 'current', 'center_ev': 1.64158,
                                'width_mev': 5.0, 'metric': 'mean',
                                'trace_values': np.zeros(5)}]
        with tempfile.TemporaryDirectory() as raw:
            first = export_mcd_analysis(snapshot, raw)
            original = Path(first['metadata_json']).read_bytes()
            snapshot['slopes'][0]['slope'] += .001
            second = export_mcd_analysis(snapshot, raw)
            self.assertEqual(first['export_id'], 'E1.64158eV_W5meV')
            self.assertEqual(second['export_id'], first['export_id'])
            self.assertEqual(first['revision_dir'], first['package_dir'])
            self.assertEqual(Path(first['xlsx']).name, 'MCD_unified_E1.64158eV_W5meV.xlsx')
            for path in Path(first['revision_dir']).iterdir():
                self.assertIn('E1.64158eV_W5meV', path.name)
            self.assertNotEqual(Path(first['metadata_json']).read_bytes(), original)
            snapshot['windows'][0]['center_ev'] = 1.65
            third = export_mcd_analysis(snapshot, raw)
            self.assertEqual(third['export_id'], 'E1.65000eV_W5meV')
            snapshot['windows'][0]['width_mev'] = 10
            fourth = export_mcd_analysis(snapshot, raw)
            self.assertEqual(fourth['export_id'], 'E1.65000eV_W10meV')
            self.assertFalse(Path(third['metadata_json']).exists())
            self.assertFalse(Path(third['xlsx']).exists())
            self.assertTrue(Path(first['metadata_json']).exists())
            self.assertEqual(len(list(Path(first['package_dir']).glob('*_MCD_settings*.json'))), 2)
            snapshot['windows'][0]['center_ev'] = 1.641581
            fifth = export_mcd_analysis(snapshot, raw)
            snapshot['windows'][0]['center_ev'] = 1.641582
            sixth = export_mcd_analysis(snapshot, raw)
            self.assertEqual(fifth['export_id'], 'E1.641581eV_W10meV')
            self.assertEqual(sixth['export_id'], 'E1.641582eV_W10meV')

    def test_retained_feature_only_has_no_implicit_mcd_window_products(self):
        snapshot = _snapshot()
        snapshot["windows"] = []
        snapshot["slopes"] = []
        snapshot["provenance"]["save_scope"] = "retained"
        with tempfile.TemporaryDirectory() as raw:
            result = export_mcd_analysis(snapshot, raw)
            metadata = json.loads(Path(result["metadata_json"]).read_text())
            self.assertEqual(metadata["windows"], [])
            self.assertNotIn("trace_csv", result)
            self.assertFalse(list(Path(result["revision_dir"]).glob("*MCD_vs_B*")))
            workbook = load_workbook(result["xlsx"], data_only=True)
            self.assertEqual(workbook.sheetnames, ["Energy", "Shift"])
            workbook.close()
            self.assertEqual({Path(path).stem.removesuffix('_' + result['export_id']) for path in result["plot_pngs"]},
                             {"MCD_map", "feature_energy", "feature_shift"})
            self.assertTrue(all((Path(result["revision_dir"]) / name).exists() for name in metadata["outputs"]))

    def test_legacy_empty_windows_still_exports_current_window(self):
        snapshot = _snapshot()
        snapshot["windows"] = []
        with tempfile.TemporaryDirectory() as raw:
            result = export_mcd_analysis(snapshot, raw)
            metadata = json.loads(Path(result["metadata_json"]).read_text())
            self.assertEqual(metadata["windows"][0]["window_id"], "current")
            self.assertTrue(Path(result["xlsx"]).is_file())
            self.assertNotIn("trace_csv", result)
            self.assertTrue(any(Path(path).name.startswith("MCD_vs_B_") for path in result["plot_pngs"]))

    def test_success_writes_current_plot_ready_sheets_and_metadata_json(self):
        snapshot = _snapshot()
        raw_before = snapshot["mcd"]["values"].copy()
        with tempfile.TemporaryDirectory() as raw:
            result = export_mcd_analysis(snapshot, raw)
            self.assertTrue(result["export_id"])
            self.assertEqual(result["export_id"], json.loads(Path(result["metadata_json"]).read_text())["export_id"])
            self.assertTrue(Path(result["xlsx"]).is_file())
            self.assertTrue(Path(result["mcd_map_dat"]).is_file())
            self.assertEqual({Path(path).stem.removesuffix('_' + result['export_id']) for path in result["plot_pngs"]}, {"MCD_map", "MCD_vs_B_w1", "MCD_vs_B_w2", "feature_energy", "feature_shift"})
            workbook = load_workbook(result["xlsx"], data_only=False)
            self.assertLessEqual(len(workbook.sheetnames), 5)
            self.assertEqual(workbook.sheetnames, ["MCD", "Slopes", "Energy", "Shift"])
            self.assertNotIn("Settings", workbook.sheetnames)
            self.assertNotIn("Links", workbook.sheetnames)
            mcd_headers = [cell.value for cell in workbook["MCD"][1]]
            self.assertTrue(any(str(value).startswith("B_increasing_T") for value in mcd_headers))
            self.assertTrue(any(str(value).startswith("MCD_corrected_signed_mean_increasing") for value in mcd_headers))
            self.assertTrue(any(str(value).startswith("B_decreasing_T") for value in mcd_headers))
            inc_value_column = mcd_headers.index(next(value for value in mcd_headers if str(value).startswith("MCD_corrected_signed_mean_increasing"))) + 1
            self.assertAlmostEqual(float(workbook["MCD"].cell(3, inc_value_column).value), 0.5)
            self.assertTrue(any(str(value).startswith("Energy_eV_peak") for value in [cell.value for cell in workbook["Energy"][1]]))
            metadata = json.loads(Path(result["metadata_json"]).read_text())
            self.assertEqual(metadata["links"][0]["state"], "ambiguous")
            self.assertEqual(metadata["settings"]["feature_method"], "Raw spectrum")
            self.assertEqual(metadata["source"]["repeat"], "repeat02")
            self.assertEqual(snapshot["mcd"]["values"].tolist(), raw_before.tolist())

    def test_successive_exports_replace_same_center_set(self):
        with tempfile.TemporaryDirectory() as raw:
            first = export_mcd_analysis(_snapshot(), raw)
            changed = _snapshot()
            changed['slopes'][0]['slope'] += .001
            second = export_mcd_analysis(changed, raw)
            self.assertEqual(first["export_id"], second["export_id"])
            self.assertEqual(Path(first["revision_dir"]), Path(second["revision_dir"]))
            self.assertTrue(Path(first["xlsx"]).is_file())
            self.assertTrue(Path(second["xlsx"]).is_file())
            package_dirs = list((Path(raw) / "Processed Data" / "MCD").glob("*_MCD"))
            self.assertEqual(len(package_dirs), 1)
            self.assertEqual([item for item in package_dirs[0].iterdir() if item.is_dir()], [])

    def test_identical_save_reuses_existing_files(self):
        with tempfile.TemporaryDirectory() as raw:
            first = export_mcd_analysis(_snapshot(), raw)
            before = {p.name: (p.read_bytes(), p.stat().st_mtime_ns) for p in Path(first['package_dir']).iterdir()}
            second = export_mcd_analysis(_snapshot(), raw)
            self.assertEqual(first['export_id'], second['export_id'])
            self.assertTrue(second['reused'])
            after = {p.name: (p.read_bytes(), p.stat().st_mtime_ns) for p in Path(first['package_dir']).iterdir()}
            self.assertEqual(before, after)

    def test_reloaded_source_reuses_identical_scientific_result(self):
        snapshot = _snapshot()
        snapshot['windows'] = [snapshot['windows'][0]]
        snapshot['windows'][0].update(id='window-1-1.5', window_id='window-1-1.5', source_generation=1)
        snapshot['provenance']['source_generation'] = 1
        with tempfile.TemporaryDirectory() as raw:
            first = export_mcd_analysis(snapshot, raw)
            snapshot['windows'][0].update(id='window-2-1.5', window_id='window-2-1.5', source_generation=2)
            snapshot['provenance']['source_generation'] = 2
            second = export_mcd_analysis(snapshot, raw)
            self.assertEqual(first['export_id'], second['export_id'])
            self.assertTrue(second['reused'])

    def test_flat_publication_failure_preserves_previous_export(self):
        import shutil
        snapshot = _snapshot()
        with tempfile.TemporaryDirectory() as raw:
            first = export_mcd_analysis(snapshot, raw)
            package = Path(first['package_dir'])
            before = {p.name: p.read_bytes() for p in package.iterdir()}
            snapshot['slopes'][0]['slope'] += .01
            copy = shutil.copyfileobj
            calls = []
            def fail_second(source, target, *args):
                if Path(getattr(target, 'name', '')).parent == package:
                    calls.append(target.name)
                    if len(calls) == 2:
                        raise OSError('simulated disk failure')
                copy(source, target, *args)
            with patch('core.mcd_unified_export.shutil.copyfileobj', side_effect=fail_second):
                with self.assertRaisesRegex(OSError, 'simulated disk failure'):
                    export_mcd_analysis(snapshot, raw)
            self.assertEqual({p.name: p.read_bytes() for p in package.iterdir()}, before)

    def test_same_filename_from_different_source_does_not_overwrite(self):
        snapshot = _snapshot()
        with tempfile.TemporaryDirectory() as raw:
            first = export_mcd_analysis(snapshot, raw)
            before = Path(first['xlsx']).read_bytes()
            snapshot['source_descriptor']['path'] = '/other/experiment/same.csv'
            snapshot['slopes'][0]['slope'] += .1
            second = export_mcd_analysis(snapshot, raw)
            self.assertNotEqual(first['package_dir'], second['package_dir'])
            self.assertEqual(Path(first['xlsx']).read_bytes(), before)
            self.assertTrue(export_mcd_analysis(snapshot, raw)['reused'])

    def test_save_replaces_old_same_center_revisions(self):
        snapshot = _snapshot()
        with tempfile.TemporaryDirectory() as raw:
            first = export_mcd_analysis(snapshot, raw)
            package = Path(first['package_dir'])
            metadata = json.loads(Path(first['metadata_json']).read_text())
            old_id = first['export_id']
            revision_id = old_id + '_r02'
            for name in metadata['outputs']:
                (package / name.replace(old_id, revision_id)).write_bytes((package / name).read_bytes())
            metadata.update(export_id=revision_id, revision=2,
                            outputs=[name.replace(old_id, revision_id) for name in metadata['outputs']])
            (package / Path(first['metadata_json']).name.replace(old_id, revision_id)).write_text(json.dumps(metadata))
            saved = export_mcd_analysis(snapshot, raw)
            self.assertEqual(saved['export_id'], old_id)
            self.assertEqual(len(list(package.glob('result_MCD_settings_*.json'))), 1)
            self.assertFalse(any('_r02' in path.name for path in package.iterdir()))

    def test_failed_rollback_keeps_recovery_files(self):
        import shutil
        snapshot = _snapshot()
        with tempfile.TemporaryDirectory() as raw:
            first = export_mcd_analysis(snapshot, raw)
            old_workbook = Path(first['xlsx']).read_bytes()
            snapshot['slopes'][0]['slope'] += .1
            replace = Path.replace
            copy = shutil.copyfileobj
            def fail_restore(path, target):
                if path.parent.name == 'previous':
                    raise OSError('restore blocked')
                return replace(path, target)
            def fail_publish(source, target, *args):
                if Path(getattr(target, 'name', '')).parent == Path(first['package_dir']):
                    raise OSError('write blocked')
                return copy(source, target, *args)
            with patch('core.mcd_unified_export.shutil.copyfileobj', side_effect=fail_publish), patch.object(Path, 'replace', fail_restore):
                with self.assertRaisesRegex(RuntimeError, 'Preserved recovery files'):
                    export_mcd_analysis(snapshot, raw)
            backups = list(Path(first['package_dir']).glob('.staging-*/previous/*.xlsx'))
            self.assertEqual(len(backups), 1)
            self.assertEqual(backups[0].read_bytes(), old_workbook)

    def test_failed_export_leaves_no_apparently_completed_revision(self):
        bad = _snapshot()
        bad["mcd"] = dict(bad["mcd"], values=np.ones((2, 2)))
        with tempfile.TemporaryDirectory() as raw:
            with self.assertRaises(ValueError):
                export_mcd_analysis(bad, raw)
            package_parent = Path(raw) / "Processed Data" / "MCD"
            revisions = [item for package in package_parent.glob("*_MCD") for item in package.iterdir() if item.is_dir()]
            self.assertEqual(revisions, [])

    def test_snapshot_builder_copies_arrays_and_marks_them_read_only(self):
        class Result:
            source_file = "sample.csv"
            energy_ev = np.asarray([1.0, 2.0])
            pair_b = np.asarray([-1.0, 1.0])
            pair_labels = np.asarray(["B decreasing", "B increasing"])
            pair_mcd_corrected = np.asarray([[1.0, 2.0], [3.0, 4.0]])

        result = Result()
        snapshot = build_mcd_export_snapshot(result, settings={"fit_requested": False})
        result.energy_ev[0] = 99.0
        self.assertEqual(float(snapshot["result"].energy_ev[0]), 1.0)
        with self.assertRaises(ValueError):
            snapshot["result"].energy_ev[0] = 0.0
        self.assertFalse(snapshot["settings"] is None)

    def test_metadata_keeps_existing_mcd_history_discovery_compatible(self):
        with tempfile.TemporaryDirectory() as raw:
            snapshot = _snapshot()
            source_name = snapshot["source_descriptor"]["filename"]
            snapshot["source_descriptor"] = dict(snapshot["source_descriptor"], path=str(Path(raw) / source_name), relative_path=source_name)
            exported = export_mcd_analysis(snapshot, raw)
            status = discover_mcd_processing_status(raw, [source_name])
            self.assertIn(source_name, status)
            self.assertTrue(status[source_name])
            self.assertIn("_MCD_settings_", Path(exported["metadata_json"]).name)

    def test_result_snapshot_orders_wavelength_columns_with_energy_axis(self):
        class Result:
            source_file = "asymmetric.csv"
            wavelength_nm = np.asarray([600.0, 800.0])
            energy_ev = np.asarray([1239.841984 / 600.0, 1239.841984 / 800.0])
            pair_b = np.asarray([1.0])
            pair_labels = np.asarray(["B increasing"])
            pair_mcd_corrected = np.asarray([[10.0, 20.0]])

        snapshot = build_mcd_export_snapshot(Result())
        self.assertTrue(np.allclose(snapshot["mcd"]["energy_ev"], [1239.841984 / 800.0, 1239.841984 / 600.0]))
        self.assertTrue(np.allclose(snapshot["mcd"]["values"], [[20.0, 10.0]]))

    def test_dat_keeps_duplicate_fields_and_branch_row_identity(self):
        snapshot = _snapshot()
        snapshot["mcd"] = dict(snapshot["mcd"], fields_t=np.asarray([-1.0, 1.0, 1.0, -1.0]), branches=np.asarray(["B increasing", "B increasing", "B decreasing", "B decreasing"]), values=np.arange(12.0).reshape(4, 3))
        with tempfile.TemporaryDirectory() as raw:
            output = export_mcd_analysis(snapshot, raw)
            lines = Path(output["mcd_map_dat"]).read_text().splitlines()
            self.assertEqual(lines[2].split("\t")[1:], ["-1", "1", "1", "-1"])
            columns = json.loads(lines[1].split("=", 1)[1])
            self.assertEqual([item["branch"] for item in columns], ["B increasing", "B increasing", "B decreasing", "B decreasing"])
            self.assertEqual(len(lines[3].split("\t")), 5)
            loaded = load_dat(output["mcd_map_dat"])
            self.assertEqual(loaded.Z.shape, (4, 3))
            self.assertTrue(np.allclose(loaded.Z[:, 0], [0.0, 3.0, 6.0, 9.0]))

    def test_feature_and_splitting_tables_keep_channel_method_and_pair_identity(self):
        snapshot = _snapshot()
        snapshot["plot_state"] = dict(snapshot["plot_state"], visible_features=None)
        snapshot["feature_points"] = [
            {"feature_id": "peak-1", "feature_kind": "peak", "channel": "K", "analysis_method": "raw", "branch": "B increasing", "field_t": 1.0, "energy_ev": 1.6, "delta_energy_ev": .01},
            {"feature_id": "peak-1", "feature_kind": "peak", "channel": "Kp", "analysis_method": "second derivative", "branch": "B increasing", "field_t": 1.0, "energy_ev": 1.7, "delta_energy_ev": .02},
        ]
        snapshot["splitting"] = [
            {"pair_id": "pair-a", "method": "raw", "selected_channel": "K", "counterpart_channel": "Kp", "branch": "B increasing", "points": [{"field_t": 1.0, "splitting_ev": .01}]},
            {"pair_id": "pair-b", "method": "raw", "selected_channel": "K", "counterpart_channel": "Kp", "branch": "B increasing", "points": [{"field_t": 1.0, "splitting_ev": .02}]},
        ]
        with tempfile.TemporaryDirectory() as raw:
            output = export_mcd_analysis(snapshot, raw)
            workbook = load_workbook(output["xlsx"], data_only=True)
            energy_headers = [cell.value for cell in workbook["Energy"][1]]
            splitting_headers = [cell.value for cell in workbook["Splitting"][1]]
            self.assertEqual(len([header for header in energy_headers if str(header).startswith("Energy_eV")]), 2)
            self.assertEqual(len([header for header in splitting_headers if str(header).startswith("Splitting_meV")]), 2)
            metadata = json.loads(Path(output["metadata_json"]).read_text())
            self.assertIn("slopes", metadata)
            self.assertEqual(len(metadata["windows"]), 2)

    def test_late_render_failure_removes_staging_revision(self):
        with tempfile.TemporaryDirectory() as raw:
            with patch("core.mcd_unified_export._plot_pngs", side_effect=RuntimeError("render failed")):
                with self.assertRaises(RuntimeError):
                    export_mcd_analysis(_snapshot(), raw)
            package_parent = Path(raw) / "Processed Data" / "MCD"
            self.assertEqual([item for package in package_parent.glob("*_MCD") for item in package.iterdir() if item.is_dir()], [])

    def test_real_slope_analysis_dict_and_peak_results_export_as_distinct_products(self):
        fields = np.asarray([-0.2, -0.1, 0.0, 0.1, 0.2])
        branches = np.asarray(["B increasing"] * 5)
        slopes = fit_mcd_slopes(fields, 2.0 * fields, branches, ranges={"low": (-0.2, 0.2)}).to_dict()
        peak_results = {
            "Raw spectrum": {
                "raw pos": PeakShiftResult(fields, branches, tuple(() for _ in fields), (PeakTrack(1, "B increasing", tuple(PeakPoint(float(b), "B increasing", 1.6, float(b), "tracked") for b in fields), 1.6, 0.0, "exact 0 T"),), "raw pos"),
                "raw neg": PeakShiftResult(fields, branches, tuple(() for _ in fields), (PeakTrack(1, "B increasing", tuple(PeakPoint(float(b), "B increasing", 1.7, float(b), "tracked") for b in fields), 1.7, 0.0, "exact 0 T"),), "raw neg"),
            }
        }
        with tempfile.TemporaryDirectory() as raw:
            snapshot = _snapshot()
            snapshot["slopes"] = slopes
            snapshot["analysis_results"] = peak_results
            snapshot["feature_points"] = ()
            output = export_mcd_analysis(snapshot, raw)
            workbook = load_workbook(output["xlsx"], data_only=True)
            slope_headers = [cell.value for cell in workbook["Slopes"][7]]
            self.assertEqual(slope_headers, ["Region", "B range (T)", "Inc slope", "Inc SE", "Dec slope", "Dec SE"])
            energy_headers = [cell.value for cell in workbook["Energy"][1]]
            self.assertEqual(len([header for header in energy_headers if str(header).startswith("Energy_eV")]), 2)
            metadata = json.loads(Path(output["metadata_json"]).read_text())
            self.assertEqual(len(metadata["slopes"]["fits"]), len(slopes["fits"]))
            self.assertIn("field_min_t", metadata["slopes"]["fits"][0])
            self.assertIn("field_max_t", metadata["slopes"]["fits"][0])

    def test_real_process_mcd_result_can_be_snapshotted_and_exported(self):
        with tempfile.TemporaryDirectory() as raw:
            source = Path(raw) / "sample_D-0.2V_T1.67K_repeat01.csv"
            frame = pd.DataFrame({
                "Bmid_T": [-1.0, -1.0, 0.0, 0.0, 1.0, 1.0],
                "rotation_angle_deg": [0.0, 90.0, 0.0, 90.0, 0.0, 90.0],
                "500.0": [10, 9, 10, 10, 11, 10],
                "600.0": [11, 10, 10, 10, 12, 10],
                "700.0": [12, 11, 10, 10, 13, 10],
            })
            frame.to_csv(source, index=False)
            processed = process_mcd(str(source), McdSettings(max_sequence_gap=1, max_delta_b=0.01))
            snapshot = build_mcd_export_snapshot(processed, settings={"mcd_center_ev": 2.0, "mcd_width_mev": 20.0})
            self.assertTrue(np.allclose(snapshot["mcd"]["energy_ev"], np.sort(processed.energy_ev)))
            acquisition_order = np.argsort(1239.841984 / np.asarray(processed.wavelength_nm), kind="stable")
            self.assertTrue(np.allclose(snapshot["mcd"]["values"], processed.pair_mcd_corrected[:, acquisition_order], equal_nan=True))
            from matplotlib.figure import Figure
            captured_figures = []
            with patch.object(Figure, "savefig", lambda figure, *_args, **_kwargs: captured_figures.append(figure)):
                _plot_pngs(Path(raw), _mcd_data(snapshot), snapshot, "r01", map_data=_selected_map_data(snapshot, _mcd_data(snapshot)))
            self.assertTrue(captured_figures[0].axes[0].collections)
            output = export_mcd_analysis(snapshot, raw)
            self.assertTrue(Path(output["xlsx"]).is_file())
            self.assertTrue(Path(output["mcd_map_dat"]).is_file())
            status = discover_mcd_processing_status(raw, [source.name])
            self.assertTrue(status.get(source.name))

    def test_unavailable_spectrum_product_is_omitted(self):
        snapshot = _snapshot()
        snapshot["spectra"] = {}
        with tempfile.TemporaryDirectory() as raw:
            output = export_mcd_analysis(snapshot, raw)
            self.assertFalse(any("spectra_" in path for path in output["plot_pngs"]))

    def test_selected_datacube_drives_map_dat_and_metadata(self):
        snapshot = _snapshot()
        snapshot["maps"] = {
            "B increasing": {
                "energy": np.asarray([1.1, 1.9]),
                "gate": np.asarray([0.25, 0.75]),
                "Z": np.asarray([[4.0, 5.0], [6.0, 7.0]]),
                "title": "B increasing",
            }
        }
        snapshot["selected_map"] = "B increasing"
        from matplotlib.figure import Figure
        captured = []
        def save_figure(figure, *_args, **_kwargs): captured.append(figure)
        with tempfile.TemporaryDirectory() as raw, patch.object(Figure, "savefig", save_figure):
            _plot_pngs(Path(raw), _mcd_data(snapshot), snapshot, "r01", map_data=_selected_map_data(snapshot, _mcd_data(snapshot)))
        self.assertTrue(captured[0].axes[0].collections)
        with tempfile.TemporaryDirectory() as raw:
            output = export_mcd_analysis(snapshot, raw)
            loaded = load_dat(output["mcd_map_dat"])
            self.assertTrue(np.allclose(loaded.energy, [1.1, 1.9]))
            self.assertTrue(np.allclose(loaded.gate, [0.25, 0.75]))
            self.assertTrue(np.allclose(loaded.Z, [[4.0, 5.0], [6.0, 7.0]]))
            metadata = json.loads(Path(output["metadata_json"]).read_text())
            self.assertEqual(metadata["map"]["name"], "B increasing")

    def test_visible_branch_keeps_direct_retained_trace_values_aligned(self):
        snapshot = _snapshot()
        snapshot["plot_state"] = {"visible_branches": ["B increasing"]}
        snapshot["windows"] = [{"window_id": "current", "values": [10.0, 20.0, 30.0, 40.0, 50.0], "metric": "mean"}]
        with tempfile.TemporaryDirectory() as raw:
            output = export_mcd_analysis(snapshot, raw)
            workbook = load_workbook(output["xlsx"], data_only=True)
            headers = [cell.value for cell in workbook["MCD"][1]]
            value_column = next(index for index, header in enumerate(headers, start=1) if str(header).startswith("MCD_corrected_signed_mean_increasing"))
            self.assertEqual([workbook["MCD"].cell(row, value_column).value for row in (2, 3, 4)], [30.0, 40.0, 50.0])

    def test_invalid_retained_window_is_rejected(self):
        snapshot = _snapshot()
        snapshot["windows"] = [{"window_id": "stale", "metric": "integral"}]
        with tempfile.TemporaryDirectory() as raw:
            with self.assertRaises(ValueError):
                export_mcd_analysis(snapshot, raw)
        snapshot = _snapshot()
        snapshot["windows"] = [{"window_id": "wrong-grid", "center_ev": 1.5, "width_mev": 20.0, "values": [1.0, 2.0], "metric": "mean"}]
        with tempfile.TemporaryDirectory() as raw:
            with self.assertRaises(ValueError):
                export_mcd_analysis(snapshot, raw)
        snapshot = _snapshot()
        snapshot["windows"] = [{"window_id": "out", "center_ev": 9.0, "width_mev": 20.0, "metric": "mean"}]
        with tempfile.TemporaryDirectory() as raw:
            with self.assertRaises(ValueError):
                export_mcd_analysis(snapshot, raw)

    def test_plot_state_selects_raw_spectrum_b_and_feature_visibility(self):
        from matplotlib.figure import Figure
        captured = []
        def save_figure(figure, *_args, **_kwargs):
            captured.append(figure)
        snapshot = _snapshot()
        snapshot["plot_state"] = {"selected_field_t": 2.0, "spectrum_channel": "K_raw", "spectrum_mode": "raw", "export_spectra": True, "visible_features": []}
        snapshot["spectra"]["K_raw"]["raw_values"] = np.asarray([[10.0, 11.0, 12.0], [20.0, 21.0, 22.0], [30.0, 31.0, 32.0]])
        with tempfile.TemporaryDirectory() as raw, patch.object(Figure, "savefig", save_figure):
            paths = _plot_pngs(Path(raw), _mcd_data(snapshot), snapshot, "r01")
        self.assertFalse(any("features_" in str(path) for path in paths))
        spectrum_axis = next(figure.axes[0] for figure in captured if figure.axes[0].get_ylabel() == "Spectrum")
        self.assertEqual(len(spectrum_axis.lines), 1)
        self.assertTrue(np.allclose(spectrum_axis.lines[0].get_ydata(), [30.0, 31.0, 32.0]))
        mcd_axes = [figure.axes[0] for figure in captured if figure.axes[0].get_ylabel() == "MCD"]
        self.assertEqual(len(mcd_axes), 2)
        self.assertEqual([len(axis.lines) for axis in mcd_axes], [2, 2])

    def test_visible_branch_filters_mcd_slopes_features_and_splitting(self):
        from matplotlib.figure import Figure
        captured = []
        def save_figure(figure, *_args, **_kwargs): captured.append(figure)
        snapshot = _snapshot()
        snapshot["plot_state"] = {"visible_branches": ["B increasing"], "visible_features": None}
        snapshot["feature_points"] = [
            {"feature_id": "inc", "feature_kind": "peak", "branch": "B increasing", "field_t": 1.0, "energy_ev": 1.5},
            {"feature_id": "dec", "feature_kind": "peak", "branch": "B decreasing", "field_t": -1.0, "energy_ev": 1.4},
        ]
        snapshot["splitting"] = [
            {"pair_id": "inc", "branch": "B increasing", "points": [{"field_t": 1.0, "splitting_ev": .01}]},
            {"pair_id": "dec", "branch": "B decreasing", "points": [{"field_t": -1.0, "splitting_ev": .02}]},
        ]
        snapshot["slopes"] = [{"branch": "B increasing", "slope": 1.0, "intercept": 0.0, "field_min_t": 0.0, "field_max_t": 1.0}, {"branch": "B decreasing", "slope": 2.0, "intercept": 0.0, "field_min_t": -1.0, "field_max_t": 0.0}]
        with tempfile.TemporaryDirectory() as raw, patch.object(Figure, "savefig", save_figure):
            _plot_pngs(Path(raw), _mcd_data(snapshot), snapshot, "r01")
        self.assertTrue(all("decreasing" not in label.casefold() for label in captured[1].axes[0].get_legend_handles_labels()[1]))
        energy_axis = next(figure.axes[0] for figure in captured if figure.axes[0].get_ylabel() == "Energy (eV)")
        self.assertEqual(len(energy_axis.lines), 1)
        with tempfile.TemporaryDirectory() as raw, patch.object(Figure, "savefig", save_figure):
            output = export_mcd_analysis(snapshot, raw)
            workbook = load_workbook(output["xlsx"], data_only=True)
            self.assertFalse(any("decreasing" in str(cell.value).casefold() for cell in workbook["MCD"][1]))
            self.assertFalse(any("decreasing" in str(cell.value).casefold() for cell in workbook["Slopes"][1]))
            dat_header = Path(output["mcd_map_dat"]).read_text().splitlines()[2]
            self.assertEqual(dat_header.split("\t")[1:], ["0", "1", "2"])


if __name__ == "__main__":
    unittest.main()
