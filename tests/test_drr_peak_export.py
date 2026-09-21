import copy
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import numpy as np
from openpyxl import load_workbook
from PIL import Image

from core.loader import DataCube
from core.drr_peak_export import export_drr_peak_analysis


class PeakExportTests(unittest.TestCase):
    def test_appending_peak_rows_does_not_scan_existing_worksheet(self):
        from openpyxl import Workbook
        from openpyxl.worksheet.worksheet import Worksheet
        from core.drr_peak_export import _append
        book = Workbook()
        self.addCleanup(book.close)
        sheet = book.active
        original = Worksheet.max_row.fget
        scans = []
        def max_row(worksheet):
            scans.append(1)
            return original(worksheet)
        with patch.object(Worksheet, 'max_row', property(max_row)):
            for index in range(100):
                _append(sheet, [index, '=literal label', 1.5, None])
        self.assertLessEqual(len(scans), 1, 'Appending rows must not repeatedly scan all previous cells')
        self.assertEqual(sheet.max_row, 100)
        self.assertEqual(sheet.cell(100, 1).value, 99)
        self.assertEqual(sheet.cell(100, 2).data_type, 's')
        self.assertEqual(sheet.cell(100, 3).value, 1.5)

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        energy = np.linspace(1.5, 1.7, 41)
        self.cube = DataCube(energy, np.array([-1., 0., 1.]),
                             np.array([(energy - 1.6) ** 2] * 3),
                             "Doping (V)", "Synthetic source", "DR/R")
        def point(row, energy, track, status="accepted"):
            return dict(row_index=row, y=float(row-1), energy=energy,
                        amplitude=0.1, prominence=0.03, polarity="peak",
                        track_id=track, status=status)
        self.result = dict(schema_version=1, y_label="Doping (V)",
                           y_values=[-1., 0., 1.],
                           settings=dict(x_min=1.55, x_max=1.65, y_min=-1., y_max=1.,
                                         source="both", polarity="both", prominence=0.05,
                                         min_distance_mev=1., max_peaks=6, max_shift_mev=3.,
                                         sg_window=9, sg_polyorder=2),
                           products={"raw": {"points": [point(0, 1.59, 2), point(2, 1.61, 2),
                                                          point(1, 1.63, 7, "uncertain")]},
                                     "second": {"points": [point(0, 1.60, 1)]}})

    def test_wide_numeric_columns_keep_track_identity_and_missing_rows(self):
        paths = export_drr_peak_analysis(self.folder, "sample", self.result, self.cube)
        book = load_workbook(paths["xlsx"], data_only=True)
        self.addCleanup(book.close)
        self.assertEqual(book.sheetnames, ["DRR_Peaks", "D2E_Peaks", "Peak_Info", "Parameters"])
        sheet = book["DRR_Peaks"]
        self.assertEqual(sheet.cell(1, 1).value, "Doping (V)")
        self.assertIn("2", sheet.cell(1, 2).value)
        self.assertEqual(sheet.max_column, 2)
        self.assertEqual(list(sheet.values)[1:], [(-1., 1.59), (0., None), (1., 1.61)])
        self.assertEqual(sheet.cell(2, 2).data_type, "n")
        self.assertIn("uncertain", str(list(book["Peak_Info"].values)))
        payload = json.loads(paths["json"].read_text(encoding="utf-8"))
        self.assertEqual(payload["products"], self.result["products"])
        self.assertEqual(payload["metadata"]["source_title"], "Synthetic source")
        with Image.open(paths["png"]) as image:
            self.assertGreater(image.width, 500)
            image.verify()

    def test_unique_safe_outputs_preserve_existing_originals_and_inputs(self):
        original = self.folder / "sample.png"
        original.write_bytes(b"original")
        before = copy.deepcopy(self.result)
        z_before = self.cube.Z.copy()
        first = export_drr_peak_analysis(self.folder, "../sample", self.result, self.cube)
        second = export_drr_peak_analysis(self.folder, "../sample", self.result, self.cube)
        self.assertTrue(set(first.values()).isdisjoint(second.values()))
        self.assertTrue(all(p.parent == self.folder for p in first.values()))
        self.assertTrue(first["xlsx"].name.endswith("_peak_analysis.xlsx"))
        self.assertTrue(first["png"].name.endswith("_peaks_comparison.png"))
        self.assertEqual(original.read_bytes(), b"original")
        self.assertEqual(self.result, before)
        np.testing.assert_array_equal(self.cube.Z, z_before)

    def test_missing_product_exports_empty_sheet_and_comparison(self):
        del self.result["products"]["second"]
        paths = export_drr_peak_analysis(self.folder, "one", self.result, self.cube)
        book = load_workbook(paths["xlsx"])
        self.addCleanup(book.close)
        self.assertEqual(list(book["D2E_Peaks"].values), [("Doping (V)",), (-1.,), (0.,), (1.,)])

    def test_failure_cleans_new_files(self):
        with patch("matplotlib.figure.Figure.savefig", side_effect=OSError("disk full")):
            with self.assertRaises(OSError):
                export_drr_peak_analysis(self.folder, "bad", self.result, self.cube)
        self.assertEqual(list(self.folder.iterdir()), [])

    def test_invalid_point_rejected_before_writing(self):
        self.result["products"]["raw"]["points"][0]["energy"] = float("nan")
        with self.assertRaises(ValueError):
            export_drr_peak_analysis(self.folder, "bad", self.result, self.cube)
        self.assertEqual(list(self.folder.iterdir()), [])

    def test_derivative_plot_preserves_missing_measurement_gaps(self):
        from matplotlib.figure import Figure
        self.cube.Z[1, 20] = np.nan
        original_save = Figure.savefig
        inspected = []
        def inspect_and_save(fig, *args, **kwargs):
            derivative_ax = next(ax for ax in fig.axes if ax.get_title() == "Second energy derivative")
            values = derivative_ax.collections[0].get_array()
            inspected.append(np.ma.getmaskarray(values)[1, 10])
            return original_save(fig, *args, **kwargs)
        with patch.object(Figure, "savefig", inspect_and_save):
            export_drr_peak_analysis(self.folder, "gaps", self.result, self.cube)
        self.assertEqual(inspected, [True])

    def test_source_provenance_in_parameters_and_json(self):
        self.result["provenance"] = {"source_files": ["measurement.csv"], "reference": "zero.csv"}
        paths = export_drr_peak_analysis(self.folder, "source", self.result, self.cube)
        book = load_workbook(paths["xlsx"])
        self.addCleanup(book.close)
        params = dict(book["Parameters"].values)
        self.assertEqual(json.loads(params["provenance"]), self.result["provenance"])
        payload = json.loads(paths["json"].read_text(encoding="utf-8"))
        self.assertEqual(payload["provenance"], self.result["provenance"])

    def test_descending_nonuniform_derivative_matches_increasing_plot(self):
        from matplotlib.figure import Figure
        energy = 1.5 + 0.2 * np.linspace(0, 1, 41) ** 1.2
        self.cube.energy = energy
        self.cube.Z = np.array([np.exp(-((energy - 1.6) / .02) ** 2)] * 3)
        original_save = Figure.savefig
        inspected = []
        def inspect_and_save(fig, *args, **kwargs):
            ax = next(ax for ax in fig.axes if ax.get_title() == "Second energy derivative")
            inspected.append(ax.collections[0].get_array().copy())
            return original_save(fig, *args, **kwargs)
        with patch.object(Figure, "savefig", inspect_and_save):
            export_drr_peak_analysis(self.folder, "ascending", self.result, self.cube)
            self.cube.energy = energy[::-1]
            self.cube.Z = self.cube.Z[:, ::-1]
            export_drr_peak_analysis(self.folder, "descending", self.result, self.cube)
        np.testing.assert_allclose(inspected[0], inspected[1])

    def test_many_tracks_use_compact_legend_and_omit_uncertain_only_columns(self):
        from matplotlib.figure import Figure
        prototype = self.result["products"]["raw"]["points"][0]
        points = [dict(prototype, track_id=index, energy=1.56 + index * .0005,
                       status="accepted" if index < 30 else "uncertain",
                       polarity="peak" if index % 2 else "dip") for index in range(100)]
        self.result["products"]["raw"]["points"] = points
        original_save = Figure.savefig
        legends = []
        line_counts = []
        def inspect_and_save(fig, *args, **kwargs):
            ax = next(ax for ax in fig.axes if ax.get_title() == "DR/R")
            legends.append([text.get_text() for text in ax.get_legend().get_texts()])
            line_counts.append(len(ax.lines))
            return original_save(fig, *args, **kwargs)
        with patch.object(Figure, "savefig", inspect_and_save):
            paths = export_drr_peak_analysis(self.folder, "many", self.result, self.cube)
        self.assertEqual(set(legends[0]), {"Accepted peaks", "Accepted dips", "Uncertain"})
        self.assertEqual(line_counts, [0])
        book = load_workbook(paths["xlsx"])
        self.addCleanup(book.close)
        self.assertEqual(book["DRR_Peaks"].max_column, 31)
        self.assertEqual(book["Peak_Info"].max_row, 102)


if __name__ == "__main__":
    unittest.main()
