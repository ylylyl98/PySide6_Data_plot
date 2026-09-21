import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
import unittest
import numpy as np
from core.drr_p2p_metrics import metric_values

class MetricTests(unittest.TestCase):
    def test_position_pick_preserves_duplicate_y_row(self):
        from PySide6.QtWidgets import QApplication,QWidget
        from core.loader import DataCube
        from core.drr_analysis_workspace import create_dataset
        from ui_qt.drr_p2p_batch import BatchRangeAmplitudePage
        from types import SimpleNamespace
        app=QApplication.instance() or QApplication([]);owner=QWidget()
        d=create_dataset(DataCube(np.array([1.83,1.84,1.85,1.86]),np.array([1.,1.]),np.array([[0.,-1.,2.,0.],[0.,2.,-1.,0.]]),'Y','A','DRR'),'A',{})
        p=BatchRangeAmplitudePage(owner,d,(1.83,1.86,1.,1.))
        try:
            p.view_mode.setCurrentIndex(1);p.comparison_metric.setCurrentIndex(1);p.canvas.draw()
            line=p.metric_axes[0].lines[0];xy=line.axes.transData.transform([1.,1.84])
            p.pick_curve(SimpleNamespace(artist=line,ind=[0,1],mouseevent=SimpleNamespace(x=xy[0],y=xy[1])))
            self.assertEqual(p.row_spin.value(),2)
        finally:p.close();owner.close()

    def test_columns_and_mev(self):
        data=np.array([[2.,.4,-.1,.3,1.84,1.85,20]])
        np.testing.assert_allclose(metric_values(data,'maximum'),[1.85])
        np.testing.assert_allclose(metric_values(data,'minimum'),[1.84])
        np.testing.assert_allclose(metric_values(data,'separation'),[10.])

    def test_views_and_origin_workbook(self):
        from PySide6.QtWidgets import QApplication,QWidget
        from core.loader import DataCube
        from core.drr_analysis_workspace import create_dataset
        from ui_qt.drr_p2p_batch import BatchRangeAmplitudePage
        from tempfile import TemporaryDirectory
        from pathlib import Path
        from openpyxl import load_workbook
        from unittest.mock import patch
        app=QApplication.instance() or QApplication([]);owner=QWidget()
        x=np.array([1.83,1.84,1.85,1.86]);y=np.array([0.,1.])
        a=create_dataset(DataCube(x,y,np.array([[0.,-1.,2.,0.]]*2),'Y','A','DRR'),'A',{})
        b=create_dataset(DataCube(x,np.array([5.]),np.array([[0.,-2.,1.,0.]]),'Y','B','DRR'),'B',{'id':'b'})
        owner.datasets={d.key:d for d in (a,b)}
        page=BatchRangeAmplitudePage(owner,a,(1.83,1.86,0.,1.))
        try:
            page.receive_records(page.compute_records([a,b],(1.83,1.86,0.,1.),False));page.view_mode.setCurrentIndex(1)
            with patch.object(page,'compute_records',side_effect=AssertionError('recalculated')):
                page.comparison_metric.setCurrentIndex(1)
                fig=page.comparison_export_figure();self.assertEqual(len(fig.axes),2)
                np.testing.assert_allclose(fig.axes[0].lines[0].get_ydata(),1.85)
                np.testing.assert_allclose(fig.axes[1].lines[0].get_ydata(),1.84)
                from types import SimpleNamespace
                page.canvas.draw()
                line=next(l for l in page.metric_axes[1].lines if getattr(l,'_p2p_record',None)==1)
                coords=line.axes.transData.transform([5.,1.84])
                page.pick_curve(SimpleNamespace(artist=line,ind=[0],mouseevent=SimpleNamespace(x=coords[0],y=coords[1])))
                self.assertEqual(page.name,'B')
                page.comparison_metric.setCurrentIndex(2)
                fig=page.comparison_export_figure();np.testing.assert_allclose(fig.axes[0].lines[0].get_ydata(),10.)
                state=page.snapshot_state();page.comparison_metric.setCurrentIndex(0);page.restore_state(state)
                self.assertEqual(page.comparison_quantity(),'separation')
            with TemporaryDirectory() as folder:
                import csv
                csv_path=Path(folder)/'metrics.csv';page.write_csv(csv_path)
                with csv_path.open(encoding='utf-8-sig') as f:rows=list(csv.DictReader(f))
                self.assertEqual(len(rows),3)
                self.assertAlmostEqual(float(rows[0]['Separation_meV']),10.)
                self.assertIn('B_T',rows[0]);self.assertIn('Emax_edge',rows[0])
                path=Path(folder)/'metrics.xlsx';page.write_xlsx(path)
                wb=load_workbook(path)
                self.assertTrue({'Amplitude','Emax','Emin','Separation','Metadata'}<=set(wb.sheetnames))
                self.assertEqual(wb['Emax']['B2'].value,1.85)
                self.assertEqual(wb['Emax']['C2'].value,5.)
                self.assertIsNone(wb['Emax']['C3'].value)
                wb.close()
        finally:page.close();owner.close()
