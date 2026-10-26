"""Isolated native DRR Analysis regression and QML lifecycle checks."""
from __future__ import annotations

import os
from pathlib import Path
import subprocess
import sys
from tempfile import TemporaryDirectory
import unittest


class QuickDrrAnalysisTests(unittest.TestCase):
    def test_native_analysis_workflow(self):
        with TemporaryDirectory(prefix="quick-drr-analysis-process-") as temporary:
            environment = dict(os.environ, QT_QPA_PLATFORM="offscreen", QT_QUICK_BACKEND="software", MPLBACKEND="agg",
                MPLCONFIGDIR=temporary, XDG_CONFIG_HOME=temporary, XDG_CACHE_HOME=temporary, PYTHONIOENCODING="utf-8")
            result = subprocess.run([sys.executable, str(Path(__file__).resolve()), "--isolated"],
                cwd=Path(__file__).resolve().parents[1], env=environment, capture_output=True, text=True, timeout=240)
        self.assertEqual(result.returncode, 0, result.stdout+result.stderr)
        self.assertIn("QUICK_DRR_ANALYSIS_TESTS=", result.stdout)
        print(result.stdout.strip())


def isolated():
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    import importlib.abc
    import json
    import threading
    import time
    import traceback
    from copy import deepcopy
    from dataclasses import replace
    from unittest.mock import patch

    class NoWidgets(importlib.abc.MetaPathFinder):
        def find_spec(self, fullname, path=None, target=None):
            if fullname in {"PySide6.QtWidgets", "PySide6.QtQuickWidgets", "ui_qt.common", "ui_qt.main_window",
                "ui_qt.drr_analysis_window", "ui_qt.drr_p2p_batch", "matplotlib.backends.backend_qt", "matplotlib.backends.backend_qtagg"}:
                raise AssertionError("Native DRR Analysis imported Widgets: "+fullname)
    sys.meta_path.insert(0, NoWidgets())
    import numpy as np
    from PySide6.QtCore import QEvent, QMetaObject, QObject, QSettings, QUrl, qInstallMessageHandler
    from PySide6.QtGui import QGuiApplication
    from PySide6.QtQml import QQmlApplicationEngine
    from PySide6.QtQuickControls2 import QQuickStyle
    from core.loader import DataCube
    from core.drr_analysis_workspace import create_dataset
    from core.drr_workspace_session import load_session, save_session
    from ui_qt.quick.drr_analysis_bridge import DrrAnalysisBridge
    from ui_qt.quick.plot_surface import register_plot_surface
    from ui_qt.quick import drr_analysis_workflow as workflow

    QQuickStyle.setStyle("Basic")
    app = QGuiApplication([])
    python_errors, qml_errors = [], []
    def exception(kind, value, trace):
        python_errors.append("".join(traceback.format_exception(kind, value, trace)))
        traceback.print_exception(kind, value, trace)
    sys.excepthook = exception
    def qt_message(kind, context, message):
        if any(marker in message for marker in ("ReferenceError", "TypeError", "Binding loop", "Cannot assign", "Unable to assign", "Error:")):
            qml_errors.append(str(message)); print(message, file=sys.stderr)
    previous_handler = qInstallMessageHandler(qt_message)

    def wait(predicate, timeout=45):
        deadline = time.monotonic()+timeout
        while time.monotonic() < deadline:
            app.processEvents()
            if predicate():
                app.processEvents(); return
            time.sleep(.005)
        raise AssertionError("Timed out waiting for native DRR Analysis state")

    class Cases(unittest.TestCase):
        def setUp(self):
            start_python, start_qml = len(python_errors), len(qml_errors)
            self.addCleanup(lambda: self.assertEqual(python_errors[start_python:], [], "Qt callback exception"))
            self.addCleanup(lambda: self.assertEqual(qml_errors[start_qml:], [], "QML runtime error"))
            self.temporary = TemporaryDirectory(prefix="quick-drr-analysis-case-")
            self.addCleanup(self.temporary.cleanup)
            self.root = Path(self.temporary.name)
            self.bridge = DrrAnalysisBridge(settings=QSettings(str(self.root/"settings.ini"), QSettings.IniFormat),
                session_path=self.root/"workspace.npz", start_folder=str(self.root))
            self.addCleanup(self.dispose)

        def dispose(self):
            self.bridge.close()
            wait(lambda: all(future.done() for future in self.bridge._futures.values()))
            app.processEvents()

        def idle(self):
            wait(lambda: not self.bridge._pending and not self.bridge._plot_timer.isActive() and not self.bridge._filter_timer.isActive() and not self.bridge._p2p_timer.isActive())
            if self.bridge._canvas:
                wait(lambda: not self.bridge._canvas.backend._timer.isActive())

        def dataset(self, name="first", rows=6):
            x = np.linspace(1.,1.1,151)
            y = np.arange(rows,dtype=float)
            z = np.array([np.exp(-((x-1.035-index*.000005)/.002)**2)-.4*np.exp(-((x-1.055)/.003)**2)+.8*np.exp(-((x-1.075)/.002)**2) for index in range(rows)])
            cube = DataCube(x,y,z,"Gate (V)",name,"DR/R")
            dataset = create_dataset(cube,name,{"measurement_files":[str(self.root/(name+".csv"))]})
            dataset.settings = replace(dataset.settings, source="raw", polarity="peaks")
            return dataset

        def add(self, *datasets):
            self.bridge.addDatasets(list(datasets)); self.bridge.setParameter("display",1); self.idle()

        def analyze(self):
            self.bridge.setParameter("automatic",False)
            self.bridge.invoke("analyze_selected"); self.idle()
            self.assertFalse(self.bridge.snapshot["error"])

        def test_candidates_pagination_filter_exclusion_and_seed_settings(self):
            dataset = self.dataset(rows=260)
            self.add(dataset); self.analyze()
            self.assertGreater(self.bridge.snapshot["retained"],500)
            self.assertEqual(self.bridge.peakResultsModel.rowCount(),500)
            self.bridge.invoke("next")
            self.assertEqual(self.bridge.snapshot["page"],2)
            self.bridge.selectPeak(510)
            revision = dataset.revision
            self.bridge.invoke("exclude"); self.idle()
            self.assertEqual(dataset.revision,revision+1)
            self.assertEqual(len(dataset.result["excluded_points"]),1)
            original = sum(len(product["points"]) for product in dataset.result["products"].values())
            self.bridge.setParameter("filter.x_min",1.065)
            self.bridge.setParameter("filter.x_max",1.09)
            self.bridge.setParameter("filter.limit_x",True)
            self.idle()
            self.assertLess(self.bridge.snapshot["retained"],original)
            self.assertEqual(sum(len(product["points"]) for product in dataset.result["products"].values()),original)
            self.bridge.invoke("reset_filter"); self.idle()
            self.assertEqual(self.bridge.snapshot["retained"],original)
            self.bridge.setParameter("peak.seed_energy",1.035)
            self.bridge.setParameter("peak.seed_y",2.)
            self.bridge.setParameter("peak.mode","seed")
            self.bridge.invoke("analyze"); self.idle()
            self.assertIsNotNone(dataset.result)
            self.assertEqual(dataset.settings.mode,"seed")
            self.assertFalse(dataset.cube.Z.flags.writeable)

        def test_revision_guard_discards_analysis_and_filter_from_old_state(self):
            dataset = self.dataset()
            self.add(dataset)
            entered, release = threading.Event(), threading.Event()
            original = workflow.analyze
            def slow(*args,**kwargs):
                entered.set(); release.wait(10); return original(*args,**kwargs)
            with patch.object(workflow,"analyze",side_effect=slow):
                self.bridge.invoke("analyze")
                wait(entered.is_set)
                self.bridge.setParameter("peak.prominence",.2)
                release.set(); self.idle()
            self.assertIsNone(dataset.result)
            self.analyze()
            from core import drr_result_filter
            original_filter = drr_result_filter.filtered_result
            entered.clear(); release.clear()
            def slow_filter(*args,**kwargs):
                entered.set(); release.wait(10); return original_filter(*args,**kwargs)
            with patch.object(drr_result_filter,"filtered_result",side_effect=slow_filter):
                self.bridge.setParameter("filter.limit_x",True)
                wait(entered.is_set)
                self.bridge.setParameter("peak.prominence",.3)
                release.set(); self.idle()
            self.assertIsNone(dataset.result)
            self.assertEqual(self.bridge.snapshot["retained"],0)

        def test_result_filter_reuses_maps_colorbars_and_navigation(self):
            self.add(self.dataset());self.analyze()
            canvas=self.bridge._canvas;axes=tuple(canvas.figure.axes);meshes=tuple(row[3] for row in self.bridge._maps)
            self.bridge.views.clear()
            signature=self.bridge._static_peak_signature()
            self.bridge._save_view()
            self.assertEqual(self.bridge._static_peak_signature(),signature)
            self.bridge._navigation.push_current();history=len(self.bridge._navigation._nav_stack)
            draws,blits=canvas.draw_count,canvas.blit_count
            self.bridge.setParameter("filter.min_prominence",.1);self.idle()
            self.assertEqual(tuple(canvas.figure.axes),axes)
            self.assertEqual(tuple(row[3] for row in self.bridge._maps),meshes)
            self.assertEqual(canvas.draw_count,draws);self.assertGreater(canvas.blit_count,blits)
            self.assertEqual(len(self.bridge._navigation._nav_stack),history)
            displayed=self.bridge._display_result(self.bridge.active_key)
            plotted=sum(len(artist._drr_peaks[2]) for artist in self.bridge._peak_overlays if hasattr(artist,"_drr_peaks"))
            self.assertEqual(plotted,len(displayed["products"]["raw"]["points"]))
            self.bridge.setParameter("showPeaks",False);self.idle()
            self.assertEqual(self.bridge._peak_overlays,[]);self.assertEqual(canvas.draw_count,draws)
            self.bridge.setParameter("display",0);self.idle()
            self.assertNotEqual(tuple(canvas.figure.axes),axes)
            signature=self.bridge._static_peak_signature()
            self.bridge._maps[0][2].set_xlim(1.02,1.08)
            self.assertNotEqual(self.bridge._static_peak_signature(),signature)

        def test_map_linecut_preview_and_batch_settings_preserve_seeds(self):
            first, second = self.dataset(), self.dataset("second")
            second.settings = replace(second.settings,seed_energy=1.08,seed_y=3.)
            self.add(first,second)
            self.bridge.setParameter("peak.prominence",.11)
            self.bridge.invoke("apply_selected")
            self.idle()
            self.assertEqual(second.settings.prominence,.11)
            self.assertEqual((second.settings.seed_energy,second.settings.seed_y),(1.08,3.))
            canvas = self.bridge._canvas
            draws = canvas.draw_count
            self.bridge.setParameter("gate",2.)
            self.idle()
            self.assertEqual(canvas.draw_count,draws)
            self.assertGreater(canvas.blit_count,0)
            line = self.bridge._linecuts[0][-1]
            np.testing.assert_array_equal(line.get_ydata(),first.cube.Z[2])
            self.bridge.setParameter("peak.sg_window",31)
            self.assertNotEqual(self.bridge.preview_sg[first.key][0],31)
            self.bridge.invoke("preview")
            self.bridge.setParameter("display",0)
            self.idle()
            self.assertEqual(self.bridge._derivatives[first.key][0],(31,2))
            self.assertEqual(len(self.bridge._linecuts),2)
            self.bridge.setParameter("range.x_min",1.02)
            self.bridge.setParameter("range.x_max",1.09)
            self.idle()
            self.assertEqual(self.bridge._maps[0][2].get_xlim(),(1.02,1.09))

        def test_peak_exports_and_live_system_appearance(self):
            self.add(self.dataset(),self.dataset("second")); self.analyze()
            for kind in ("peaks","summary"):
                folder = self.root/kind
                self.bridge.exportTo(kind,str(folder)); self.idle()
                self.assertFalse(self.bridge.snapshot["error"])
                self.assertTrue(all(Path(path).exists() for path in self.bridge.snapshot["resultPaths"]))
            self.bridge._appearance.setMode("high-contrast"); self.idle()
            self.assertEqual(self.bridge.palette["inset"],"#000000")
            self.assertEqual(self.bridge._canvas.figure.get_facecolor(),(0.,0.,0.,1.))
            self.bridge._appearance.setMode("system"); self.idle()
            self.assertFalse(self.bridge.snapshot["error"])

        def test_duplicate_y_p2p_curve_pick_keeps_exact_source_row(self):
            from types import SimpleNamespace
            cube = DataCube(np.array([1.83,1.84,1.85,1.86]),np.array([1.,1.]),np.array([[0.,-1.,2.,0.],[0.,2.,-1.,0.]]),"Y","duplicate","DRR")
            dataset = create_dataset(cube,"duplicate",{})
            self.add(dataset)
            self.bridge.setParameter("tab",1); self.idle()
            self.bridge.setParameter("p2p.view_mode",1)
            self.bridge.setParameter("p2p.comparison_metric",1)
            self.bridge.setParameter("p2p.show_inspection",True); self.idle()
            line = next(line for axis in self.bridge._canvas.figure.axes for line in axis.lines if hasattr(line,"_drr_p2p"))
            coordinates = line.axes.transData.transform([1.,1.84])
            self.bridge._picked(SimpleNamespace(artist=line,ind=[0,1],mouseevent=SimpleNamespace(x=coordinates[0],y=coordinates[1])))
            self.assertEqual(self.bridge.p2p["inspection_rows"][dataset.key],2)
            spectrum = self.bridge._p2p_inspection[4]
            np.testing.assert_array_equal(spectrum.lines[0].get_ydata(),cube.Z[1])
            self.assertIn("P2P = 3",spectrum.get_title())

        def test_inspection_export_builds_captured_figure_on_worker(self):
            from ui_qt.quick import drr_analysis_plot as plots
            self.add(self.dataset());self.bridge.setParameter("tab",1);self.idle()
            self.bridge.setParameter("p2p.row",2);self.idle()
            axes=tuple(self.bridge._canvas.figure.axes);original=plots.inspection_export_figure
            gui_thread=threading.get_ident();entered=threading.Event();release=threading.Event();captured=[]
            def delayed(datasets,state,active_key,values,locked_colors):
                captured.append((threading.get_ident(),state["inspection_rows"][active_key]))
                entered.set();release.wait(10)
                return original(datasets,state,active_key,values,locked_colors)
            target=self.root/"worker-inspection.png"
            with patch.object(plots,"inspection_export_figure",side_effect=delayed):
                self.bridge.exportTo("inspection",str(target));wait(entered.is_set)
                self.assertNotEqual(captured[0][0],gui_thread)
                self.assertEqual(captured[0][1],2)
                self.bridge.setParameter("p2p.row",4)
                release.set();self.idle()
            self.assertTrue(target.is_file(),self.bridge.snapshot["error"])
            self.assertEqual(tuple(self.bridge._canvas.figure.axes),axes)
            self.assertEqual(self.bridge.p2p["inspection_rows"][self.bridge.active_key],4)

        def test_complete_p2p_state_range_smoothing_export_and_schema_one_roundtrip(self):
            first, second = self.dataset(), self.dataset("second",rows=4)
            self.add(first,second)
            self.bridge.setParameter("tab",1); self.idle()
            state = self.bridge.p2p
            self.assertEqual(len(state["records"]),2)
            self.assertFalse(state["pending"])
            self.bridge.setParameter("p2p.x_min",1.02)
            self.bridge.setParameter("p2p.x_max",1.08)
            self.bridge.setParameter("p2p.smoothing.method","SG")
            self.bridge.setParameter("p2p.smoothing.window_points",11)
            self.idle()
            self.assertTrue(all(record["processed_Z"] is not None for record in state["records"]))
            self.assertTrue(all(not record["processed_Z"].flags.writeable for record in state["records"]))
            snapshot = self.bridge._workspace_snapshot()
            self.assertIs(snapshot[3]["p2p"]["records"][0]["processed_Z"],state["records"][0]["processed_Z"])
            self.bridge.setParameter("p2p.view_mode",1)
            self.bridge.setParameter("p2p.comparison_metric",1)
            self.bridge.setParameter("p2p.show_inspection",True)
            self.bridge.setParameter("p2p.fixed_calculation",True)
            self.bridge.setParameter("p2p.fixed_display",True)
            self.bridge.setParameter("p2p.style.offset",True)
            self.idle()
            draws = self.bridge._canvas.draw_count
            self.bridge.setParameter("p2p.row",3); self.idle()
            self.assertEqual(self.bridge._canvas.draw_count,draws)
            self.assertGreater(self.bridge._canvas.blit_count,0)
            self.bridge.setParameter("p2p.smoothing.unit","meV"); self.idle()
            self.assertFalse(self.bridge.snapshot["error"])
            self.bridge.setParameter("p2p.smoothing.unit","points"); self.idle()
            self.assertGreaterEqual(len(self.bridge._canvas.figure.axes),5)
            for extension in ("csv","xlsx","png","inspection"):
                path = self.root/("result_"+extension+(".png" if extension == "inspection" else "."+extension))
                self.bridge.exportTo(extension,str(path)); self.idle()
                self.assertTrue(path.is_file(),self.bridge.snapshot["error"])
                self.assertGreater(path.stat().st_size,50)
            from openpyxl import load_workbook
            book = load_workbook(self.root/"result_xlsx.xlsx")
            self.assertEqual(book.sheetnames,["Amplitude","Emax","Emin","Separation","Metadata","Quality"])
            book.close()
            self.bridge.saveWorkspace(""); self.idle()
            saved = load_session(self.root/"workspace.npz",with_display=True)
            self.assertEqual(len(saved[0]),2)
            self.assertEqual(saved[3]["p2p"]["bounds"],[1.02,1.08,0.,5.])
            self.assertTrue(saved[3]["p2p"]["fixed_calculation"])
            self.assertEqual(saved[3]["p2p"]["inspection_rows"][first.key],3)
            self.bridge.invoke("remove_selected")
            self.bridge.restoreWorkspace(str(self.root/"workspace.npz")); self.idle()
            self.assertEqual(len(self.bridge.datasets),2)
            np.testing.assert_array_equal(self.bridge.p2p["records"][0]["processed_Z"],saved[3]["p2p"]["records"][0]["processed_Z"])

        def processed(self,field,spot="p5n2"):
            folder = self.root/"Processed Data"/"DRR"
            folder.mkdir(parents=True,exist_ok=True)
            path = folder/f"{spot}_{field}.metadata.json"
            source = f"YZ365_{spot}_{field}T_1.67KREF_720nmc_Rot49deg_TG-1.087BG=18.csv"
            meta = dict(operation="DR/R",processing=dict(derivative_order=None,baseline_selection="External",average_count=1),
                sources=[dict(role="measurement",filename=source,source_path=str(self.root/source))],plot=dict(ylabel="Gate (V)"))
            path.write_text(json.dumps(meta),encoding="utf-8")
            x = np.linspace(1,1.1,31)
            z = np.sin(x*100+field)
            dat = path.with_name(path.name.removesuffix(".metadata.json")+".dat")
            dat.write_text("Energy\t0\t1\t2\n"+"\n".join(f"{energy}\t{value}\t{value*2}\t{value*3}" for energy,value in zip(x,z)))
            return path

        def test_comparison_catalog_members_history_cache_and_source_versions(self):
            one,two,other = self.processed(0),self.processed(1),self.processed(2,"p3n2")
            self.bridge.refreshCatalog(); self.idle()
            self.assertEqual(self.bridge.catalogModel.rowCount(),3)
            key = next(key for key,group in self.bridge.comparison_catalog.items() if len(group["members"]) == 2)
            self.bridge.setParameter("selectedGroup",key)
            self.bridge.openComparison(key); self.idle()
            self.assertEqual(len(self.bridge.datasets),2)
            first_identity = self.bridge.dataset
            self.bridge.openComparison(key); self.idle()
            self.assertIs(self.bridge.dataset,first_identity)
            self.bridge.setParameter("tab",1); self.idle()
            self.bridge.saveWorkspace(""); self.idle()
            self.assertTrue(self.bridge.historyModel.rows)
            history = self.bridge.historyModel.rows[0]["id"]
            from core.drr_comparison_groups import product_id
            self.bridge.setMember(product_id(two),False)
            self.bridge.invoke("apply_members"); self.idle()
            self.assertEqual(len(self.bridge.datasets),1)
            self.assertFalse(next(row for row in self.bridge.membersModel.rows if row["id"] == product_id(two))["included"])
            self.bridge.restoreWorkspace(history); self.idle()
            self.assertEqual(len(self.bridge.datasets),2)
            self.bridge.setMember(product_id(two),True)
            self.bridge.openComparison(key); self.idle()
            old = {dataset.key for dataset in self.bridge.datasets.values()}
            dat = one.with_name(one.name.removesuffix(".metadata.json")+".dat")
            dat.write_text(dat.read_text().replace("\t0\t1\t2","\t0\t1\t3"))
            self.bridge.openComparison(key); self.idle()
            self.assertNotEqual({dataset.key for dataset in self.bridge.datasets.values()},old)

        def test_raw_recipe_reload_and_snapshot_handoff(self):
            sample = self.root/"sample_TG-1.087BG=18.csv"
            sample.write_text("Vbg,Vtg,700,710,720,730,740\n0,18,2,4,6,4,2\n1,19.087,4,8,12,8,4\n")
            background = self.root/"background.csv"
            background.write_text("Vbg,Vtg,700,710,720,730,740\n0,0,1,2,3,2,1\n0,0,3,6,9,6,3\n")
            before = sample.read_bytes(),background.read_bytes()
            self.bridge.openFiles("backgrounds",[str(background)])
            self.bridge.openFiles("measurements",[str(sample)]); self.idle()
            dataset = self.bridge.dataset
            self.assertEqual(dataset.provenance["background_mode"],"external")
            self.bridge.setParameter("peak.seed_y",1.)
            self.bridge.setParameter("recipeMode","self_last")
            # Self recipes intentionally ignore stale external selection.
            self.bridge.invoke("reload_selected"); self.idle()
            self.assertNotEqual(self.bridge.dataset.key,dataset.key)
            self.assertEqual(self.bridge.dataset.settings.seed_y,1.)
            self.assertEqual((sample.read_bytes(),background.read_bytes()),before)
            snapshot = self.root/"handoff.npz"
            save_session(snapshot,[self.dataset("snapshot")])
            self.bridge.receiveSnapshot(str(snapshot)); self.idle()
            self.assertEqual(len(self.bridge.datasets),2)
            self.assertTrue(snapshot.exists())

        def test_dialog_folders_follow_source_workspace_and_last_export(self):
            source = self.root/"experiment #1 測量"
            sessions = self.root/"saved workspaces"
            exports = self.root/"exports %1"
            for folder in (source, sessions, exports): folder.mkdir()
            self.bridge.start_folder = str(source)
            self.bridge.session_path = sessions/"workspace.npz"
            self.bridge._result_paths = [str(exports/"curve.csv")]
            register_plot_surface()
            engine = QQmlApplicationEngine()
            engine.rootContext().setContextProperty("drrAnalysisBridge", self.bridge)
            engine.load(QUrl.fromLocalFile(str(Path(__file__).resolve().parents[1]/"ui_qt/quick/qml/DrrAnalysisApplication.qml")))
            self.assertTrue(engine.rootObjects())
            window = engine.rootObjects()[0]
            try:
                for kind, expected in (("measurements", source), ("backgrounds", source), ("processed", source),
                        ("members", source), ("restore", sessions), ("save", sessions), ("snapshot", sessions),
                        ("export_png", exports), ("folder", source), ("export_peaks", exports), ("export_summary", exports)):
                    with self.subTest(kind=kind):
                        self.bridge.sourceDialogRequested.emit(kind, kind in {"measurements", "backgrounds", "members", "processed", "snapshot"})
                        folder_role = kind in {"folder", "export_peaks", "export_summary"}
                        picker = window.findChild(QObject, "drr_analysis_folder_dialog" if folder_role else "drr_analysis_file_dialog")
                        wait(lambda: picker.property("visible"))
                        self.assertEqual(Path(picker.property("currentFolder").toLocalFile()).resolve(), expected.resolve())
                        if kind == "save":
                            picker.setProperty("selectedFile", QUrl.fromLocalFile(str(sessions/"cancelled-copy.npz")))
                        if kind == "snapshot":
                            self.assertNotIn("cancelled-copy.npz", [url.fileName() for url in picker.property("selectedFiles")])
                        self.assertTrue(QMetaObject.invokeMethod(picker, "reject"))
                        wait(lambda: not picker.property("visible"))
                self.assertEqual(self.bridge.datasets, {})
                self.assertFalse(self.bridge._dirty)
                self.bridge.start_folder = str(self.root)
                self.bridge._publish()
                app.processEvents()
                self.bridge.sourceDialogRequested.emit("measurements", True)
                picker = window.findChild(QObject, "drr_analysis_file_dialog")
                wait(lambda: picker.property("visible"))
                self.assertEqual(Path(picker.property("currentFolder").toLocalFile()).resolve(), self.root.resolve())
                self.assertTrue(QMetaObject.invokeMethod(picker, "reject"))
            finally:
                window.setProperty("allowClose", True)
                window.close()
                engine.deleteLater()
                app.sendPostedEvents(None, QEvent.DeferredDelete)
                app.processEvents()

        def test_qml_full_workspace_and_nonnative_picker_cancel(self):
            self.processed(0); self.processed(1)
            self.bridge.refreshCatalog(); self.idle()
            self.bridge.setParameter("selectedGroup",next(iter(self.bridge.comparison_catalog)))
            register_plot_surface()
            engine = QQmlApplicationEngine()
            engine.rootContext().setContextProperty("drrAnalysisBridge",self.bridge)
            engine.load(QUrl.fromLocalFile(str(Path(__file__).resolve().parents[1]/"ui_qt/quick/qml/DrrAnalysisApplication.qml")))
            self.assertTrue(engine.rootObjects())
            window = engine.rootObjects()[0]
            self.add(self.dataset())
            self.analyze()
            self.bridge.setParameter("tab",1); self.idle()
            self.bridge.setParameter("p2p.view_mode",1)
            self.bridge.setParameter("p2p.comparison_metric",2); self.idle()
            library = window.findChild(QObject,"drr_analysis_library")
            tabs = window.findChild(QObject,"drr_analysis_library_tabs")
            self.assertTrue(QMetaObject.invokeMethod(library,"open"))
            for tab in (0,1,2): tabs.setProperty("currentIndex",tab); app.processEvents()
            self.assertTrue(QMetaObject.invokeMethod(library,"close"))
            from tests.quick_selection_checks import assert_text_selection_palette
            for mode in ("light", "dark", "high-contrast"):
                self.bridge._appearance.setMode(mode); self.idle(); app.processEvents()
                assert_text_selection_palette(self, window, self.bridge.palette)
            self.bridge._appearance.setMode("light"); self.idle()
            for action,name in (("add_files","drr_analysis_file_dialog"),("folder","drr_analysis_folder_dialog")):
                self.bridge.invoke(action)
                picker = window.findChild(QObject,name)
                wait(lambda: picker.property("visible"))
                self.assertTrue(QMetaObject.invokeMethod(picker,"reject"))
                wait(lambda: not picker.property("visible"))
                for _ in range(10): app.processEvents()
            window.setProperty("allowClose",True); window.close()
            engine.deleteLater(); app.sendPostedEvents(None,QEvent.DeferredDelete); app.processEvents()

        def test_entry_legacy_ipc_forwarding_and_native_snapshot_startup(self):
            from PySide6.QtNetwork import QLocalServer
            from run_drr_analysis_quick import server_name
            snapshot = self.root/"incoming.npz"
            save_session(snapshot,[self.dataset("incoming")])
            server = QLocalServer()
            server.setSocketOptions(QLocalServer.UserAccessOption)
            self.assertTrue(server.listen(server_name(self.root/"ipc-session.npz")))
            received, clients = [], []
            def connected():
                client = server.nextPendingConnection(); clients.append(client)
                buffer = bytearray()
                def read():
                    buffer.extend(bytes(client.readAll()))
                    if b"\n" in buffer:
                        received.append(json.loads(bytes(buffer).split(b"\n")[0])); client.disconnectFromServer()
                client.readyRead.connect(read); read()
            server.newConnection.connect(connected)
            args = [sys.executable,"run_drr_analysis_quick.py",str(self.root),"--session",str(self.root/"ipc-session.npz"),
                "--snapshot",str(snapshot),"--settings-file",str(self.root/"ipc-settings.ini")]
            process = subprocess.Popen(args,stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True)
            wait(lambda: process.poll() is not None)
            stdout,stderr = process.communicate()
            self.assertEqual(process.returncode,0,stdout+stderr)
            wait(lambda: bool(received))
            self.assertEqual(received[0]["snapshot"],str(snapshot.resolve()))
            self.assertFalse((self.root/"ipc-session.npz").exists())
            server.close()
            report = self.root/"entry-report.json"
            startup = subprocess.Popen([sys.executable,"run_drr_analysis_quick.py",str(self.root),"--session",str(self.root/"native-session.npz"),
                "--snapshot",str(snapshot),"--settings-file",str(self.root/"native-settings.ini"),"--check","--report",str(report)],
                stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True)
            wait(lambda: startup.poll() is not None,timeout=90)
            stdout,stderr = startup.communicate()
            self.assertEqual(startup.returncode,0,stdout+stderr)
            state = json.loads(report.read_text())
            self.assertEqual(state["application"],"QGuiApplication")
            self.assertFalse(state["widgetsImported"])
            self.assertEqual(state["warnings"],[])
            self.assertEqual(state["datasets"],1)
            self.assertGreater(state["renderedFrames"],0)
            self.assertEqual(len(load_session(self.root/"native-session.npz")[0]),1)

        def test_save_failure_preserves_file_and_close_waits_for_save(self):
            self.add(self.dataset())
            target = self.root/"workspace.npz"
            target.write_bytes(b"original")
            with patch("core.drr_workspace_session.save_session",side_effect=OSError("disk unavailable")):
                self.bridge.saveWorkspace(""); self.idle()
            self.assertEqual(target.read_bytes(),b"original")
            self.assertTrue(self.bridge._dirty)
            ready = []
            self.bridge.closeReady.connect(lambda: ready.append(True))
            self.bridge.requestClose()
            # A previous IO error must be surfaced; explicit retry can recover.
            self.assertFalse(ready)
            self.bridge.saveWorkspace(""); self.idle()
            self.bridge.requestClose(); wait(lambda: bool(ready))
            self.assertFalse(self.bridge._dirty)
            self.assertEqual(len(load_session(target)[0]),1)

        def test_corrupt_session_is_preserved_and_copy_does_not_move_owner(self):
            target = self.root/"workspace.npz"
            target.write_bytes(b"corrupt original")
            self.bridge._session_protected = True
            self.add(self.dataset())
            self.assertFalse(self.bridge._autosave_timer.isActive())
            dialogs,ready = [],[]
            self.bridge.sourceDialogRequested.connect(lambda kind,multiple:dialogs.append(kind))
            self.bridge.closeReady.connect(lambda:ready.append(True))
            self.bridge.requestClose()
            self.assertEqual(dialogs,["save"])
            self.bridge.dialogRejected("save")
            self.assertFalse(self.bridge._closing)
            alternate = self.root/"saved-copy.npz"
            self.bridge.saveWorkspace(str(alternate));self.idle()
            self.assertEqual(target.read_bytes(),b"corrupt original")
            self.assertEqual(self.bridge.session_path,target)
            self.assertEqual(len(load_session(alternate)[0]),1)
            self.bridge.requestClose();wait(lambda:bool(ready))

    suite = unittest.defaultTestLoader.loadTestsFromTestCase(Cases)
    result = unittest.TextTestRunner(verbosity=2).run(suite)
    qInstallMessageHandler(previous_handler)
    print("QUICK_DRR_ANALYSIS_TESTS="+json.dumps(dict(tests=result.testsRun,failures=len(result.failures),errors=len(result.errors),widgetsImported="PySide6.QtWidgets" in sys.modules)))
    return 0 if result.wasSuccessful() else 1


if __name__ == "__main__":
    raise SystemExit(isolated()) if "--isolated" in sys.argv else unittest.main()
