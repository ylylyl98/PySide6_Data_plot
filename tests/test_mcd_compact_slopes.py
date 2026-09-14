import json
from pathlib import Path
import tempfile
import unittest
from openpyxl import load_workbook
from core.mcd_unified_export import export_mcd_analysis
from tests.test_mcd_unified_export import _snapshot


class CompactSlopesTests(unittest.TestCase):
    def test_compact_sheet_keeps_diagnostics_in_json_and_low_minus_high(self):
        snapshot = _snapshot()
        snapshot['windows'] = [snapshot['windows'][0]]
        snapshot['settings'].update(slope_low_t=-.2, slope_low_end_t=.2,
            slope_high_negative_t=-2., slope_high_negative_end_t=-1.5,
            slope_high_positive_t=1.5, slope_high_positive_end_t=2.)
        snapshot['slopes'] = [
            {'window_id': 'w1', 'region': region, 'branch': branch, 'slope': value,
             'slope_se': .01, 'status': 'ok', 'n': 7, 'residual_rms': .001,
             'jump_flag': False, 'curvature_flag': True, 'intercept': .2}
            for region, branch, value in (
                ('low', 'B increasing', .08), ('low', 'B decreasing', .07),
                ('high_negative', 'B increasing', .02), ('high_negative', 'B decreasing', .03),
                ('high_positive', 'B increasing', .01), ('high_positive', 'B decreasing', .04))]
        snapshot['slopes'].append({'window_id': 'w1', 'comparison': 'high_minus_low',
            'region': 'high_negative', 'reference_region': 'low', 'branch_a': 'B increasing',
            'branch_b': 'B increasing', 'slope_difference': -.06, 'status': 'unsupported'})
        with tempfile.TemporaryDirectory() as folder:
            result = export_mcd_analysis(snapshot, folder)
            book = load_workbook(result['xlsx'], data_only=True)
            sheet = book['Slopes']
            self.assertEqual(sheet.max_column, 6)
            self.assertEqual([c.value for c in sheet[7]], ['Region', 'B range (T)', 'Inc slope', 'Inc SE', 'Dec slope', 'Dec SE'])
            self.assertEqual(sheet['A8'].value, 'Near zero')
            self.assertEqual(sheet['B8'].value, '-0.2 to 0.2')
            self.assertAlmostEqual(sheet['C8'].value, .08)
            self.assertEqual(sheet['A13'].value, 'Near-zero minus negative high field')
            self.assertAlmostEqual(sheet['B13'].value, .06)
            self.assertAlmostEqual(sheet['C13'].value, .04)
            self.assertAlmostEqual(sheet['B14'].value, .07)
            self.assertAlmostEqual(sheet['C14'].value, .03)
            text = str(list(sheet.values))
            for hidden in ('window_id', 'fit_id', 'residual_rms', 'curvature_flag', 'unsupported'):
                self.assertNotIn(hidden, text)
            book.close()
            meta = json.loads(Path(result['metadata_json']).read_text())
            self.assertEqual(meta['slopes'][0]['n'], 7)
            self.assertTrue(meta['slopes'][0]['curvature_flag'])
            self.assertEqual(meta['slopes'][-1]['comparison'], 'low_minus_high')
            self.assertAlmostEqual(meta['slopes'][-1]['slope_difference'], .06)
