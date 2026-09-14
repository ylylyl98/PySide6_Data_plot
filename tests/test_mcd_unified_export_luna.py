import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import numpy as np
from PIL import Image

from core.mcd_unified_export import _plot_pngs, export_mcd_analysis


def _snapshot():
    fields = np.asarray([-1.0, 0.0, 1.0])
    energy = np.asarray([1.5, 1.6, 1.7])
    branches = np.asarray(["B decreasing", "B increasing", "B increasing"])
    values = np.asarray([[-0.1, -0.08, -0.06], [0.0, 0.01, 0.02], [0.1, 0.11, 0.12]])
    return {
        "source_descriptor": {
            "filename": "YZ365_p5n2_MCD_2p5K_G01_D6p3_F20_r1p087.csv",
            "name": "YZ365_p5n2_MCD_2p5K_G01_D6p3_F20_r1p087.csv",
        },
        "mcd": {
            "energy_ev": energy,
            "fields_t": fields,
            "branches": branches,
            "values": values,
            "metric": "corrected signed mean",
            "channel": "MCD",
        },
        "windows": [{"window_id": "w1", "values": [0.0, 0.01, 0.11], "metric": "mean", "center_ev": 1.6, "width_mev": 5.0}],
        "slopes": [{"branch": "B increasing", "region": "low", "slope": 0.1, "intercept": 0.01, "actual_range_t": [0.0, 1.0], "status": "ok"}],
        "analysis_results": [{
            "tracking_method": "Raw spectrum",
            "source": "Raw",
            "tracks": [
                {"feature_id": "p+", "feature_kind": "peak", "channel": "raw pos", "branch": "B increasing", "reference_energy_ev": 1.6,
                 "points": [{"field_t": -1.0, "energy_ev": 1.59, "delta_energy_ev": -0.01, "status": "tracked"},
                            {"field_t": 0.0, "energy_ev": 1.6, "delta_energy_ev": 0.0, "status": "tracked"},
                            {"field_t": 1.0, "energy_ev": 1.61, "delta_energy_ev": 0.01, "status": "tracked"}]},
                {"feature_id": "p-", "feature_kind": "peak", "channel": "raw neg", "branch": "B decreasing", "reference_energy_ev": 1.61,
                 "points": [{"field_t": -1.0, "energy_ev": 1.62, "delta_energy_ev": 0.01, "status": "tracked"},
                            {"field_t": 0.0, "energy_ev": None, "delta_energy_ev": None, "status": "missing"},
                            {"field_t": 1.0, "energy_ev": 1.60, "delta_energy_ev": -0.01, "status": "tracked"}]},
            ],
        }],
        "splitting": [{"branch": "B increasing", "pair_id": "ambiguous", "status": "ambiguous", "points": [{"field_t": 0.0, "splitting_ev": 0.01}]}],
        "plot_state": {"show_raw": True, "show_corrected": False, "visible_branches": ["B increasing", "B decreasing"]},
        "settings": {"mcd_center_ev": 1.6, "mcd_width_mev": 5.0},
    }


class LunaMcdExportTests(unittest.TestCase):
    def test_default_pngs_are_fixed_size_and_split_feature_products(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = export_mcd_analysis(_snapshot(), tmp)
            names = {Path(path).name for path in output["plot_pngs"]}
            self.assertEqual(len(names), 4)
            self.assertTrue(any("MCD_map" in name for name in names))
            self.assertTrue(any("MCD_vs_B" in name for name in names))
            self.assertTrue(any("feature_energy" in name for name in names))
            self.assertTrue(any("feature_shift" in name for name in names))
            self.assertFalse(any("spectra" in name for name in names))
            for path in output["plot_pngs"]:
                with Image.open(path) as image:
                    self.assertEqual(image.size, (1200, 900))

    def test_map_is_corrected_even_when_spectrum_toggle_is_raw_only(self):
        snapshot = _snapshot()
        snapshot["mcd"]["raw_values"] = snapshot["mcd"]["values"] + 100.0
        snapshot["plot_state"] = {"show_raw": True, "show_corrected": False}
        from unittest.mock import patch
        from matplotlib.figure import Figure
        captured = []
        with tempfile.TemporaryDirectory() as tmp, patch.object(Figure, "savefig", lambda figure, *_a, **_k: captured.append(figure)):
            paths = _plot_pngs(Path(tmp), snapshot["mcd"], snapshot, "r01")
            self.assertEqual(Path(paths[0]).name, "MCD_map_r01.png")
            self.assertTrue(captured[0].axes[0].collections)
            plotted = np.concatenate([collection.get_array() for collection in captured[0].axes[0].collections])
            self.assertTrue(np.allclose(plotted, snapshot["mcd"]["values"].ravel()))

    def test_invalid_splitting_pair_is_not_exported(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = export_mcd_analysis(_snapshot(), tmp)
            from openpyxl import load_workbook
            workbook = load_workbook(output["xlsx"], data_only=True)
            self.assertNotIn("Splitting", workbook.sheetnames)

    def test_valid_paired_splitting_gets_its_own_png_and_sheet(self):
        snapshot = _snapshot()
        snapshot["splitting"] = [{
            "branch": "B increasing", "pair_id": "p+-",
            "selected_feature_id": "p+", "counterpart_feature_id": "p-",
            "selected_channel": "raw pos", "counterpart_channel": "raw neg",
            "method": "Raw spectrum", "status": "valid",
            "points": [{"field_t": 0.0, "splitting_ev": 0.01, "status": "tracked"}],
        }]
        with tempfile.TemporaryDirectory() as tmp:
            output = export_mcd_analysis(snapshot, tmp)
            self.assertEqual(len(output["plot_pngs"]), 5)
            splitting_png = next(path for path in output["plot_pngs"] if "splitting_" in path)
            with Image.open(splitting_png) as image:
                self.assertEqual(image.size, (1200, 900))
            from openpyxl import load_workbook
            workbook = load_workbook(output["xlsx"], data_only=True)
            self.assertIn("Splitting", workbook.sheetnames)

    def test_feature_labels_keep_channel_method_identity_and_truthful_e0_method(self):
        snapshot = _snapshot()
        snapshot["feature_points"] = [
            {"feature_id": "manual-track", "feature_kind": "peak", "channel": "K",
             "analysis_method": "raw", "branch": "B increasing", "field_t": 0.0,
             "energy_ev": 1.6, "delta_energy_ev": 0.0, "status": "tracked",
             "reference_energy_ev": 1.6, "reference_field_t": 0.0,
             "reference_method": "manual"},
            {"feature_id": "nearest-track", "feature_kind": "dip", "channel": "Kp",
             "analysis_method": "second derivative", "branch": "B increasing", "field_t": 1.0,
             "energy_ev": 1.7, "delta_energy_ev": 0.1, "status": "tracked",
             "reference_energy_ev": 1.6, "reference_field_t": 0.0,
             "reference_method": "nearest"},
        ]
        snapshot["plot_state"] = {"visible_branches": ["B increasing"], "visible_features": None}
        from matplotlib.figure import Figure
        captured = []
        with tempfile.TemporaryDirectory() as tmp, patch.object(Figure, "savefig", lambda figure, *_a, **_k: captured.append(figure)):
            _plot_pngs(Path(tmp), snapshot["mcd"], snapshot, "r01")
        energy_axis = next(figure.axes[0] for figure in captured if figure.axes[0].get_ylabel() == "Energy (eV)")
        labels = energy_axis.get_legend_handles_labels()[1]
        self.assertTrue(any("manual-track" in label and "raw" in label and "manual" in label for label in labels))
        self.assertTrue(any("nearest-track" in label and "2nd deriv" in label and "near" in label for label in labels))
        self.assertNotIn("interpolated", " ".join(labels).casefold())

    def test_multiple_windows_label_window_identity_metric_and_units(self):
        snapshot = _snapshot()
        snapshot["windows"] = [
            {"window_id": "near_peak", "center_ev": 1.5, "width_mev": 20.0, "metric": "mean"},
            {"window_id": "broad_area", "center_ev": 1.6, "width_mev": 20.0, "metric": "integral"},
        ]
        with tempfile.TemporaryDirectory() as tmp:
            output = export_mcd_analysis(snapshot, tmp)
            from openpyxl import load_workbook
            workbook = load_workbook(output["xlsx"], data_only=True)
            headers = [cell.value for cell in workbook["MCD"][1]]
            self.assertTrue(any("near_peak" in str(header) and "mean" in str(header) and "MCD" in str(header) for header in headers))
            self.assertTrue(any("broad_area" in str(header) and "integral" in str(header) and "MCD" in str(header) for header in headers))
            self.assertEqual(workbook["Slopes"]["A5"].value, "Slope unit")
            self.assertTrue(workbook["Slopes"]["B5"].value)


if __name__ == "__main__":
    unittest.main()
