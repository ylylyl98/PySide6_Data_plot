"""Native Slides parity against temporary plots/decks and isolated settings."""
from __future__ import annotations

import hashlib
import json
import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ.setdefault("QT_QUICK_BACKEND", "software")
os.environ.setdefault("QT_QUICK_CONTROLS_STYLE", "Basic")

from pathlib import Path
from tempfile import TemporaryDirectory
from threading import Event
import unittest
from unittest.mock import patch

from PIL import Image
from pptx import Presentation
from core.presentation import BuildResult
from PySide6.QtCore import QMetaObject, QObject, QPoint, QSettings, Qt, QUrl
from PySide6.QtGui import QFont, QFontDatabase, QGuiApplication
from PySide6.QtQuick import QQuickItem, QQuickView
from PySide6.QtQuickControls2 import QQuickStyle
from PySide6.QtTest import QTest

import ui_qt.quick.presentation_bridge as slides
from ui_qt.quick.presentation_bridge import PresentationBridge
from ui_qt.quick.appearance import quick_colors


COLORS = quick_colors("light")


class QuickPresentationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if QGuiApplication.instance() is None:
            QQuickStyle.setStyle("Basic")
        cls.app = QGuiApplication.instance() or QGuiApplication([])
        if os.name == "nt" and not QFontDatabase.families():
            font = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "segoeui.ttf"
            if font.is_file():
                QFontDatabase.addApplicationFont(str(font))
                cls.app.setFont(QFont("Segoe UI", 9))

    def setUp(self):
        self.temporary = TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.settings = QSettings(str(self.root / "settings.ini"), QSettings.IniFormat)
        self.patches = [patch.object(slides, "powerpoint_integration_available", return_value=(False, "Live integration is unavailable in this test.")),
                        patch.object(slides, "powerpoint_presentation_is_open", return_value=False)]
        for item in self.patches:
            item.start()
        self.bridge = PresentationBridge(settings=self.settings)

    def tearDown(self):
        self.bridge.close()
        self.bridge._executor.shutdown(wait=True)
        self.app.processEvents()
        for item in reversed(self.patches):
            item.stop()
        self.temporary.cleanup()

    def wait(self, predicate, message="worker did not finish"):
        for _ in range(600):
            self.app.processEvents()
            if predicate():
                return
            QTest.qWait(10)
        self.fail(message)

    def png(self, relative, number):
        path = self.root / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        Image.new("RGB", (160, 90), (number % 255, (number * 23) % 255, (number * 71) % 255)).save(path)
        os.utime(path, (100 + number, 100 + number))
        return path.resolve()

    def discover(self, path=None):
        self.bridge.setOption("imageRoot", str(path or self.root))
        self.wait(lambda: not self.bridge.snapshot["scanning"])
        self.assertFalse(self.bridge.snapshot["error"], self.bridge.snapshot["error"])

    def build(self, operation="copy"):
        self.bridge.build(operation)
        self.wait(lambda: not self.bridge.snapshot["building"], "presentation build did not finish")
        self.assertFalse(self.bridge.snapshot["error"], self.bridge.snapshot["error"])

    def deck(self):
        path = self.root / "source.pptx"
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[6])
        deck.save(str(path))
        return path

    def test_discovery_filters_complete_names_and_newest_mcd_pair(self):
        old = self.png("MCD/older/run_MCD_Combo_map.png", 1)
        self.png("MCD/older/run_MCD_vs_B_E1.png", 2)
        combo = self.png("MCD/newer/run_MCD_Combo_map.png", 3)
        trace = self.png("MCD/newer/run_MCD_vs_B_E2.png", 4)
        long_name = "complete_measurement_" * 9 + ".png"
        long = self.png("PL/run/" + long_name, 5)
        self.discover()
        self.assertEqual(self.bridge.availableModel.rows[0]["path"], str(long))
        self.assertEqual(self.bridge.availableModel.rows[0]["name"], long_name)
        self.bridge.setOption("plotKind", "mcd_combo")
        self.assertEqual({row["path"] for row in self.bridge.availableModel.rows}, {str(combo), str(old)})
        self.bridge.addNewestMcdPair()
        self.assertEqual(self.bridge._queue, [str(combo), str(trace)])
        self.assertIn("same folder", self.bridge.snapshot["mcdStatus"])
        self.assertTrue(all(row["badge"].startswith("F") for row in self.bridge.queueModel.rows))
        self.assertTrue(list((self.root / "slides-catalog").glob("*.json")))

    def test_selection_order_unique_queue_grouping_and_cached_preview(self):
        first = self.png("MCD/D1/run_D1_F0_MCD_vs_B_E1.57eV_W5meV.png", 1)
        other = self.png("MCD/D2/run_D2_F0_MCD_vs_B_E1.60eV_W5meV.png", 2)
        second = self.png("MCD/D1_F20/run_D1_F20_MCD_vs_B_E1.64eV_W5meV.png", 3)
        unselected = self.png("MCD/D1_F20/run_D1_F20_MCD_vs_B_E1.66eV_W5meV.png", 4)
        self.discover()
        indices = {row["path"]: index for index, row in enumerate(self.bridge.availableModel.rows)}
        self.bridge.selectAvailableRow(indices[str(first)], 0)
        self.bridge.selectAvailableRow(indices[str(other)], Qt.ControlModifier.value)
        self.bridge.selectAvailableRow(indices[str(second)], Qt.ControlModifier.value)
        with patch.object(slides, "plan_presentation_slides", wraps=slides.plan_presentation_slides) as planner:
            self.bridge.addSelected()
            self.assertEqual(planner.call_count, 1)
            self.bridge.setOption("title", "Experiment")
            self.bridge.setOption("caption", "full")
            self.bridge.setPreviewSlide(0)
            self.assertEqual(planner.call_count, 1)
        self.assertEqual(self.bridge._queue, [str(first), str(other), str(second)])
        self.assertNotIn(str(unselected), self.bridge._queue)
        self.assertEqual([len(plan.images) for plan in self.bridge._plans], [2, 1])
        self.assertIn("Doping = 1 V", self.bridge.snapshot["preview"]["title"])
        self.bridge.addSelected()
        self.assertEqual(len(self.bridge._queue), 3)
        self.bridge.selectQueueRow(0, 0)
        self.bridge.moveSelected(1)
        self.assertEqual(self.bridge._queue[1], str(first))
        self.bridge.moveQueue(1, 2)
        self.assertEqual(self.bridge._queue[2], str(first))
        self.bridge.removeSelected()
        self.assertNotIn(str(first), self.bridge._queue)
        self.bridge.clearQueue()
        self.assertEqual(self.bridge.snapshot["preview"]["images"], [])

    def test_actual_powerpoint_copy_append_deduplication_and_png_integrity(self):
        paths = [self.png(f"PL/run/plot_{index:02d}.png", index) for index in range(13)]
        png_hashes = [hashlib.sha256(path.read_bytes()).hexdigest() for path in paths]
        source = self.deck()
        source_bytes = source.read_bytes()
        self.discover()
        self.bridge.addAll()
        self.bridge.setOption("source", str(source))
        self.bridge.setOption("output", str(self.root / "copy.pptx"))
        self.bridge.setOption("groupBy", "queue")
        self.bridge.setOption("caption", "full")
        self.bridge.setOption("panelLabels", True)
        self.build()
        output = Path(self.bridge.snapshot["lastOutput"])
        self.assertEqual(len(Presentation(str(output)).slides), 3)
        manifest = json.loads(slides.manifest_path_for(output).read_text(encoding="utf-8"))
        self.assertEqual(len(manifest["images"]), 13)
        self.assertEqual(manifest["source_presentation"], str(source))
        self.assertEqual(manifest["output_presentation"], str(output))
        self.assertEqual(source.read_bytes(), source_bytes)
        self.assertEqual([hashlib.sha256(path.read_bytes()).hexdigest() for path in paths], png_hashes)
        self.build()
        self.assertEqual(len(Presentation(str(output)).slides), 3)
        self.assertIn("Nothing duplicated", self.bridge.snapshot["status"])
        self.assertFalse(list(self.root.glob(".dptk-slide-build-*")))

    def test_in_place_build_preserves_recovery_backup_and_continuation(self):
        self.png("PL/run/plot.png", 7)
        source = self.deck()
        before = source.read_bytes()
        self.discover()
        self.bridge.addAll()
        self.bridge.setOption("source", str(source))
        self.build("auto_save")
        self.assertEqual(len(Presentation(str(source)).slides), 2)
        backup = source.with_name("source.dptk-backup.pptx")
        self.assertEqual(backup.read_bytes(), before)
        self.build("auto_save")
        self.assertEqual(len(Presentation(str(source)).slides), 2)
        self.assertEqual(backup.read_bytes(), before)
        manifest = json.loads(slides.manifest_path_for(source).read_text(encoding="utf-8"))
        self.assertEqual(len(manifest["images"]), 1)
        self.assertEqual(manifest["source_presentation"], str(source))

    def test_cancelled_offline_build_does_not_publish_or_change_destination(self):
        self.png("PL/run/plot.png", 7)
        source = self.deck()
        before = source.read_bytes()
        self.discover()
        self.bridge.addAll()
        self.bridge.setOption("source", str(source))
        entered, release = Event(), Event()
        original = slides.build_presentation
        def blocked(**kwargs):
            entered.set()
            release.wait(5)
            return original(**kwargs)
        with patch.object(slides, "build_presentation", side_effect=blocked):
            self.bridge.build("auto_save")
            self.wait(entered.is_set)
            self.bridge.cancel()
            release.set()
            self.wait(lambda: not self.bridge.snapshot["building"])
        self.assertEqual(source.read_bytes(), before)
        self.assertFalse(slides.manifest_path_for(source).exists())
        self.assertIn("cancelled", self.bridge.snapshot["status"])
        self.assertFalse(list(self.root.glob(".dptk-slide-build-*")))

    def test_live_route_uses_domain_service_and_error_messages_are_visible(self):
        self.bridge.build("copy")
        self.assertIn("Add at least one", self.bridge.snapshot["error"])
        self.png("PL/run/plot.png", 1)
        self.discover()
        self.bridge.addAll()
        self.bridge.build("auto_save")
        self.assertIn("existing PowerPoint", self.bridge.snapshot["error"])
        self.bridge.build("live")
        self.assertTrue(self.bridge.snapshot["error"])
        source = self.deck()
        self.bridge.setOption("source", str(source))
        self.bridge._live_available = True
        with patch.object(slides, "insert_plots_into_open_powerpoint", side_effect=RuntimeError("PowerPoint is not currently open.")) as live:
            self.bridge.build("live")
            self.wait(lambda: not self.bridge.snapshot["building"])
        live.assert_called_once()
        self.assertFalse(live.call_args.kwargs["save"])
        self.assertIn("not currently open", self.bridge.snapshot["error"])
        unsaved = BuildResult(source, 1, 1, 0, 2, slides.manifest_path_for(source), live_edit=True, saved=False)
        with patch.object(slides, "insert_plots_into_open_powerpoint", return_value=unsaved) as live:
            self.bridge.build("live")
            self.wait(lambda: not self.bridge.snapshot["building"])
        self.assertFalse(live.call_args.kwargs["save"])
        self.assertIn("Save it in PowerPoint", self.bridge.snapshot["status"])
        saved = slides.replace(unsaved, saved=True)
        with (patch.object(slides, "powerpoint_presentation_is_open", return_value=True),
              patch.object(slides, "insert_plots_into_open_powerpoint", return_value=saved) as live,
              patch.object(slides, "_offline_build") as offline):
            self.bridge.build("auto_save")
            self.wait(lambda: not self.bridge.snapshot["building"])
        self.assertTrue(live.call_args.kwargs["save"])
        offline.assert_not_called()

    def test_settings_and_auto_root_match_session_and_splitter_semantics(self):
        processed = self.root / "experiment" / "Processed Data"
        processed.mkdir(parents=True)
        self.bridge.setExperimentFolder(str(processed.parent))
        self.wait(lambda: not self.bridge.snapshot["scanning"])
        self.assertEqual(self.bridge.snapshot["imageRoot"], str(processed))
        custom = self.root / "chosen"
        custom.mkdir()
        self.bridge.setOption("imageRoot", str(custom))
        self.wait(lambda: not self.bridge.snapshot["scanning"])
        self.bridge.setExperimentFolder(str(self.root))
        self.assertEqual(self.bridge.snapshot["imageRoot"], str(custom))
        self.bridge.setSplitterSizes([500, 250, 350])
        self.bridge.setOption("caption", "full")
        self.settings.sync()
        self.assertEqual(self.settings.allKeys(), ["slides/workspace_splitter_sizes"])
        other = PresentationBridge(settings=self.settings)
        try:
            self.assertEqual(other.snapshot["splitterSizes"], [500, 250, 350])
            self.assertEqual(other.snapshot["caption"], "none")
        finally:
            other.close()

    def test_dialog_folders_follow_deck_copy_and_plot_paths(self):
        folders = [self.root / name for name in ("decks #1 測量", "copies %1", "plots", "other deck")]
        for folder in folders:
            folder.mkdir()
        source, output = folders[0] / "source.pptx", folders[1] / "copy.pptx"
        self.bridge.setOption("source", str(source))
        self.assertEqual(Path(QUrl(self.bridge.snapshot["outputFolderUrl"]).toLocalFile()).resolve(), source.parent.resolve())
        self.assertEqual(Path(QUrl(self.bridge.snapshot["outputFileUrl"]).toLocalFile()).resolve(), source.with_name("source_with_plots.pptx").resolve())
        self.bridge.setOption("output", str(output))
        self.discover(folders[2])
        view = QQuickView()
        warnings = []
        view.engine().warnings.connect(lambda items: warnings.extend(item.toString() for item in items))
        view.setInitialProperties({"bridge": self.bridge, "colors": COLORS})
        view.setSource(QUrl.fromLocalFile(str(Path(__file__).resolve().parents[1] / "ui_qt/quick/qml/Presentation.qml")))
        self.assertIsNotNone(view.rootObject(), "\n".join(warnings + [error.toString() for error in view.errors()]))
        view.resize(1180, 760)
        view.show()
        before = {key: self.bridge.snapshot[key] for key in ("source", "output", "imageRoot")}
        try:
            for method, name, expected in (("chooseSource", "slides_source_dialog", folders[0]),
                    ("chooseOutput", "slides_output_dialog", folders[1]),
                    ("choosePlotFolder", "slides_plot_folder_dialog", folders[2])):
                with self.subTest(role=method):
                    self.assertTrue(QMetaObject.invokeMethod(view.rootObject(), method))
                    dialog = view.rootObject().findChild(QObject, name)
                    self.wait(lambda: dialog.property("visible"))
                    self.assertEqual(Path(dialog.property("currentFolder").toLocalFile()).resolve(), expected.resolve())
                    if method == "chooseOutput":
                        self.assertEqual(Path(dialog.property("selectedFile").toLocalFile()).resolve(), output.resolve())
                    self.assertTrue(QMetaObject.invokeMethod(dialog, "reject"))
                    self.wait(lambda: not dialog.property("visible"))
            self.assertEqual({key: self.bridge.snapshot[key] for key in before}, before)
            self.bridge.setOption("source", str(folders[3] / "another.pptx"))
            self.app.processEvents()
            self.assertTrue(QMetaObject.invokeMethod(view.rootObject(), "chooseSource"))
            dialog = view.rootObject().findChild(QObject, "slides_source_dialog")
            self.wait(lambda: dialog.property("visible"))
            self.assertEqual(Path(dialog.property("currentFolder").toLocalFile()).resolve(), folders[3].resolve())
            self.assertTrue(QMetaObject.invokeMethod(dialog, "reject"))
            self.assertEqual(warnings, [])
        finally:
            view.close()
            view.deleteLater()
            self.app.processEvents()

    def test_real_qml_selection_queue_preview_and_narrow_layout(self):
        self.png("PL/run/first_plot.png", 1)
        self.png("PL/run/second_plot.png", 2)
        self.discover()
        view = QQuickView()
        warnings = []
        view.engine().warnings.connect(lambda items: warnings.extend(item.toString() for item in items))
        view.setInitialProperties({"bridge": self.bridge, "colors": COLORS})
        view.setResizeMode(QQuickView.SizeRootObjectToView)
        view.setSource(QUrl.fromLocalFile(str(Path(__file__).resolve().parents[1] / "ui_qt/quick/qml/Presentation.qml")))
        self.assertIsNotNone(view.rootObject(), "\n".join(warnings + [error.toString() for error in view.errors()]))
        view.resize(1180, 760)
        view.show()
        QTest.qWait(40)
        def find(name, parent=None):
            parent = parent or view.rootObject()
            if parent.objectName() == name:
                return parent
            for child in parent.childItems():
                found = find(name, child)
                if found is not None:
                    return found
            return None
        def click(name):
            item = find(name)
            self.assertIsNotNone(item, name)
            QTest.mouseClick(view, Qt.LeftButton, Qt.NoModifier, item.mapToScene(item.boundingRect().center()).toPoint())
            QTest.qWait(20)
        try:
            click("slides_available_row_0")
            click("slides_add_selected")
            self.assertEqual(self.bridge.snapshot["queueCount"], 1)
            self.assertEqual(len(self.bridge.snapshot["preview"]["images"]), 1)
            self.wait(lambda: find("slides_preview_image_0") is not None and find("slides_preview_image_0").property("imageReady"))
            self.wait(lambda: find("slides_thumbnail_0") is not None and find("slides_thumbnail_0").property("imageReady"))
            click("slides_add_all")
            self.assertEqual(self.bridge.snapshot["queueCount"], 2)
            click("slides_queue_row_0")
            click("slides_down")
            self.assertEqual(self.bridge._queue[1], self.bridge.availableModel.rows[0]["path"])
            first_before_drag = self.bridge._queue[0]
            row = find("slides_queue_row_1")
            target = find("slides_queue_row_0")
            start = row.mapToScene(row.boundingRect().center()).toPoint()
            finish = target.mapToScene(target.boundingRect().center()).toPoint()
            QTest.mousePress(view, Qt.LeftButton, Qt.NoModifier, start)
            QTest.mouseMove(view, (start + finish) / 2, 15)
            QTest.mouseMove(view, finish, 15)
            QTest.mouseRelease(view, Qt.LeftButton, Qt.NoModifier, finish)
            QTest.qWait(20)
            self.assertNotEqual(self.bridge._queue[0], first_before_drag)
            click("slides_advanced")
            self.assertTrue(find("slides_output").isVisible())
            click("slides_remove")
            self.assertEqual(self.bridge.snapshot["queueCount"], 1)
            view.resize(640, 620)
            QTest.qWait(30)
            for name in ("slides_available_pane", "slides_queue_pane", "slides_preview_pane"):
                self.assertGreater(find(name).width(), 150)
                self.assertGreater(find(name).height(), 80)
            self.assertEqual(warnings, [])
        finally:
            view.close()
            view.deleteLater()
            self.app.processEvents()


if __name__ == "__main__":
    unittest.main()
