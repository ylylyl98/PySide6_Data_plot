"""Qt-free, snapshot based export for the unified MCD workflow.

The exporter deliberately accepts a small dictionary protocol.  Controllers can
capture their existing :class:`McdResult` and peak analysis objects with
``build_mcd_export_snapshot`` and queue the returned value on their worker.  All
file generation then uses a private copy of that value; no widget, figure, or
fit operation is consulted.
"""

from __future__ import annotations

from dataclasses import fields as dataclass_fields, is_dataclass
from datetime import datetime, timezone
import copy
import hashlib
import json
from pathlib import Path
import re
import shutil
import tempfile
from types import SimpleNamespace
from typing import Any, Mapping, Sequence

import numpy as np


SCHEMA_VERSION = 1
_BRANCHES = ("B increasing", "B decreasing")


def _freeze(value: Any) -> Any:
    """Deep-copy a snapshot value and make every numerical array read-only."""
    if isinstance(value, np.ndarray):
        result = np.array(value, copy=True)
        result.setflags(write=False)
        return result
    if isinstance(value, Mapping):
        return {str(key): _freeze(item) for key, item in value.items()}
    if isinstance(value, list):
        return tuple(_freeze(item) for item in value)
    if isinstance(value, tuple):
        return tuple(_freeze(item) for item in value)
    if is_dataclass(value):
        result = copy.deepcopy(value)
        for field in dataclass_fields(result):
            try:
                current = getattr(result, field.name)
                frozen = _freeze(current)
                try:
                    setattr(result, field.name, frozen)
                except (AttributeError, TypeError):
                    # Frozen result dataclasses still contain independent arrays
                    # after deepcopy; setting their flags is sufficient.
                    if isinstance(current, np.ndarray):
                        current.setflags(write=False)
            except (AttributeError, TypeError):
                continue
        return result
    if hasattr(value, "__dict__") and not isinstance(value, type) and vars(value):
        result = copy.deepcopy(value)
        for name, current in vars(result).items():
            try:
                setattr(result, name, _freeze(current))
            except (AttributeError, TypeError):
                if isinstance(current, np.ndarray):
                    current.setflags(write=False)
        return result
    if not isinstance(value, type):
        # Lightweight test doubles and a few legacy result wrappers expose
        # numerical arrays as class attributes.  A namespace captures values
        # at queue time instead of retaining those mutable class attributes.
        names = (
            "source_file", "energy_ev", "pair_b", "pair_labels",
            "pair_mcd_corrected", "pair_mcd_raw", "wavelength_nm",
        )
        captured = {name: _freeze(getattr(value, name)) for name in names if hasattr(value, name)}
        if captured:
            return SimpleNamespace(**captured)
    return copy.deepcopy(value)


def freeze_mcd_export_snapshot(snapshot: Mapping[str, Any]) -> dict[str, Any]:
    """Return an isolated snapshot suitable for a worker thread.

    This is public so a controller may freeze at queue time and use the same
    object for progress/cancellation bookkeeping.  The return value remains a
    normal dictionary for straightforward integration with existing controllers.
    """
    if not isinstance(snapshot, Mapping):
        raise TypeError("MCD export snapshot must be a mapping.")
    return _freeze(snapshot)


def build_mcd_export_snapshot(
    result: Any,
    *,
    analysis_results: Any = (),
    windows: Any = (),
    slopes: Any = (),
    links: Any = (),
    settings: Mapping[str, Any] | Any | None = None,
    source_descriptor: Mapping[str, Any] | None = None,
    provenance: Mapping[str, Any] | None = None,
    plot_state: Mapping[str, Any] | None = None,
    spectra: Any = (),
    features: Any = (),
    splitting: Any = (),
    maps: Any = None,
    selected_map: str | None = None,
) -> dict[str, Any]:
    """Capture existing numerical results in the explicit export protocol.

    ``windows`` may carry the UI's already reduced ``values`` per field, or a
    center/width/metric tuple that the worker can reduce without fitting.
    ``maps``/``selected_map`` optionally carry the selected ``DataCube`` for
    map DAT/PNG output; when omitted, result maps are captured automatically.
    """
    source_file = str(getattr(result, "source_file", ""))
    descriptor = dict(source_descriptor or {})
    if source_file and not descriptor.get("name"):
        descriptor["name"] = source_file
    if source_file and not descriptor.get("filename"):
        descriptor["filename"] = Path(source_file).name
    if source_file and not descriptor.get("path"):
        descriptor["path"] = source_file
    mcd = {
        "energy_ev": np.asarray(getattr(result, "energy_ev"), dtype=float),
        "fields_t": np.asarray(getattr(result, "pair_b"), dtype=float),
        "branches": np.asarray(getattr(result, "pair_labels"), dtype=str),
        "values": np.asarray(getattr(result, "pair_mcd_corrected"), dtype=float),
        "raw_values": np.asarray(getattr(result, "pair_mcd_raw"), dtype=float)
        if hasattr(result, "pair_mcd_raw") else None,
        "metric": "corrected signed mean",
        "channel": "MCD",
    }
    # McdResult retains spectra in acquisition wavelength order. Match the
    # canonical ascending-energy axis used by feature analysis when available.
    if hasattr(result, "wavelength_nm"):
        wavelength = np.asarray(getattr(result, "wavelength_nm"), dtype=float).ravel()
        if wavelength.size == mcd["energy_ev"].size and np.all(np.isfinite(wavelength)):
            order = np.argsort(1239.841984 / wavelength, kind="stable")
            # The McdResult coordinate is already the canonical physical
            # energy axis; source columns may still be wavelength ordered.
            # Sort the coordinate independently, then apply acquisition order
            # only to numerical columns.
            mcd["energy_ev"] = np.sort(mcd["energy_ev"], kind="stable")
            mcd["values"] = mcd["values"][:, order]
            if mcd["raw_values"] is not None:
                mcd["raw_values"] = mcd["raw_values"][:, order]
    payload: dict[str, Any] = {
        "source_file": source_file,
        "source_descriptor": descriptor,
        "acquisition_conditions": getattr(result, "acquisition_conditions", {}),
        "result": result,
        "mcd_result": result,
        "mcd": mcd,
        "analysis_results": analysis_results,
        "feature_results": analysis_results,
        "windows": windows,
        "slopes": slopes,
        "links": links,
        "settings": settings or {},
        "provenance": provenance or {},
        "plot_state": plot_state or {},
        "spectra": spectra,
        "features": features,
        "splitting": splitting,
        "maps": getattr(result, "maps", {}) if maps is None else maps,
        "selected_map": selected_map,
    }
    return freeze_mcd_export_snapshot(payload)


def _jsonable(value: Any) -> Any:
    if isinstance(value, np.ndarray):
        return [_jsonable(item) for item in value.tolist()]
    if isinstance(value, (np.integer, np.floating, np.bool_)):
        value = value.item()
    if isinstance(value, float):
        return value if np.isfinite(value) else None
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, Mapping):
        return {str(key): _jsonable(item) for key, item in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [_jsonable(item) for item in value]
    if is_dataclass(value):
        return {field.name: _jsonable(getattr(value, field.name)) for field in dataclass_fields(value)}
    if hasattr(value, "__dict__"):
        return {str(key): _jsonable(item) for key, item in vars(value).items()}
    return str(value)


def _safe_name(value: Any) -> str:
    text = re.sub(r"[^A-Za-z0-9._-]+", "_", str(value or "")).strip("._-")
    return text or "MCD"


def _source_prefix(snapshot: Mapping[str, Any]) -> str:
    descriptor = snapshot.get("source_descriptor", {})
    descriptor = descriptor if isinstance(descriptor, Mapping) else {}
    requested = snapshot.get("output_prefix") or descriptor.get("output_prefix")
    result = snapshot.get("result", snapshot.get("mcd_result"))
    source = requested or descriptor.get("filename") or descriptor.get("name") or snapshot.get("source_file") or getattr(result, "source_file", "MCD")
    prefix = Path(str(source or "MCD")).stem
    # The source stem is the strongest readable identity and normally already
    # contains sample/point/temperature/gate/field/repeat tokens.  Do not append
    # those fields a second time when a descriptor also exposes them.
    if prefix.casefold().endswith("_mcd"):
        prefix = prefix[:-4].rstrip("._-")
    prefix = _safe_name(prefix)
    if len(prefix) > 120:
        digest = hashlib.sha256(prefix.encode("utf-8")).hexdigest()[:8]
        prefix = f"{prefix[:110].rstrip('._-')}_{digest}"
    return prefix or "MCD"


def _prefix_for_root(prefix: str, root: Path) -> str:
    """Keep the package and revision leaf comfortably below Win32 limits."""
    # Reserve space for the energy label in both the directory and filename.
    # This budget must not depend on the window: all windows share one package.
    available = max(24, 128 - len(str(root)))
    if len(prefix) <= available:
        return prefix
    digest = hashlib.sha256(prefix.encode("utf-8")).hexdigest()[:8]
    return f"{prefix[:max(15, available - 9)].rstrip('._-')}_{digest}"


def _mcd_data(snapshot: Mapping[str, Any]) -> dict[str, Any]:
    source = snapshot.get("mcd")
    if not isinstance(source, Mapping):
        result = snapshot.get("result", snapshot.get("mcd_result"))
        source = {
            "energy_ev": getattr(result, "energy_ev", None),
            "fields_t": getattr(result, "pair_b", None),
            "branches": getattr(result, "pair_labels", None),
            "values": getattr(result, "pair_mcd_corrected", None),
            "raw_values": getattr(result, "pair_mcd_raw", None),
            "metric": "corrected signed mean",
            "channel": "MCD",
        }
    energy = np.asarray(source.get("energy_ev", ()), dtype=float).ravel()
    fields = np.asarray(source.get("fields_t", ()), dtype=float).ravel()
    branches = np.asarray(source.get("branches", ()), dtype=str).ravel()
    values = np.asarray(source.get("values", ()), dtype=float)
    if not energy.size or not fields.size or values.ndim != 2:
        raise ValueError("MCD snapshot needs non-empty energy, field, and two-dimensional values.")
    # The documented snapshot orientation is fields x energy.  Only a
    # non-square alternate shape can be identified safely and transposed.
    if values.shape != (fields.size, energy.size) and values.shape == (energy.size, fields.size):
        values = values.T
    if values.shape != (fields.size, energy.size) or branches.size != fields.size:
        raise ValueError("MCD snapshot arrays do not share a fields x energy grid.")
    raw = source.get("raw_values")
    raw_values = None if raw is None else np.asarray(raw, dtype=float)
    if raw_values is not None:
        if raw_values.shape != (fields.size, energy.size) and raw_values.shape == (energy.size, fields.size):
            raw_values = raw_values.T
        if raw_values.shape != values.shape:
            raise ValueError("MCD raw values do not match the corrected fields x energy grid.")
    return {"energy_ev": energy, "fields_t": fields, "branches": branches, "values": values, "raw_values": raw_values, "metric": str(source.get("metric", "corrected signed mean")), "channel": str(source.get("channel", "MCD"))}


def _record(value: Any) -> dict[str, Any]:
    if isinstance(value, Mapping):
        return dict(value)
    if is_dataclass(value):
        return {field.name: getattr(value, field.name) for field in dataclass_fields(value)}
    if hasattr(value, "__dict__"):
        return dict(vars(value))
    return {"value": value}


def _feature_rows(snapshot: Mapping[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    # ``features`` is the candidate inventory (representative energies/support)
    # and is metadata-only.  Numerical trajectories must be explicitly retained
    # as ``feature_points`` or come from a PeakShiftResult.
    explicit = snapshot.get("feature_points", ())
    if isinstance(explicit, Mapping):
        explicit = [explicit]
    for item in explicit or ():
        record = _record(item)
        rows.append(record)
    def analysis_items(value: Any, context: Mapping[str, Any] | None = None):
        context = dict(context or {})
        if isinstance(value, Mapping) and "tracks" not in value:
            for key, nested in value.items():
                next_context = dict(context)
                if str(key).casefold() in {"method", "tracking_method", "analysis_method"}:
                    next_context["analysis_method"] = str(key) if str(key) != "method" else str(nested)
                elif str(key).casefold() in {"channel", "source", "selected_channel", "pos", "neg", "k", "kp"}:
                    next_context["channel"] = str(key) if str(key) != "channel" else str(nested)
                else:
                    next_context.setdefault("analysis_id", str(key))
                yield from analysis_items(nested, next_context)
        elif isinstance(value, (list, tuple)):
            for nested in value:
                yield from analysis_items(nested, context)
        elif value is not None:
            yield value, context

    analyses = snapshot.get("analysis_results", snapshot.get("feature_results", ()))
    for analysis, analysis_context in analysis_items(analyses):
        tracks = getattr(analysis, "tracks", None)
        if tracks is None and isinstance(analysis, Mapping):
            tracks = analysis.get("tracks")
        for track_index, track in enumerate(tracks or (), start=1):
            track_record = _record(track)
            track_id = track_record.get("feature_id", f"{track_record.get('feature_kind', 'feature')}-{track_record.get('peak_id', track_index)}")
            analysis_source = getattr(analysis, "source", "") if not isinstance(analysis, Mapping) else analysis.get("source", "")
            analysis_method = getattr(analysis, "tracking_method", "") if not isinstance(analysis, Mapping) else analysis.get("tracking_method", analysis.get("method", ""))
            resolved_channel = track_record.get("channel", analysis_context.get("channel", analysis_source or track_record.get("source", "")))
            resolved_analysis_id = analysis_context.get("analysis_id", "")
            if analysis_source:
                resolved_analysis_id = f"{resolved_analysis_id}:{analysis_source}" if resolved_analysis_id else str(analysis_source)
            for point in track_record.get("points", ()):
                point_record = _record(point)
                rows.append({
                    "feature_id": str(track_id),
                    "feature_kind": track_record.get("feature_kind", "peak"),
                    "branch": point_record.get("branch", track_record.get("branch", "")),
                    "field_t": point_record.get("field_t"),
                    "energy_ev": point_record.get("energy_ev"),
                    "delta_energy_ev": point_record.get("delta_energy_ev"),
                    "status": point_record.get("status", track_record.get("quality", "")),
                    "reference_energy_ev": track_record.get("reference_energy_ev"),
                    "reference_field_t": track_record.get("reference_field_t"),
                    "reference_method": track_record.get("reference_method", track_record.get("reference_label", "")),
                    "analysis_method": track_record.get("analysis_method", analysis_context.get("analysis_method", analysis_method)),
                    "channel": resolved_channel,
                    "analysis_id": track_record.get("analysis_id", resolved_analysis_id),
                })
    return rows


def _valid_splitting_record(record: Mapping[str, Any]) -> bool:
    """Return whether a splitting row carries an explicit accepted pairing."""
    pair_id = record.get("pair_id") or record.get("mapping_id") or record.get("link_id")
    feature_ids = {
        str(record.get(name, ""))
        for name in ("selected_feature_id", "counterpart_feature_id", "selected_peak_id", "counterpart_peak_id")
    } - {""}
    explicit_pair = bool(pair_id) or len(feature_ids) >= 2
    status = str(record.get("status", record.get("state", "ok"))).casefold()
    accepted = status in {"", "ok", "complete", "completed", "ready", "valid", "tracked"}
    points = record.get("points", ())
    return explicit_pair and accepted and any(
        _record(point).get("field_t") is not None and _record(point).get("splitting_ev") is not None
        for point in (points or ())
    )


def _series_table(snapshot: Mapping[str, Any]) -> tuple[list[str], list[list[Any]]]:
    spectra = snapshot.get("spectra", {})
    items = list(spectra.items()) if isinstance(spectra, Mapping) else []
    columns: list[str] = []
    vectors: list[np.ndarray] = []
    for label, raw in items:
        record = _record(raw)
        energy = np.asarray(record.get("energy_ev", ()), float).ravel()
        fields = np.asarray(record.get("fields_t", ()), float).ravel()
        values = np.asarray(record.get("values", ()), float)
        if values.ndim != 2:
            continue
        if values.shape != (fields.size, energy.size) and values.shape == (energy.size, fields.size):
            values = values.T
        if values.shape != (fields.size, energy.size):
            continue
        for index, field in enumerate(fields):
            columns.extend([f"Energy_eV_{_safe_name(label)}_{_safe_name(field)}T", f"Spectrum_{_safe_name(label)}_{_safe_name(field)}T"])
            vectors.extend([energy, values[index]])
    if not vectors:
        return [], []
    rows = max(vector.size for vector in vectors)
    return columns, [[vector[row] if row < vector.size else None for vector in vectors] for row in range(rows)]


def _window_values(data: Mapping[str, Any], window: Mapping[str, Any]) -> np.ndarray:
    """Reduce one retained MCD window to one value per measured field."""
    center = window.get("center_ev")
    width = window.get("width_mev")
    # Validate coordinates even when a caller supplied pre-reduced values.
    # This prevents a stale center/width from silently being paired with a
    # trace computed from another energy window.
    if center is not None or width is not None:
        if center is None or width is None:
            raise ValueError("A retained MCD window needs both center_ev and width_mev.")
        if not np.isfinite(float(center)) or not np.isfinite(float(width)) or float(width) <= 0.0:
            raise ValueError("A retained MCD window must have finite center_ev and positive width_mev.")
        if float(center) < float(np.nanmin(data["energy_ev"])) or float(center) > float(np.nanmax(data["energy_ev"])):
            raise ValueError("A retained MCD window center is outside the captured energy range.")
    direct = window.get("values", window.get("trace_values"))
    if direct is not None:
        values = np.asarray(direct, dtype=float).ravel()
        if values.size == data["fields_t"].size:
            return values
        raise ValueError("A retained MCD window trace must have one value per captured field.")
    if center is None or width is None:
        raise ValueError("A retained MCD window needs center_ev and width_mev or one value per field.")
    half = max(float(width), 0.0) * 5e-4
    mask = np.abs(data["energy_ev"] - float(center)) <= half
    if not np.any(mask):
        mask[int(np.argmin(np.abs(data["energy_ev"] - float(center))))] = True
    selected = data["values"][:, mask]
    metric = str(window.get("metric", "mean")).casefold().replace(" ", "_")
    if metric in {"integral", "signed_integral"}:
        return np.trapezoid(selected, x=data["energy_ev"][mask], axis=1)
    if metric in {"absolute_mean", "unsigned_absolute_mean"}:
        return np.nanmean(np.abs(selected), axis=1)
    if metric in {"field_signed_absolute_mean", "signed_absolute_mean"}:
        return np.sign(data["fields_t"]) * np.nanmean(np.abs(selected), axis=1)
    return np.nanmean(selected, axis=1)


def _metric_unit(metric: Any) -> str:
    """Return the physical unit for one retained MCD metric."""
    normalized = str(metric or "mean").casefold().replace("−", "-")
    return "MCD·eV" if "integral" in normalized else "MCD"


def _retained_windows(snapshot: Mapping[str, Any], data: Mapping[str, Any]) -> list[dict[str, Any]]:
    """Return retained windows, deriving a validated current window if given."""
    source = snapshot.get("windows", ()) or ()
    if isinstance(source, Mapping):
        source = [source]
    retained = [_record(item) for item in source]
    for index, window in enumerate(retained, start=1):
        window.setdefault("window_id", window.get("id", f"window-{index}"))
    if retained:
        return retained
    provenance = snapshot.get("provenance", {})
    if isinstance(provenance, Mapping) and provenance.get("save_scope") == "retained":
        # An explicit feature-only batch must not silently acquire the live window.
        return []
    settings = snapshot.get("settings", {})
    settings = settings if isinstance(settings, Mapping) else {}
    center, width = settings.get("mcd_center_ev"), settings.get("mcd_width_mev")
    if center is not None or width is not None:
        return [{
            "center_ev": center,
            "width_mev": width,
            "metric": settings.get("metric", data["metric"]),
            "window_id": "current",
        }]
    # Legacy callers without an explicit window still get a reproducible
    # current trace.  New callers should always provide center/width.
    return [{"values": np.nanmean(data["values"], axis=1), "metric": data["metric"], "window_id": "current"}]


def _visible_filters(snapshot: Mapping[str, Any]) -> tuple[set[str] | None, set[str] | None]:
    state = snapshot.get("plot_state", {})
    state = state if isinstance(state, Mapping) else {}
    branches = state.get("visible_branches", state.get("branches", ()))
    if isinstance(branches, str):
        branches = [branches]
    allowed_branches = {str(value) for value in branches} if branches else None
    features = state.get("visible_features")
    if isinstance(features, str):
        features = [features]
    allowed_features = None if features is None else {str(value) for value in features}
    return allowed_branches, allowed_features


def _branch_filtered_data(data: Mapping[str, Any], allowed_branches: set[str] | None) -> dict[str, Any]:
    if allowed_branches is None:
        return dict(data)
    mask = np.asarray([str(branch) in allowed_branches for branch in data["branches"]], dtype=bool)
    if not np.any(mask):
        raise ValueError("Visible branch selection does not match any captured MCD rows.")
    return {
        **data,
        "fields_t": data["fields_t"][mask],
        "branches": data["branches"][mask],
        "values": data["values"][mask],
        "raw_values": None if data.get("raw_values") is None else data["raw_values"][mask],
    }


def _selected_map_data(snapshot: Mapping[str, Any], fallback: Mapping[str, Any]) -> dict[str, Any]:
    """Adapt a captured ``DataCube`` for the map/DAT product.

    ``McdResult.maps`` contains already binned map grids.  Reusing the chosen
    cube preserves its field axis and avoids rebuilding a branch mesh from the
    pair rows.  Plain snapshots may omit maps and continue to use the pair
    grid as an explicit fallback.
    """
    maps = snapshot.get("maps", {})
    if not isinstance(maps, Mapping) or not maps:
        return dict(fallback)
    state = snapshot.get("plot_state", {})
    state = state if isinstance(state, Mapping) else {}
    selected = snapshot.get("selected_map") or state.get("selected_map") or state.get("map_name") or "Combo"
    if not isinstance(selected, str):
        raise TypeError("selected_map must be a map name string; capture the selected DataCube in maps.")
    source = maps.get(selected)
    if source is None:
        source = maps.get("Corrected") or maps.get("Combo")
    if source is None:
        return dict(fallback)
    record = _record(source)
    energy = np.asarray(record.get("energy_ev", record.get("energy", ())), dtype=float).ravel()
    fields = np.asarray(record.get("fields_t", record.get("gate", ())), dtype=float).ravel()
    values = np.asarray(record.get("values", record.get("Z", ())), dtype=float)
    if not energy.size or not fields.size or values.ndim != 2:
        return dict(fallback)
    if values.shape != (fields.size, energy.size) and values.shape == (energy.size, fields.size):
        values = values.T
    if values.shape != (fields.size, energy.size):
        raise ValueError(f"Selected MCD map {selected!r} does not match its field x energy grid.")
    map_name = str(selected)
    branch_value = record.get("branch", map_name if map_name.casefold() in {"b increasing", "b decreasing"} else "")
    branches = np.asarray(record.get("branches", ()), dtype=str).ravel()
    if branches.size != fields.size:
        branches = np.full(fields.size, str(branch_value), dtype=object)
    return {
        "energy_ev": energy,
        "fields_t": fields,
        "branches": branches,
        "values": values,
        "raw_values": None,
        "metric": str(record.get("metric", fallback.get("metric", "corrected signed mean"))),
        "channel": str(record.get("channel", fallback.get("channel", "MCD"))),
        "map_name": map_name,
    }


def _feature_table(rows: Sequence[dict[str, Any]], *, value_key: str) -> tuple[list[str], list[list[Any]]]:
    groups: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        value = row.get(value_key)
        if row.get("field_t") is None or value is None:
            continue
        key = f"{row.get('feature_kind', 'feature')}_{row.get('feature_id', 'unknown')}_{row.get('channel', row.get('source', ''))}_{row.get('analysis_method', row.get('method', ''))}_{row.get('analysis_id', '')}_{row.get('branch', '')}"
        groups.setdefault(_safe_name(key), []).append(row)
    if not groups:
        return [], []
    columns: list[str] = []
    vectors: list[np.ndarray] = []
    for key, values in groups.items():
        values = sorted(values, key=lambda item: float(item.get("field_t")))
        columns.extend([f"B_{key}_T", f"{('Energy_eV' if value_key == 'energy_ev' else 'Shift_meV')}_{key}"])
        vectors.extend([np.asarray([float(item["field_t"]) for item in values]), np.asarray([float(item[value_key]) * (1000.0 if value_key == "delta_energy_ev" else 1.0) for item in values])])
    rows_out = max(vector.size for vector in vectors)
    return columns, [[vector[row] if row < vector.size else None for vector in vectors] for row in range(rows_out)]


def _slope_table(snapshot: Mapping[str, Any], *, allowed_branches: set[str] | None = None) -> tuple[list[str], list[list[Any]]]:
    source = snapshot.get("slopes", ())
    if isinstance(source, Mapping) and ("fits" in source or "differences" in source):
        flattened: list[dict[str, Any]] = []
        for category in ("fits", "differences"):
            entries = source.get(category, {})
            if isinstance(entries, Mapping):
                for key, value in entries.items():
                    record = _record(value)
                    record.setdefault("fit_id", str(key)); record.setdefault("window_id", source.get("window_id", "")); record.setdefault("category", category)
                    if category == "differences" and "difference" not in record:
                        record["difference"] = record.get("value", record.get("slope_difference"))
                    flattened.append(record)
            elif isinstance(entries, (list, tuple)):
                for value in entries:
                    record = _record(value); record.setdefault("category", category); flattened.append(record)
        source = flattened
    elif isinstance(source, Mapping):
        source = [source]
    records = [_record(item) for item in source or ()]
    windows = snapshot.get("windows", ()) or ()
    if isinstance(windows, Mapping):
        windows = [windows]
    window_records = [_record(item) for item in windows]
    window_by_id = {str(item.get("window_id")): item for item in window_records if item.get("window_id") is not None}
    default_window_id = next(iter(window_by_id), None) if len(window_by_id) == 1 else None
    for record in records:
        window_id = str(record.get("window_id", "") or "")
        if not window_id:
            # A lone fit can safely inherit the lone retained window.  With
            # several windows an untagged fit is explicitly marked instead of
            # being silently attributed to incompatible traces.
            window_id = default_window_id or ("unassigned" if len(window_by_id) > 1 else "")
            record["window_id"] = window_id
        window = window_by_id.get(window_id)
        if window is not None:
            record.setdefault("metric", window.get("metric"))
            record.setdefault("center_ev", window.get("center_ev"))
            record.setdefault("width_mev", window.get("width_mev"))
    if allowed_branches is not None:
        records = [record for record in records if not record.get("branch") or str(record.get("branch")) in allowed_branches or str(record.get("branch_a", "")) in allowed_branches or str(record.get("branch_b", "")) in allowed_branches]
    if not records:
        return [], []
    keys = [
        "window_id", "center_ev", "width_mev", "metric", "unit", "fit_id", "branch", "region", "method",
        "slope", "slope_se", "difference", "difference_se", "n", "requested_range_min_t", "requested_range_max_t",
        "actual_range_min_t", "actual_range_max_t", "covariance_supported", "residual_rms", "jump_flag", "curvature_flag", "status",
    ]
    columns = [
        "window_id", "center_eV", "width_meV", "metric", "unit", "fit_id", "branch", "region", "method",
        "slope_MCD_per_T", "slope_SE_MCD_per_T", "difference_MCD_per_T", "difference_SE_MCD_per_T", "n",
        "requested_range_min_T", "requested_range_max_T", "actual_range_min_T", "actual_range_max_T", "covariance_supported", "residual_rms", "jump_flag", "curvature_flag", "status",
    ]
    normalized: list[dict[str, Any]] = []
    for record in records:
        value = dict(record)
        value.setdefault("slope_se", value.get("standard_error"))
        value.setdefault("difference", value.get("slope_difference"))
        value.setdefault("difference_se", value.get("standard_error"))
        value.setdefault("branch", f"{value.get('branch_a', '')}->{value.get('branch_b', '')}" if value.get("branch_a") is not None else "")
        requested = value.get("requested_range_t")
        actual = value.get("actual_range_t")
        if isinstance(requested, (list, tuple)):
            value["requested_range_min_t"], value["requested_range_max_t"] = (list(requested) + [None, None])[:2]
        if isinstance(actual, (list, tuple)):
            value["actual_range_min_t"], value["actual_range_max_t"] = (list(actual) + [None, None])[:2]
        value.setdefault("actual_range_min_t", value.get("field_min_t"))
        value.setdefault("actual_range_max_t", value.get("field_max_t"))
        value.setdefault("unit", _metric_unit(value.get("metric", "mean")))
        normalized.append(value)
    return columns, [[record.get(key) for key in keys] for record in normalized]


def _splitting_table(snapshot: Mapping[str, Any], *, allowed_branches: set[str] | None = None, allowed_features: set[str] | None = None) -> tuple[list[str], list[list[Any]]]:
    source = snapshot.get("splitting", ())
    if isinstance(source, Mapping):
        source = [source]
    groups: dict[str, list[dict[str, Any]]] = {}
    for item in source or ():
        record = _record(item)
        if not _valid_splitting_record(record):
            continue
        if allowed_branches is not None and record.get("branch") and str(record.get("branch")) not in allowed_branches:
            continue
        if allowed_features is not None:
            identity_values = {str(record.get(name, "")) for name in ("pair_id", "selected_feature_id", "counterpart_feature_id", "selected_peak_id", "counterpart_peak_id")} - {""}
            # Preserve legacy splitting rows without feature identity; they
            # cannot be filtered meaningfully, while an explicit empty
            # selection still hides every splitting product.
            if not allowed_features or (identity_values and identity_values.isdisjoint(allowed_features)):
                continue
        points = record.get("points", ())
        for point in points or ():
            point_record = _record(point)
            if allowed_branches is not None and point_record.get("branch") and str(point_record.get("branch")) not in allowed_branches:
                continue
            if point_record.get("field_t") is None or point_record.get("splitting_ev") is None:
                continue
            identity = "_".join(str(record.get(name, point_record.get(name, ""))) for name in (
                "method", "pair_id", "selected_channel", "counterpart_channel",
                "selected_peak_id", "counterpart_peak_id", "branch",
            ))
            key = _safe_name(identity or point_record.get("branch", "branch"))
            groups.setdefault(key, []).append(point_record)
    if not groups:
        return [], []
    columns: list[str] = []
    vectors: list[np.ndarray] = []
    for key, points in groups.items():
        points = sorted(points, key=lambda item: float(item["field_t"]))
        columns.extend([f"B_{key}_T", f"Splitting_meV_{key}"])
        vectors.extend([np.asarray([float(item["field_t"]) for item in points]), np.asarray([float(item["splitting_ev"]) * 1000.0 for item in points])])
    rows = max(vector.size for vector in vectors)
    return columns, [[vector[row] if row < vector.size else None for vector in vectors] for row in range(rows)]


def _compact_slope_table(snapshot: Mapping[str, Any], *, allowed_branches: set[str] | None = None) -> tuple[list[Any], list[list[Any]]]:
    columns, values = _slope_table(snapshot, allowed_branches=allowed_branches)
    if not values:
        return [], []
    records = [dict(zip(columns, row)) for row in values]
    windows = snapshot.get("windows", ()) or ()
    if isinstance(windows, Mapping):
        windows = [windows]
    windows = [_record(window) for window in windows]
    if not windows:
        windows = [{"window_id": key} for key in dict.fromkeys(row["window_id"] for row in records)]
    descriptor = snapshot.get("source_descriptor", {}) or {}
    source = Path(str(descriptor.get("filename") or snapshot.get("source_file") or descriptor.get("name") or "Unavailable")).name
    rows: list[list[Any]] = []
    names = (("low", "Near zero"), ("high-", "Negative high field"), ("high+", "Positive high field"))
    errors = {"insufficient": "Insufficient points", "constant_field": "Constant field", "out_of_range": "No data in range"}
    for window in windows:
        key = str(window.get("window_id", window.get("id", "")))
        owned = [record for record in records if str(record["window_id"] or "") == key]
        if not owned and len(windows) == 1:
            owned = records
        metric = window.get("metric") or (owned[0].get("metric") if owned else None) or "mean"
        settings = dict(snapshot.get("settings", {}) or {})
        settings.update(window.get("settings", {}) or {})
        fits = {( _slope_region_name(record.get("region")), str(record.get("branch"))): record
                for record in owned if str(record.get("branch")) in _BRANCHES}
        def value(region: str, branch: str, *, error: bool = False) -> Any:
            if allowed_branches is not None and branch not in allowed_branches:
                return "Not selected"
            fit = fits.get((region, branch), {})
            slope = fit.get("slope_MCD_per_T")
            if slope is None or not np.isfinite(float(slope)):
                return errors.get(str(fit.get("status")), "Unavailable")
            number = fit.get("slope_SE_MCD_per_T") if error else slope
            return number if number is not None and np.isfinite(float(number)) else "Unavailable"
        rows.extend([
            ["Source", source], ["Energy (eV)", window.get("center_ev", owned[0].get("center_eV") if owned else None)],
            ["Width (meV)", window.get("width_mev", owned[0].get("width_meV") if owned else None)],
            ["Metric", metric], ["Slope unit", "eV/T" if "integral" in str(metric).casefold() else "1/T"], [],
            ["Region", "B range (T)", "Inc slope", "Inc SE", "Dec slope", "Dec SE"],
        ])
        for region, name in names:
            setting = {"low": "low", "high-": "high_negative", "high+": "high_positive"}[region]
            low, high = settings.get(f"slope_{setting}_t"), settings.get(f"slope_{setting}_end_t")
            if low is None or high is None:
                fit = next((record for (reg, _), record in fits.items() if reg == region), {})
                low, high = fit.get("requested_range_min_T"), fit.get("requested_range_max_T")
            bounds = f"{min(float(low), float(high)):.6g} to {max(float(low), float(high)):.6g}" if low is not None and high is not None else "Unavailable"
            rows.append([name, bounds, value(region, _BRANCHES[0]), value(region, _BRANCHES[0], error=True),
                         value(region, _BRANCHES[1]), value(region, _BRANCHES[1], error=True)])
        rows.extend([[], ["Comparison", "Inc", "Dec"]])
        for region, label in (("high-", "negative"), ("high+", "positive")):
            differences = []
            for branch in _BRANCHES:
                low, high = value("low", branch), value(region, branch)
                differences.append(low - high if not isinstance(low, str) and not isinstance(high, str) else (low if isinstance(low, str) else high))
            rows.append([f"Near-zero minus {label} high field", *differences])
        rows.extend([["Values = near-zero slope minus high-field slope."], ["SE = standard error. Full diagnostics and difference uncertainty are in JSON."], [], []])
    return rows[0], rows[1:]


def _write_xlsx(path: Path, tables: Sequence[tuple[str, list[str], list[list[Any]]]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, columns, rows in tables:
        if not columns or not rows:
            continue
        sheet = workbook.create_sheet(name)
        sheet.freeze_panes = "A2"
        sheet.append([str(column) for column in columns])
        for row in rows:
            sheet.append([None if isinstance(value, float) and not np.isfinite(value) else value for value in row])
        if name == "Slopes":
            sheet.freeze_panes = "C8"
            for index, width in enumerate((48, 25, 23, 17, 23, 17), start=1):
                sheet.column_dimensions[get_column_letter(index)].width = width
            for row in sheet.iter_rows():
                label = row[0].value
                for cell in row:
                    cell.font = Font(name="Calibri", size=11)
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = "0.000000"
                sheet.row_dimensions[row[0].row].height = 25
                if label in {"Region", "Comparison"}:
                    for cell in row:
                        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                        cell.fill = PatternFill("solid", fgColor="24476A")
                elif label in {"Source", "Metric", "Slope unit"}:
                    sheet.merge_cells(start_row=row[0].row, start_column=2, end_row=row[0].row, end_column=6)
                    row[0].font = Font(name="Calibri", size=11, bold=True)
                    if label == "Source":
                        sheet.row_dimensions[row[0].row].height = 36
                elif isinstance(label, str) and (label.startswith("Values =") or label.startswith("SE =")):
                    sheet.merge_cells(start_row=row[0].row, start_column=1, end_row=row[0].row, end_column=6)
                    row[0].font = Font(name="Calibri", size=10, color="555555")
            sheet.sheet_view.showGridLines = False
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_setup.orientation = "landscape"
            sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
            continue
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for index, column in enumerate(columns, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = min(38, max(12, len(str(column)) + 2))
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    if not workbook.sheetnames:
        raise ValueError("MCD export contains no plot-ready numerical table.")
    workbook.save(path)


def _write_dat(path: Path, data: Mapping[str, Any], export_id: str) -> None:
    energy = np.asarray(data["energy_ev"], float)
    fields = np.asarray(data["fields_t"], float)
    values = np.asarray(data["values"], float)
    with path.open("w", encoding="utf-8", newline="") as handle:
        handle.write(f"# workflow=MCD export_id={export_id} channel={data['channel']} metric={data['metric']}\n")
        # Every acquired row remains a DAT column, including equal B values
        # from the opposite round-trip branch.  The comment is machine readable
        # and preserves branch and measured-row identity while numeric headers
        # remain consumable by core.loader.load_dat.
        columns = [
            {"column": index + 1, "field_t": float(field), "branch": str(branch), "row": index}
            for index, (field, branch) in enumerate(zip(fields, data["branches"]))
        ]
        handle.write("# columns=" + json.dumps(columns, separators=(",", ":")) + "\n")
        handle.write("Energy_eV\t" + "\t".join(f"{field:.12g}" for field in fields) + "\n")
        for energy_index, energy_value in enumerate(energy):
            row = [f"{energy_value:.12g}"] + [f"{values[field_index, energy_index]:.12g}" for field_index in range(fields.size)]
            handle.write("\t".join(row) + "\n")


def _slope_region_name(value: Any) -> str:
    name = str(value or "").casefold().replace("−", "-")
    return {"high_positive": "high+", "high_negative": "high-"}.get(name, name)


def _plot_pngs(directory: Path, data: Mapping[str, Any], snapshot: Mapping[str, Any], export_id: str, *, map_data: Mapping[str, Any] | None = None) -> list[Path]:
    """Render the default MCD products from the frozen numerical snapshot."""
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.figure import Figure
    from matplotlib.ticker import MaxNLocator, ScalarFormatter
    from matplotlib.colors import Normalize

    FIGSIZE = (8.0, 6.0)
    DPI = 150
    # Keep the same fixed geometry for every product.  The lower top edge
    # leaves a measured 6–10 px gap below the compact detail line at 1200x900,
    # while the larger left margin keeps the 22 pt y labels inside the canvas.
    AXES_RECT = (0.21, 0.13, 0.74, 0.772)
    MAP_RECT = (0.21, 0.13, 0.59, 0.772)
    CBA_RECT = (0.83, 0.13, 0.025, 0.772)
    state = snapshot.get("plot_state", {})
    state = state if isinstance(state, Mapping) else {}
    energy = np.asarray(data["energy_ev"], float).ravel()
    fields = np.asarray(data["fields_t"], float).ravel()
    values = np.asarray(data["values"], float)
    branches = np.asarray(data["branches"], str).ravel()
    map_data = map_data or data
    map_energy = np.asarray(map_data["energy_ev"], float).ravel()
    map_fields = np.asarray(map_data["fields_t"], float).ravel()
    map_values = np.asarray(map_data["values"], float)
    map_branches = np.asarray(map_data["branches"], str).ravel()
    selected = state.get("visible_branches", state.get("branches", ()))
    if isinstance(selected, str): selected = [selected]
    selected_branches = {str(item) for item in selected} if selected else set(dict.fromkeys(branches.tolist()))
    paths: list[Path] = []

    def _source_stem() -> str:
        descriptor = snapshot.get("source_descriptor", {})
        if isinstance(descriptor, Mapping):
            source = descriptor.get("filename") or descriptor.get("name")
        else:
            source = None
        source = source or snapshot.get("source_file") or "MCD"
        return Path(str(source)).stem

    def _token(pattern: str) -> str:
        match = re.search(pattern, _source_stem(), flags=re.IGNORECASE)
        return match.group(1).replace("p", ".") if match else ""

    def _header(detail: str) -> tuple[str, str]:
        stem = _source_stem()
        # Gxx is an internal grouping token.  It remains in the source
        # descriptor and metadata, but is intentionally absent from visible
        # figure identity text.
        stem = re.sub(r"(?:^|_)G[^_]+(?=_|$)", "_", stem, flags=re.IGNORECASE)
        sample = stem.split("_", 1)[0]
        point_match = re.search(r"(?:^|_)(p\d+(?:n\d+)?)(?:_|$)", stem, re.IGNORECASE)
        point = point_match.group(1) if point_match else ""
        temp = _token(r"(?:^|_)(\d+(?:p\d+|\.\d+)?)K(?:_|$)")
        d_match = re.search(r"(?:^|_)(D[^_]+)", stem, re.IGNORECASE)
        f_match = re.search(r"(?:^|_)(F[^_]+)", stem, re.IGNORECASE)
        r_match = re.search(r"(?:^|_)(r\d+p\d+)(?:_|$)", stem, re.IGNORECASE)
        tokens = [sample]
        if point: tokens.append(point)
        if temp: tokens.append(f"{temp} K")
        if d_match or f_match:
            tokens.append("/".join(match.group(1).replace("p", ".") for match in (d_match, f_match) if match))
        if r_match: tokens.append(r_match.group(1).replace("p", "."))
        return " · ".join(tokens), detail

    def _figure(detail: str, *, map_axes: bool = False):
        fig = Figure(figsize=FIGSIZE, dpi=DPI)
        FigureCanvasAgg(fig)
        axis = fig.add_axes(MAP_RECT if map_axes else AXES_RECT)
        first, second = _header(detail)
        fig.text(0.21, 0.985, first, fontsize=16, va="top")
        fig.text(0.21, 0.948, second, fontsize=16, va="top", color="0.30")
        axis.tick_params(direction="in", top=True, right=True, labelsize=20)
        axis.xaxis.set_major_locator(MaxNLocator(5))
        axis.yaxis.set_major_locator(MaxNLocator(5))
        return fig, axis

    def _limits(values_in: Any, *, padding: float = 0.04, minimum: float = 0.05) -> tuple[float, float]:
        finite = np.asarray(values_in, float).ravel()
        finite = finite[np.isfinite(finite)]
        if not finite.size: return (-1.0, 1.0)
        low, high = float(np.min(finite)), float(np.max(finite))
        span = high - low
        pad = max(span * padding, minimum if span == 0 else span * padding)
        return low - pad, high + pad

    def _explicit(*names: str) -> tuple[float, float] | None:
        for name in names:
            value = state.get(name)
            if isinstance(value, (list, tuple)) and len(value) == 2:
                try:
                    low, high = float(value[0]), float(value[1])
                except (TypeError, ValueError):
                    continue
                if np.isfinite(low) and np.isfinite(high) and high > low: return low, high
        return None

    retained = _retained_windows(snapshot, data)
    multiple_windows = len(retained) > 1
    center = next((window.get("center_ev") for window in retained if window.get("center_ev") is not None), None)
    width = next((window.get("width_mev") for window in retained if window.get("width_mev") is not None), None)
    # The MCD map always uses the corrected captured map; the reflection
    # spectrum selector is intentionally independent of this product.
    map_detail = "Corrected MCD map"
    if not multiple_windows and center is not None and width is not None:
        map_detail = f"Corrected · E = {float(center):.6g} eV · ΔE = {float(width):.6g} meV"
    map_fig, map_axis = _figure(map_detail, map_axes=True)
    map_sort = np.argsort(map_fields, kind="stable")
    map_fields_sorted = map_fields[map_sort]
    map_values_sorted = map_values[map_sort]
    map_branches_sorted = map_branches[map_sort]
    finite_map = np.isfinite(map_values_sorted)
    displayed_map_values = []
    displayed_map_fields = []
    for branch in dict.fromkeys(str(item) for item in map_branches_sorted.tolist()):
        if branch and branch not in selected_branches: continue
        mask = map_branches_sorted == branch if branch else np.ones(map_fields_sorted.size, dtype=bool)
        finite_branch = np.isfinite(map_values_sorted[mask])
        displayed_map_values.extend(map_values_sorted[mask][finite_branch].tolist())
        displayed_map_fields.extend(map_fields_sorted[mask].tolist())
    explicit_vmin = state.get("map_vmin", state.get("vmin")); explicit_vmax = state.get("map_vmax", state.get("vmax"))
    try:
        map_vmin, map_vmax = float(explicit_vmin), float(explicit_vmax)
        if not np.isfinite([map_vmin, map_vmax]).all() or map_vmax <= map_vmin: raise ValueError
    except (TypeError, ValueError):
        map_vmin, map_vmax = _limits(displayed_map_values, padding=0.0, minimum=0.01)
    map_norm = Normalize(vmin=map_vmin, vmax=map_vmax)
    image = None
    for branch in dict.fromkeys(str(item) for item in map_branches_sorted.tolist()):
        if branch and branch not in selected_branches: continue
        mask = map_branches_sorted == branch if branch else np.ones(map_fields_sorted.size, dtype=bool)
        valid = mask[:, None] & finite_map
        if not np.any(valid): continue
        x_grid = np.broadcast_to(map_energy[None, :], map_values_sorted.shape)[valid]
        y_grid = np.broadcast_to(map_fields_sorted[:, None], map_values_sorted.shape)[valid]
        image = map_axis.scatter(x_grid, y_grid, c=map_values_sorted[valid], s=24, marker="s", linewidths=0, cmap="RdBu_r", norm=map_norm, label=branch or None)
    map_axis.set_xlabel("Energy (eV)", fontsize=22)
    map_axis.set_ylabel("Magnetic field B (T)", fontsize=22)
    map_axis.set_xlim(_explicit("map_xlim", "energy_xlim") or _limits(map_energy, padding=0.02, minimum=0.01))
    map_axis.set_ylim(_explicit("map_ylim", "field_ylim") or _limits(displayed_map_fields, padding=0.04, minimum=0.05))
    if not multiple_windows and center is not None and width is not None:
        half = float(width) * 5e-4
        map_axis.axvspan(float(center) - half, float(center) + half, ec="0.15", fc="none", lw=1.4)
    if image is not None:
        colorbar = map_fig.add_axes(CBA_RECT)
        cbar = map_fig.colorbar(image, cax=colorbar)
        cbar.ax.tick_params(labelsize=16)
        cbar.ax.set_title("MCD", fontsize=16, pad=10)
    map_path = directory / f"MCD_map_{export_id}.png"
    map_fig.savefig(map_path, dpi=DPI)
    paths.append(map_path)

    # Each retained window gets its own plot and fit table.  This keeps
    # different physical metrics off a shared axis and preserves fit ownership.
    slope_source = snapshot.get("slopes", ()) or ()
    if isinstance(slope_source, Mapping):
        fit_items = slope_source.get("fits", ()) if "fits" in slope_source else (slope_source,)
    else:
        fit_items = slope_source
    if isinstance(fit_items, Mapping):
        fit_items = list(fit_items.values())
    else:
        fit_items = list(fit_items or ())
    fit_colors = {"low": "#806015", "high-": "#bd3651", "high−": "#bd3651", "high+": "#7950a3"}
    for window_index, window in enumerate(retained, start=1):
        window_id = str(window["window_id"])
        metric = str(window.get("metric", data["metric"]))
        integral = "integral" in metric.casefold()
        metric_label = "" if metric.casefold().replace("_", " ") in {"mean", "signed mean", "corrected signed mean"} else metric
        detail = metric_label or "MCD"
        if window.get("center_ev") is not None and window.get("width_mev") is not None:
            detail = f"E = {float(window['center_ev']):.8g} eV · W = {float(window['width_mev']):.6g} meV"
            if metric_label:
                detail += f" · {metric_label}"
        mcd_fig, mcd_axis = _figure(detail)
        trace = _window_values(data, window)
        displayed_trace_values = []
        x_data = fields[np.isin(branches, list(selected_branches))]
        x_limits = _explicit("mcd_b_xlim", "b_xlim", "xlim") or _limits(x_data, padding=.04, minimum=.05)
        for branch in dict.fromkeys(branches.tolist()):
            if branch not in selected_branches:
                continue
            mask = branches == branch
            order = np.argsort(fields[mask], kind="stable")
            x, y = fields[mask][order], trace[mask][order]
            good = np.isfinite(x) & np.isfinite(y)
            if not np.any(good):
                continue
            increasing = "inc" in branch.casefold()
            label = "Inc" if increasing else "Dec" if "dec" in branch.casefold() else branch
            color = "#1967d2" if increasing else "#dc8500"
            mcd_axis.plot(x, y, color=color, ls="-" if increasing else "--", lw=1.1, alpha=.55, label=label)
            mcd_axis.scatter(x[good], y[good], s=12, marker="o" if increasing else "s",
                             facecolors=color if increasing else "none", edgecolors=color,
                             linewidths=.75, zorder=3)
            displayed_trace_values.append(y[good])
        owned_fits = [
            _record(item) for item in fit_items
            if str(_record(item).get("window_id") or (window_id if not multiple_windows else "")) == window_id
            and (not _record(item).get("branch") or str(_record(item)["branch"]) in selected_branches)
        ]
        for fit in owned_fits:
            actual = fit.get("actual_range_t")
            if not isinstance(actual, (list, tuple)) or len(actual) != 2:
                actual = (fit.get("actual_range_min_t", fit.get("field_min_t")), fit.get("actual_range_max_t", fit.get("field_max_t")))
            try:
                low, high = float(actual[0]), float(actual[1])
                slope, intercept = float(fit.get("slope")), float(fit.get("intercept"))
            except (TypeError, ValueError):
                continue
            if not np.isfinite([low, high, slope, intercept]).all() or high <= low:
                continue
            color = fit_colors.get(_slope_region_name(fit.get("region")), "#5f6368")
            linestyle = "--" if "dec" in str(fit.get("branch", "")).casefold() else "-"
            for left, right, alpha, linewidth in ((x_limits[0], low, .32, 2.), (low, high, 1., 3.), (high, x_limits[1], .32, 2.)):
                left, right = max(left, x_limits[0]), min(right, x_limits[1])
                if right <= left:
                    continue
                mcd_axis.plot([left, right], [slope * left + intercept, slope * right + intercept], color=color, linestyle=linestyle, alpha=alpha, lw=linewidth, zorder=5)
        y_data = np.concatenate(displayed_trace_values) if displayed_trace_values else []
        mcd_axis.set_xlabel("Magnetic field B (T)", fontsize=22)
        mcd_axis.set_ylabel("Integrated MCD (eV)" if integral else "MCD", fontsize=22)
        mcd_axis.set_xlim(x_limits)
        mcd_axis.set_ylim(_explicit("mcd_b_ylim", "b_ylim", "ylim") or _limits(y_data, padding=.05, minimum=.01))
        slope_unit = "eV T⁻¹" if integral else "T⁻¹"
        mcd_axis.text(.03, .965, f"Slope ({slope_unit})", transform=mcd_axis.transAxes, va="top", fontsize=16, color=".25")
        mcd_axis.text(.31, .965, "Inc", transform=mcd_axis.transAxes, va="top", fontsize=16, color=".25")
        mcd_axis.text(.48, .965, "Dec", transform=mcd_axis.transAxes, va="top", fontsize=16, color=".25")
        for row, region in enumerate(("Low", "High−", "High+")):
            y = .885 - .072 * row
            mcd_axis.text(.03, y, region, transform=mcd_axis.transAxes, va="top", fontsize=16, color=fit_colors.get(region.casefold(), ".25"))
            for xloc, branch_name in ((.31, "B increasing"), (.48, "B decreasing")):
                candidates = [fit for fit in owned_fits if str(fit.get("branch", "")) == branch_name and _slope_region_name(fit.get("region")) == _slope_region_name(region)]
                value = _record(candidates[0]).get("slope") if candidates else None
                mcd_axis.text(xloc, y, f"{float(value):.3f}" if value is not None else "N/A", transform=mcd_axis.transAxes, va="top", fontsize=16, color=".25")
        handles, labels = mcd_axis.get_legend_handles_labels()
        if handles:
            mcd_axis.legend(loc="lower right", frameon=False, fontsize=16)
        suffix = f"_w{window_index}" if multiple_windows else ""
        mcd_path = directory / f"MCD_vs_B{suffix}_{export_id}.png"
        mcd_fig.savefig(mcd_path, dpi=DPI)
        paths.append(mcd_path)

    # Feature products use only retained, completed tracks.  Every track keeps
    # its own B axis and NaN rows so missing measurements never get bridged.
    feature_rows = _feature_rows(snapshot)
    visible_features = state.get("visible_features")
    if isinstance(visible_features, str): visible_features = [visible_features]
    if visible_features is not None: feature_rows = [row for row in feature_rows if str(row.get("feature_id", "")) in {str(value) for value in visible_features}]
    feature_rows = [row for row in feature_rows if not row.get("branch") or str(row.get("branch")) in selected_branches]
    def identity(row: Mapping[str, Any]) -> tuple[str, ...]:
        return tuple(str(row.get(name, "")) for name in ("feature_id", "feature_kind", "channel", "analysis_method", "analysis_id", "branch"))
    accepted_statuses = {"", "tracked", "ok", "complete", "completed", "ready", "valid"}
    available_groups = {
        identity(row) for row in feature_rows
        if str(row.get("status", "tracked")).casefold() in accepted_statuses
        and any(row.get(metric) is not None for metric in ("energy_ev", "delta_energy_ev"))
    }
    # Keep missing rows belonging to an available track so matplotlib renders a
    # gap instead of connecting the measurements on either side.
    feature_rows = [row for row in feature_rows if identity(row) in available_groups]
    groups = list(dict.fromkeys(identity(row) for row in feature_rows))
    display_ids: dict[tuple[str, ...], str] = {}
    used_display_ids: set[str] = set()
    for key in groups:
        raw_id = str(key[0] or "track")
        candidate = _safe_name(raw_id if len(raw_id) <= 14 else raw_id[-10:])
        if candidate in used_display_ids:
            ordinal = 2
            while f"{candidate}~{ordinal}" in used_display_ids:
                ordinal += 1
            candidate = f"{candidate}~{ordinal}"
        display_ids[key] = candidate
        used_display_ids.add(candidate)
    def _channel_style(channel: str, branch: str) -> tuple[str, str, str]:
        lowered = channel.casefold(); color = "#1967d2" if "pos" in lowered or "+" in channel else "#dc8500" if "neg" in lowered or "−" in channel or "-" in channel else "#5f6368"
        inc = "inc" in branch.casefold(); return color, "-" if inc else "--", "o" if "pos" in lowered or "+" in channel else "s"
    def _reference_label(method: str) -> str:
        method = method.casefold()
        return "interp" if "interp" in method else "manual" if "manual" in method else "near" if "nearest" in method else "exact 0 T" if "exact" in method else "unknown"

    methods = {key[3] or "method?" for key in groups}
    references = {_reference_label(str(row.get("reference_method", ""))) for row in feature_rows}
    common_context = len(methods) == 1 and len(references) == 1
    distinct_channels = len({(key[2], key[5]) for key in groups}) == len(groups)
    def _feature_plot(metric: str, ylabel: str, filename: str) -> Path:
        detail = f"{len(groups)} retained tracks"
        if common_context:
            method = next(iter(methods))
            method_label = "raw" if "raw" in method.casefold() else "2nd deriv" if "second" in method.casefold() else method[:16]
            detail = f"{len(groups)} {method_label} tracks · E₀: {next(iter(references))}"
        fig, axis = _figure(detail)
        finite_x: list[float] = []; finite_y: list[float] = []
        for key in groups:
            rows = sorted([row for row in feature_rows if identity(row) == key], key=lambda row: float(row.get("field_t", 0.0)))
            x = np.asarray([row.get("field_t", np.nan) for row in rows], float)
            y = np.asarray([row.get(metric, np.nan) if str(row.get("status", "tracked")).casefold() in accepted_statuses else np.nan for row in rows], float)
            if metric == "delta_energy_ev": y *= 1000.0
            good = np.isfinite(x) & np.isfinite(y); finite_x.extend(x[good]); finite_y.extend(y[good])
            if not np.any(good): continue
            channel, branch = key[2], key[5]; color, linestyle, marker = _channel_style(channel, branch)
            method = key[3] or "method?"
            method_lower = method.casefold()
            method_label = "raw" if "raw" in method_lower else "2nd deriv" if "second" in method_lower else _safe_name(method)[:12]
            track_rows = rows
            reference_method = next((str(row.get("reference_method", "")).strip() for row in track_rows if str(row.get("reference_method", "")).strip()), "unknown")
            reference_label = _reference_label(reference_method)
            channel_label = "+" if "pos" in channel.casefold() else "−" if "neg" in channel.casefold() else (channel or "channel?")
            branch_label = "Inc" if "inc" in branch.casefold() else "Dec" if "dec" in branch.casefold() else (branch or "branch?")
            label = f"{channel_label} {branch_label} · {display_ids[key]} · {method_label} · E₀:{reference_label}"
            if common_context:
                label = f"{channel_label} {branch_label}"
                if not distinct_channels:
                    label += f" · {display_ids[key]}"
            axis.plot(x, y, color=color, linestyle=linestyle, marker=marker, ms=4, mfc=color if linestyle == "-" else "white", lw=1.45, label=label)
        axis.set_xlabel("Magnetic field B (T)", fontsize=22); axis.set_ylabel(ylabel, fontsize=22)
        axis.set_xlim(_explicit("feature_xlim", "b_xlim") or _limits(finite_x, padding=.04, minimum=.05))
        axis.set_ylim(_explicit("feature_ylim") or _limits(finite_y, padding=.05, minimum=.01))
        if finite_x:
            axis.legend(loc="upper center" if common_context else "best", frameon=False, fontsize=16,
                        ncol=2 if common_context else 1, columnspacing=.8, handlelength=1.5)
        path = directory / f"{filename}_{export_id}.png"; fig.savefig(path, dpi=DPI); return path
    energy_rows = [row for row in feature_rows if row.get("energy_ev") is not None and row.get("field_t") is not None]
    shift_rows = [row for row in feature_rows if row.get("delta_energy_ev") is not None and row.get("field_t") is not None]
    if energy_rows: paths.append(_feature_plot("energy_ev", "Energy (eV)", "feature_energy"))
    if shift_rows: paths.append(_feature_plot("delta_energy_ev", "Shift (meV)", "feature_shift"))

    # A splitting image is meaningful only for an explicitly accepted feature
    # pair.  Ambiguous/unvalidated candidates remain available in metadata but
    # do not become a visual product.
    splitting_groups: dict[str, list[dict[str, Any]]] = {}
    for item in (snapshot.get("splitting", ()) or ()):
        record = _record(item)
        if not _valid_splitting_record(record):
            continue
        record_branch = str(record.get("branch", ""))
        if record_branch and record_branch not in selected_branches:
            continue
        if visible_features is not None:
            identities = {str(record.get(name, "")) for name in (
                "pair_id", "selected_feature_id", "counterpart_feature_id", "selected_peak_id", "counterpart_peak_id",
            )} - {""}
            if not identities or not identities.intersection({str(value) for value in visible_features}):
                continue
        method_text = str(record.get("method", "") or "method?")
        method_lower = method_text.casefold()
        method_label = "raw" if "raw" in method_lower else "2nd deriv" if "second" in method_lower else _safe_name(method_text)[:12]
        pair_text = str(record.get("pair_id", record.get("mapping_id", "pair")) or "pair")
        pair_label = _safe_name(pair_text if len(pair_text) <= 14 else pair_text[-10:])
        selected_text = str(record.get("selected_channel", "") or "?")
        counterpart_text = str(record.get("counterpart_channel", "") or "?")
        channel_label = f"{selected_text}↔{counterpart_text}"
        branch_label = "Inc" if "inc" in record_branch.casefold() else "Dec" if "dec" in record_branch.casefold() else (record_branch or "branch?")
        group_key = f"{method_label} · {pair_label} · {channel_label} · {branch_label}"
        for point in record.get("points", ()) or ():
            point_record = _record(point)
            point_branch = str(point_record.get("branch", record_branch))
            if point_branch and point_branch not in selected_branches:
                continue
            if point_record.get("field_t") is None or point_record.get("splitting_ev") is None:
                continue
            splitting_groups.setdefault(group_key, []).append({
                "field_t": point_record.get("field_t"),
                "splitting_ev": point_record.get("splitting_ev"),
            })
    if splitting_groups:
        split_fig, split_axis = _figure("Valid paired splitting")
        split_x: list[float] = []; split_y: list[float] = []
        for group_key, points in splitting_groups.items():
            points = sorted(points, key=lambda point: float(point["field_t"]))
            x = np.asarray([float(point["field_t"]) for point in points], float)
            y = np.asarray([float(point["splitting_ev"]) * 1000.0 for point in points], float)
            good = np.isfinite(x) & np.isfinite(y)
            if not np.any(good):
                continue
            split_x.extend(x[good].tolist()); split_y.extend(y[good].tolist())
            split_axis.plot(x, y, marker="o", ms=4, lw=1.45, label=group_key)
        split_axis.set_xlabel("Magnetic field B (T)", fontsize=22)
        split_axis.set_ylabel("Splitting (meV)", fontsize=22)
        split_axis.set_xlim(_explicit("splitting_xlim", "b_xlim") or _limits(split_x, padding=.04, minimum=.05))
        split_axis.set_ylim(_explicit("splitting_ylim") or _limits(split_y, padding=.05, minimum=.01))
        if split_x:
            split_axis.legend(loc="upper left", frameon=False, fontsize=16, ncol=1, columnspacing=.8, handlelength=1.5)
        split_path = directory / f"splitting_{export_id}.png"
        split_fig.savefig(split_path, dpi=DPI)
        paths.append(split_path)

    # Reflection spectra remain available for callers that explicitly request
    # them, but are not part of the default MCD result package.
    if state.get("export_spectra"):
        spectra = snapshot.get("spectra", {}) or {}
        selected_channel = state.get("spectrum_channel", state.get("selected_channel"))
        selected_channel = None if selected_channel in (None, "") else str(selected_channel)
        selected_field = state.get("spectrum_field_t", state.get("selected_field_t"))
        try:
            selected_field = float(selected_field) if selected_field is not None else None
        except (TypeError, ValueError):
            selected_field = None
        selected_mode = state.get("spectrum_mode", state.get("selected_spectrum_mode", "corrected"))
        selected_mode = str(selected_mode or "corrected").casefold()
        spectrum_fig, spectrum_axis = _figure("Reflection spectra")
        for label, raw in spectra.items() if isinstance(spectra, Mapping) else ():
            record = _record(raw)
            record_channel = str(record.get("channel", record.get("source", label)))
            if selected_channel and str(label) != selected_channel and record_channel != selected_channel:
                continue
            e = np.asarray(record.get("energy_ev", ()), float).ravel()
            if "raw" in selected_mode:
                selected_values = record.get("raw_values", record.get("raw", record.get("values")))
            else:
                selected_values = record.get("corrected_values", record.get("corrected", record.get("values")))
            z = np.asarray(selected_values if selected_values is not None else (), float)
            sf = np.asarray(record.get("fields_t", ()), float).ravel()
            if z.ndim != 2:
                continue
            # The protocol is fields x energy.  Transpose only when the
            # rectangular shape proves the input is energy x fields; a square
            # grid is ambiguous and therefore remains in captured order.
            if z.shape != (sf.size, e.size) and z.shape == (e.size, sf.size):
                z = z.T
            if z.shape != (sf.size, e.size):
                continue
            indices = range(sf.size)
            if selected_field is not None and sf.size:
                index = int(np.nanargmin(np.abs(sf - selected_field)))
                indices = (index,)
            for index in indices:
                spectrum_axis.plot(e, z[index], label=f"{record_channel} · {selected_mode} · B={sf[index]:.6g} T")
        spectrum_axis.set_xlabel("Energy (eV)", fontsize=22); spectrum_axis.set_ylabel("Spectrum", fontsize=22)
        if spectrum_axis.lines:
            spectrum_axis.legend(loc="upper left", frameon=False, fontsize=16, ncol=1, columnspacing=.8, handlelength=1.5)
            paths.append(directory / f"spectra_{export_id}.png")
            spectrum_fig.savefig(paths[-1], dpi=DPI)
    return paths


def _progress(progress: Any, value: int) -> None:
    if progress is None:
        return
    callback = getattr(progress, "emit", progress if callable(progress) else None)
    if callback is not None:
        callback(int(value))


def _package_root(output_root: Path) -> Path:
    if output_root.name.casefold() == "mcd":
        return output_root
    if output_root.name.casefold() == "processed data":
        return output_root / "MCD"
    return output_root / "Processed Data" / "MCD"


def _window_export_label(windows: Sequence[Mapping[str, Any]]) -> str:
    if len(windows) == 1:
        center, width = windows[0].get("center_ev"), windows[0].get("width_mev")
        if center is not None and width is not None:
            # The UI accepts six decimals; keep the sixth when significant.
            center_text = f"{float(center):.6f}"
            if center_text.endswith("0"):
                center_text = center_text[:-1]
            return f"E{center_text}eV_W{float(width):.6g}meV"
    if len(windows) > 1:
        coordinates = sorted((float(item.get("center_ev", 0)), float(item.get("width_mev", 0))) for item in windows)
        digest = hashlib.sha256(json.dumps(coordinates).encode("utf-8")).hexdigest()[:8]
        return f"Windows{len(windows)}_{digest}"
    return "Analysis"


def _mcd_export_fingerprint(snapshot: Mapping[str, Any], windows: Sequence[Mapping[str, Any]]) -> str:
    identities = {str(item[key]): f"window-{index}" for index, item in enumerate(windows)
                  for key in ("id", "window_id") if item.get(key)}

    def normalize(value: Any) -> Any:
        if isinstance(value, np.ndarray):
            if value.dtype.hasobject:
                return {"array_shape": value.shape, "items": [normalize(item) for item in value.flat]}
            array = np.ascontiguousarray(value)
            return {"array_dtype": array.dtype.str, "array_shape": array.shape,
                    "array_digest": hashlib.sha256(array.tobytes()).hexdigest()}
        if isinstance(value, Mapping):
            return {str(key): normalize(item) for key, item in value.items() if key not in {"source_generation", "generation"}}
        if isinstance(value, (list, tuple, set)):
            return [normalize(item) for item in value]
        if is_dataclass(value):
            return normalize({field.name: getattr(value, field.name) for field in dataclass_fields(value)})
        if hasattr(value, "__dict__"):
            return normalize(vars(value))
        if isinstance(value, str):
            return identities.get(value, value)
        return _jsonable(value)

    content = {"render_version": 5, "snapshot": normalize(snapshot)}
    return hashlib.sha256(json.dumps(content, sort_keys=True).encode("utf-8")).hexdigest()


def _reuse_mcd_export(package: Path, fingerprint: str) -> dict[str, Any] | None:
    for metadata_path in package.glob("result_MCD_settings_*.json"):
        try:
            metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
            if metadata.get("content_fingerprint") != fingerprint:
                continue
            names = metadata["outputs"]
            if not names or any(Path(name).name != name or not (package / name).is_file() for name in names):
                continue
            export_id = metadata["export_id"]
            xlsx = str(package / next(name for name in names if name.endswith(".xlsx")))
            dat = str(package / next(name for name in names if name.endswith(".dat")))
            pngs = tuple(str(package / name) for name in names if name.endswith(".png"))
            paths = {"export_id": export_id, "revision": metadata["revision"], "reused": True,
                     "package_dir": str(package), "revision_dir": str(package),
                     "xlsx": xlsx, "xlsx_path": xlsx,
                     "metadata_json": str(metadata_path), "metadata_path": str(metadata_path), "json_path": str(metadata_path),
                     "mcd_map_dat": dat, "map_dat": dat, "plot_pngs": pngs, "png_paths": pngs}
            csv = next((name for name in names if name.endswith(".csv")), None)
            if csv:
                paths["trace_csv"] = str(package / csv)
            paths["output_paths"] = {key: value for key, value in paths.items() if key.endswith("_path") or key in {"xlsx", "metadata_json", "mcd_map_dat", "trace_csv"}}
            return paths
        except (OSError, ValueError, KeyError, TypeError, StopIteration):
            continue
    return None


def _saved_window_centers(windows: Sequence[Mapping[str, Any]]) -> tuple[float, ...]:
    try:
        values = [float(item['center_ev']) for item in windows]
        return tuple(sorted({round(value, 6) for value in values})) if all(np.isfinite(values)) else ()
    except (KeyError, TypeError, ValueError):
        return ()


def _mcd_save_source_identity(snapshot: Mapping[str, Any]) -> str:
    descriptor = snapshot.get('source_descriptor', snapshot.get('source', {})) or {}
    result = snapshot.get('result', snapshot.get('mcd_result'))
    identity = (descriptor.get('path') or snapshot.get('source_path')
                or descriptor.get('relative_path') or snapshot.get('source_relative_path')
                or descriptor.get('filename') or descriptor.get('name')
                or snapshot.get('source_file') or getattr(result, 'source_file', 'MCD'))
    return str(identity).replace('\\', '/').casefold()


def _saved_center_exports(package: Path, windows: Sequence[Mapping[str, Any]], source_identity: str) -> list[Path]:
    """Find committed products for this center set, independent of width/settings."""
    wanted = _saved_window_centers(windows)
    if not wanted:
        return []
    files = set()
    for metadata_path in package.glob('result_MCD_settings_*.json'):
        try:
            metadata = json.loads(metadata_path.read_text(encoding='utf-8'))
            if (metadata.get('workflow') != 'MCD' or _mcd_save_source_identity(metadata) != source_identity
                    or _saved_window_centers(metadata.get('windows', [])) != wanted):
                continue
            names = metadata['outputs']
            if not isinstance(names, list) or any(not isinstance(name, str) or Path(name).name != name for name in names):
                continue
            files.add(metadata_path)
            files.update(package / name for name in names if (package / name).is_file())
        except (OSError, ValueError, KeyError, TypeError):
            continue
    return sorted(files)


def export_mcd_analysis(snapshot: Mapping[str, Any], output_root: str | Path, *, progress: Any = None) -> dict[str, Any]:
    """Replace saved results for the same centers, publishing metadata last."""
    captured = freeze_mcd_export_snapshot(snapshot)
    def normalize_differences(value: Any) -> Any:
        if isinstance(value, Mapping):
            record = {key: normalize_differences(item) for key, item in value.items()}
            if record.get("comparison") == "high_minus_low":
                record["comparison"] = "low_minus_high"
                for key in ("slope_difference", "difference", "value"):
                    if isinstance(record.get(key), (float, int, np.number)):
                        record[key] = -record[key]
            return record
        if isinstance(value, (list, tuple)):
            return tuple(normalize_differences(item) for item in value)
        if type(value).__name__ in {"SlopeAnalysis", "SlopeDifference", "SlopeFit"} and hasattr(value, "to_dict"):
            return normalize_differences(value.to_dict())
        return value
    captured = normalize_differences(captured)
    data = _mcd_data(captured)
    allowed_branches, allowed_features = _visible_filters(captured)
    # Keep the captured pair grid intact.  Retained traces are aligned to this
    # original row axis; each product applies the visibility mask locally so a
    # direct trace is never shortened and then accidentally recomputed.
    map_data = _selected_map_data(captured, data)
    if allowed_branches is not None:
        map_branches = {str(branch) for branch in map_data["branches"]}
        matching = map_branches & allowed_branches
        if matching:
            map_data = _branch_filtered_data(map_data, allowed_branches)
        elif map_branches != {""}:
            raise ValueError("Visible branch selection does not match the selected MCD map.")
    retained_windows = _retained_windows(captured, data)
    window_values = [_window_values(data, window) for window in retained_windows]
    label = _window_export_label(retained_windows)
    mcd_root = _package_root(Path(output_root).expanduser().resolve())
    prefix = _prefix_for_root(_source_prefix(captured), mcd_root)
    package = mcd_root / f"{prefix}_MCD"
    source_identity = _mcd_save_source_identity(captured)
    for metadata_path in package.glob('result_MCD_settings_*.json'):
        try:
            existing = json.loads(metadata_path.read_text(encoding='utf-8'))
            if _mcd_save_source_identity(existing) == source_identity:
                continue
        except (OSError, ValueError, TypeError, AttributeError):
            pass
        source_tag = hashlib.sha256(source_identity.encode()).hexdigest()[:10]
        package = mcd_root / f'{prefix}_{source_tag}_MCD'
        break
    package.mkdir(parents=True, exist_ok=True)
    fingerprint = _mcd_export_fingerprint(captured, retained_windows)
    reused = _reuse_mcd_export(package, fingerprint)
    previous_files = _saved_center_exports(package, retained_windows, source_identity)
    if reused is not None and sum(path.suffix == '.json' for path in previous_files) <= 1:
        _progress(progress, 100)
        return reused
    revision = 1
    export_id = label
    center_key = _saved_window_centers(retained_windows)
    if center_key:
        lock_key = hashlib.sha256(json.dumps(center_key).encode()).hexdigest()[:16]
        reservation = package / f'.centers-{lock_key}.lock'
        try:
            with reservation.open('x'):
                pass
        except FileExistsError:
            raise RuntimeError('A save for this MCD center is already in progress.') from None
        previous_files = _saved_center_exports(package, retained_windows, source_identity)
    while not center_key:
        reservation = package / f".{export_id}.lock"
        if not (package / export_id).exists() and not any(package.glob(f"*_{export_id}.*")):
            try:
                with reservation.open("x"):
                    pass
                break
            except FileExistsError:
                pass
        revision += 1
        export_id = f"{label}_r{revision:02d}"
    final_dir = package
    staging = None
    published: list[Path] = []
    backups: list[tuple[Path, Path]] = []
    committed = False
    recovery_failed = False
    try:
        staging = Path(tempfile.mkdtemp(prefix=".staging-", dir=str(package)))
        _progress(progress, 5)
        workbook_traces: list[dict[str, Any]] = []
        settings_mapping = captured.get("settings", {}) if isinstance(captured.get("settings"), Mapping) else {}
        center = settings_mapping.get("mcd_center_ev")
        width = settings_mapping.get("mcd_width_mev")
        if center is None or width is None:
            fallback_windows = captured.get("windows", ()) or ()
            if isinstance(fallback_windows, Mapping):
                fallback_windows = [fallback_windows]
            for window in fallback_windows:
                window_record = _record(window)
                center = window_record.get("center_ev", center)
                width = window_record.get("width_mev", width)
                if center is not None and width is not None:
                    break
        center_token = f"_E{float(center):.6g}eV" if center is not None else ""
        width_token = f"_W{float(width):.6g}meV" if width is not None else ""
        mcd_columns: list[str] = []
        mcd_vectors: list[np.ndarray] = []
        branch_names = list(_BRANCHES)
        if allowed_branches is not None:
            branch_names = [branch for branch in branch_names if branch in allowed_branches]
        if not any(np.any(data["branches"] == branch) for branch in branch_names):
            branch_names = list(dict.fromkeys(str(item) for item in data["branches"]))
        for branch in branch_names:
            mask = data["branches"] == branch
            lowered = str(branch).casefold()
            suffix = "increasing" if branch == "B increasing" or "inc" in lowered else ("decreasing" if branch == "B decreasing" or "dec" in lowered else _safe_name(branch).casefold())
            if not np.any(mask):
                continue
            b_values = data["fields_t"][mask]
            for window_index, (window, reduced) in enumerate(zip(retained_windows, window_values), start=1):
                metric = str(window.get("metric", data["metric"]))
                metric_normalized = metric.casefold().replace(" ", "_")
                metric_unit = _metric_unit(metric)
                window_id = str(window.get("window_id", f"window-{window_index}"))
                window_center = window.get("center_ev", center)
                window_width = window.get("width_mev", width)
                center_label = f"_E{float(window_center):.6g}eV" if window_center is not None else ""
                width_label = f"_W{float(window_width):.6g}meV" if window_width is not None else ""
                window_label = f"_window-{_safe_name(window_id)}_w{window_index}"
                metric_label = _safe_name(data["metric"])
                if metric_normalized not in metric_label.casefold():
                    metric_label = f"{metric_label}_{_safe_name(metric)}"
                label = f"{data['channel']}_{metric_label}_{suffix}{window_label}_metric-{_safe_name(metric)}_unit-{_safe_name(metric_unit)}{center_label}{width_label}"
                y_values = reduced[mask]
                mcd_columns.extend([f"B_{suffix}_T{window_label}", label]); mcd_vectors.extend([b_values, y_values])
                columns = {"B_T": f"B_{suffix}_T{window_label}"}
                # Keep the Organizer's three physical metrics in the workbook
                # itself. Reuse the selected curve instead of duplicating it.
                metrics = {
                    "corrected_signed_mean": "mean",
                    "corrected_field_signed_absolute_mean": "field_signed_absolute_mean",
                    "corrected_integral": "integral",
                }
                aliases = {
                    "signed_mean": "mean", "corrected_signed_mean": "mean",
                    "signed_absolute_mean": "field_signed_absolute_mean",
                    "signed_integral": "integral",
                }
                selected_metric = aliases.get(metric_normalized, metric_normalized)
                for key, reduction in metrics.items():
                    if selected_metric == reduction:
                        columns[key] = label
                    elif window_center is not None and window_width is not None:
                        metric_column = f"{key}_{suffix}{window_label}"
                        metric_window = {"center_ev": window_center, "width_mev": window_width, "metric": reduction}
                        mcd_columns.append(metric_column)
                        mcd_vectors.append(_window_values(data, metric_window)[mask])
                        columns[key] = metric_column
                workbook_traces.append({
                    "window_id": window_id, "center_ev": window_center,
                    "width_mev": window_width, "branch": branch, "columns": columns,
                })
        rows = max((vector.size for vector in mcd_vectors), default=0)
        mcd_table = [[vector[row] if row < vector.size else None for vector in mcd_vectors] for row in range(rows)]
        feature_rows = _feature_rows(captured)
        if allowed_branches is not None:
            feature_rows = [row for row in feature_rows if not row.get("branch") or str(row.get("branch")) in allowed_branches]
        energy_columns, energy_rows = _feature_table(feature_rows, value_key="energy_ev")
        shift_columns, shift_rows = _feature_table(feature_rows, value_key="delta_energy_ev")
        # Feature checkboxes control rendered overlays.  The workbook remains
        # a complete numerical snapshot of all retained feature/splitting
        # products, while the branch visibility is applied uniformly to every
        # numerical product.
        split_columns, split_rows = _splitting_table(captured, allowed_branches=allowed_branches)
        slope_columns, slope_rows = _compact_slope_table(captured, allowed_branches=allowed_branches)
        xlsx = staging / f"MCD_unified_{export_id}.xlsx"
        _write_xlsx(xlsx, (
            ("MCD", mcd_columns, mcd_table), ("Slopes", slope_columns, slope_rows),
            ("Energy", energy_columns, energy_rows), ("Shift", shift_columns, shift_rows),
            ("Splitting", split_columns, split_rows),
        ))
        _progress(progress, 35)
        map_dat = staging / f"MCD_map_{export_id}.dat"; _write_dat(map_dat, map_data, export_id)
        pngs = _plot_pngs(staging, data, captured, export_id, map_data=map_data)
        _progress(progress, 70)
        descriptor = captured.get("source_descriptor", {})
        settings_value = captured.get("settings", {})
        source_value = str(captured.get("source_file") or (descriptor.get("filename") if isinstance(descriptor, Mapping) else ""))
        # History matching treats source_file as the portable legacy identity.
        # Keep the absolute location in source_path while avoiding an absolute
        # legacy name that conflicts with a relative source selection.
        source_name = Path(source_value).name
        source_path = str(descriptor.get("path", "")) if isinstance(descriptor, Mapping) else ""
        source_relative_path = str(descriptor.get("relative_path", descriptor.get("source_relative_path", ""))) if isinstance(descriptor, Mapping) else ""
        measurement_descriptor = {
            "role": "measurement", "filename": source_name, "source_file": source_name,
            "path": source_path, "source_path": source_path,
            "relative_path": source_relative_path, "source_relative_path": source_relative_path,
        }
        metadata = {
            "content_fingerprint": fingerprint,
            "schema_version": SCHEMA_VERSION, "workflow": "MCD", "dataset_type": "unified_mcd_analysis",
            "export_id": export_id, "revision": revision, "created_utc": datetime.now(timezone.utc).isoformat(),
            "source_file": source_name, "source_path": source_path, "source_relative_path": source_relative_path,
            "source": _jsonable(descriptor), "source_descriptor": _jsonable(descriptor),
            "sources": [_jsonable({**measurement_descriptor, **(dict(descriptor) if isinstance(descriptor, Mapping) else {})})],
            "provenance": _jsonable(captured.get("provenance", {})), "settings": _jsonable(settings_value),
            "acquisition_conditions": _jsonable(captured.get("acquisition_conditions", {})),
            "slopes": _jsonable(captured.get("slopes", ())),
            "links": _jsonable(captured.get("links", ())), "windows": _jsonable(retained_windows),
            "candidates": _jsonable(captured.get("candidates", captured.get("features", ()))), "feature_analysis": _jsonable(captured.get("analysis_results", ())),
            "retained_feature_points": _jsonable(captured.get("feature_points", ())),
            "fit_diagnostics": _jsonable(captured.get("fit_diagnostics", captured.get("diagnostics", {}))),
            "plot_state": _jsonable(captured.get("plot_state", {})),
            "mcd": {"channel": data["channel"], "metric": data["metric"], "energy_points": int(data["energy_ev"].size), "field_points": int(data["fields_t"].size)},
            "map": {"name": map_data.get("map_name", "pair grid"), "energy_points": int(map_data["energy_ev"].size), "field_points": int(map_data["fields_t"].size)},
            "mcd_b": {
                "center_ev": center, "width_mev": width, "primary_metric": str(settings_mapping.get("metric", "mean")),
                "fit_near_zero": bool(settings_mapping.get("fit_near_zero", False)),
                "low_field_mcd_slope_increasing_per_T": settings_mapping.get("low_field_mcd_slope_increasing_per_T"),
                "low_field_mcd_slope_decreasing_per_T": settings_mapping.get("low_field_mcd_slope_decreasing_per_T"),
            },
            "trace_workbook": {"filename": xlsx.name, "sheet": "MCD", "traces": workbook_traces},
            "outputs": [xlsx.name, map_dat.name, *[path.name for path in pngs]],
        }
        metadata_json = staging / f"result_MCD_settings_{export_id}.json"
        metadata_json.write_text(json.dumps(_jsonable(metadata), indent=2, sort_keys=True) + "\n", encoding="utf-8")
        _progress(progress, 90)
        # Stage everything before touching prior results. Keep rollback copies
        # until all products and the final history metadata have been published.
        staged_files = [path for path in staging.iterdir() if path != metadata_json] + [metadata_json]
        backup_dir = staging / 'previous'
        backup_dir.mkdir()
        for old_path in previous_files:
            backup = backup_dir / old_path.name
            old_path.replace(backup)
            backups.append((old_path, backup))
        for path in staged_files:
            destination = final_dir / path.name
            with destination.open("xb") as target:
                published.append(destination)
                with path.open("rb") as source:
                    shutil.copyfileobj(source, target)
        # Match legacy saves: make the new result visible to an already-built
        # Organizer catalog without requiring a directory rescan.
        try:
            from core.mcd_extract import index_processed_mcd_settings
            index_processed_mcd_settings(final_dir / metadata_json.name, payload=metadata)
        except (OSError, ValueError):
            pass
        _progress(progress, 100)
        paths = {
            "export_id": export_id, "revision": revision, "package_dir": str(package), "revision_dir": str(final_dir),
            "xlsx": str(final_dir / xlsx.name), "xlsx_path": str(final_dir / xlsx.name),
            "metadata_json": str(final_dir / metadata_json.name), "metadata_path": str(final_dir / metadata_json.name), "json_path": str(final_dir / metadata_json.name),
            "mcd_map_dat": str(final_dir / map_dat.name), "map_dat": str(final_dir / map_dat.name),
            "plot_pngs": tuple(str(final_dir / path.name) for path in pngs), "png_paths": tuple(str(final_dir / path.name) for path in pngs),
        }
        paths["output_paths"] = {key: value for key, value in paths.items() if key.endswith("_path") or key in {"xlsx", "metadata_json", "mcd_map_dat", "trace_csv"}}
        committed = True
        return paths
    except Exception as error:
        recovery_errors = []
        for path in reversed(published):
            try:
                path.unlink(missing_ok=True)
            except OSError as recovery_error:
                recovery_errors.append(str(recovery_error))
        for original, backup in reversed(backups):
            try:
                backup.replace(original)
            except OSError as recovery_error:
                recovery_errors.append(str(recovery_error))
        if recovery_errors:
            recovery_failed = True
            raise RuntimeError(f'MCD save failed; recovery was incomplete. Preserved recovery files: {staging}. '
                               + '; '.join(recovery_errors)) from error
        raise
    finally:
        if staging is not None and not recovery_failed and (committed or all(not backup.exists() for _, backup in backups)):
            shutil.rmtree(staging, ignore_errors=True)
        reservation.unlink(missing_ok=True)
