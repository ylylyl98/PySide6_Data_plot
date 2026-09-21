"""Independent DRR datasets, explicit loading recipes and scientific exports."""
from __future__ import annotations

from copy import deepcopy
from dataclasses import asdict, dataclass, replace
import hashlib
import json
from pathlib import Path
from uuid import uuid4

import numpy as np

from core.loader import DataCube
from core.drr_peak_analysis import PeakAnalysisSettings, analyze_drr_peaks


@dataclass
class AnalysisDataset:
    key: str
    name: str
    cube: DataCube
    provenance: dict
    settings: PeakAnalysisSettings
    result: dict | None = None
    revision: int = 0


def create_dataset(cube, name, provenance, settings=None):
    arrays = [np.array(value, dtype=float, copy=True) for value in (cube.energy, cube.gate, cube.Z)]
    x, y, z = arrays
    if x.ndim != 1 or y.ndim != 1 or z.shape != (y.size, x.size) or not x.size or not y.size:
        raise ValueError('Dataset requires nonempty energy/Y axes and a matching matrix.')
    if not np.all(np.isfinite(x)) or not np.all(np.isfinite(y)):
        raise ValueError('Dataset axes must be finite.')
    provenance = deepcopy(provenance)
    digest = hashlib.sha256(json.dumps(provenance, sort_keys=True, default=str).encode('utf-8'))
    for array in arrays:
        digest.update(str(array.shape).encode('ascii'))
        digest.update(array.tobytes())
        array.setflags(write=False)
    digest.update(str((cube.gate_label, cube.gate_unit, cube.y_axis_semantic)).encode('utf-8'))
    snapshot = replace(cube, energy=x, gate=y, Z=z)
    settings = settings or PeakAnalysisSettings(float(x.min()), float(x.max()), float(y.min()), float(y.max()))
    return AnalysisDataset(digest.hexdigest(), str(name), snapshot, provenance, replace(settings))


def apply_settings(dataset, settings):
    if dataset.settings == settings:
        return False
    dataset.settings = replace(settings)
    dataset.result = None
    dataset.revision += 1
    return True


def apply_common_settings(source, targets):
    count = 0
    for target in targets:
        if target is source:
            continue
        settings = replace(source.settings, seed_energy=target.settings.seed_energy, seed_y=target.settings.seed_y)
        count += int(apply_settings(target, settings))
    return count


def analyze_dataset(dataset):
    result = analyze_drr_peaks(dataset.cube, dataset.settings)
    result.update(dataset_id=dataset.key, dataset_name=dataset.name,
                  dataset_revision=dataset.revision, provenance=deepcopy(dataset.provenance))
    return result


def load_dataset(path, background_files, background_mode='external', background_which='all'):
    """Load one raw measurement using only the explicitly selected recipe."""
    from core.data_io import load_drr_external_cube, load_drr_self_cube
    path = Path(path).resolve()
    backgrounds = [str(Path(value).resolve()) for value in background_files]
    mode = str(background_mode).lower()
    which = str(background_which).lower()
    if mode == 'external':
        if not backgrounds:
            raise ValueError('External background mode requires selected background files.')
        if which not in ('all', 'first', 'last'):
            raise ValueError('Background frame method must be all, first or last.')
        cube = load_drr_external_cube(str(path.parent), [str(path)], backgrounds, baseline_which=which)
    elif mode in ('self_first', 'self_last'):
        if backgrounds:
            raise ValueError('Self background mode cannot also specify external background files.')
        which = 'first' if mode == 'self_first' else 'last'
        cube = load_drr_self_cube(str(path.parent), [str(path)], use_first_frame=which == 'first')
    else:
        raise ValueError('Background mode must be external, self_first or self_last.')
    provenance = {'measurement_files': [str(path)], 'background_files': backgrounds,
                  'background_mode': mode, 'background_which': which, 'y_axis': 'auto'}
    return create_dataset(cube, path.stem, provenance)


def _valid_result(dataset):
    from core.drr_peak_export import _validate
    result = dataset.result
    if result is None:
        raise ValueError(f'{dataset.name}: analyze the dataset before exporting.')
    if (result.get('dataset_id') != dataset.key or result.get('dataset_revision') != dataset.revision
            or result.get('settings') != asdict(dataset.settings)):
        raise ValueError(f'{dataset.name}: result is stale; analyze the dataset again.')
    _validate(result, dataset.cube)
    return result


def export_dataset(folder, dataset):
    from core.drr_peak_export import export_drr_peak_analysis
    return export_drr_peak_analysis(folder, dataset.name, _valid_result(dataset), dataset.cube)


def export_summary(folder, datasets):
    """Return a unique Origin-friendly XLSX with independent track namespaces."""
    from openpyxl import Workbook
    from core.drr_peak_export import _FIELDS, _METRIC_FIELDS, _append, _tracks
    datasets = list(datasets)
    if not datasets:
        raise ValueError('Select at least one analyzed dataset.')
    results = [_valid_result(dataset) for dataset in datasets]
    book = Workbook()
    book.remove(book.active)
    info = book.create_sheet('Peak_Info')
    _append(info, ['dataset_id', 'dataset_name', 'product', *_FIELDS, *_METRIC_FIELDS])
    params = book.create_sheet('Parameters')
    _append(params, ['dataset_id', 'dataset_name', 'parameter', 'value'])
    for index, (dataset, result) in enumerate(zip(datasets, results), start=1):
        for product, suffix in (('raw', 'DRR_Peaks'), ('second', 'D2E_Peaks')):
            if product not in result['products']:
                continue
            points = result['products'][product]['points']
            accepted = [point for point in points if point['status'] == 'accepted']
            tracks = _tracks(accepted)
            sheet = book.create_sheet(f'D{index:03d}_{suffix}')
            _append(sheet, [result['y_label']] + [f'Track {track} energy (eV)' for track in tracks])
            lookup = {(point['y'], point['track_id']): point['energy'] for point in accepted}
            for y in result['y_values']:
                _append(sheet, [y] + [lookup.get((y, track)) for track in tracks])
            _append(params, [dataset.key, dataset.name, f'{product}_sheet', sheet.title])
            for point in points:
                _append(info, [dataset.key, dataset.name, product] + [point[field] for field in _FIELDS] + [point.get(field) for field in _METRIC_FIELDS])
        values = {**asdict(dataset.settings), 'revision': dataset.revision, 'provenance': dataset.provenance,
                  'result_filter': result.get('result_filter'),
                  'y_label': result['y_label'], 'energy_unit': 'eV',
                  'wide_table_policy': 'Accepted energies only; missing and uncertain detections are blank.'}
        for name, value in values.items():
            if isinstance(value, (dict, list, tuple)):
                value = json.dumps(value, ensure_ascii=False, default=str)
            _append(params, [dataset.key, dataset.name, name, value])
    from openpyxl.styles import Font
    for sheet in book:
        sheet.freeze_panes = 'B2' if sheet.title.endswith('Peaks') else 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 23
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = '0.##########'
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f'DRR_analysis_summary_{uuid4().hex[:12]}.xlsx'
    try:
        with path.open('xb') as stream:
            book.save(stream)
    except Exception:
        path.unlink(missing_ok=True)
        raise
    finally:
        book.close()
    return path
