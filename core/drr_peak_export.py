"""Separate, lossless peak tables and comparison figures for DRR analysis."""
from __future__ import annotations

from dataclasses import replace
from datetime import datetime, timezone
import json
from pathlib import Path
import re
import tempfile
from uuid import uuid4

import numpy as np
from matplotlib.backends.backend_agg import FigureCanvasAgg
from matplotlib.figure import Figure
from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Font, PatternFill

from core.drr_peak_analysis import PeakAnalysisSettings, compute_second_derivative_row


_FIELDS = ("row_index", "y", "energy", "amplitude", "prominence", "polarity", "track_id", "status")
_METRIC_FIELDS = ('prominence_fraction','width_mev','confidence','noise_ratio','scale_support','neighbor_support','quality_reason','auto_energy','auto_prominence','auto_prominence_fraction','auto_width_mev')


def _validate(result, cube):
    """Reject malformed data rather than silently producing misleading tables."""
    try:
        json.dumps(result, allow_nan=False)
        settings = result["settings"]
        ys = np.asarray(result["y_values"], float)
        energy, gate, z = np.asarray(cube.energy), np.asarray(cube.gate), np.asarray(cube.Z)
        if result["schema_version"] != 1 or not isinstance(result["y_label"], str):
            raise ValueError("Unsupported peak analysis schema")
        if ys.ndim != 1 or not ys.size or not np.all(np.isfinite(ys)):
            raise ValueError("Selected Y values must be finite and nonempty")
        if len(set(ys)) != len(ys):
            raise ValueError("Selected Y values must be unique")
        if z.shape != (gate.size, energy.size) or energy.ndim != 1 or gate.ndim != 1:
            raise ValueError("Source cube dimensions do not match its axes")
        if not all(y in gate for y in ys):
            raise ValueError("Selected Y values are absent from the source")
        for field in ("x_min", "x_max", "y_min", "y_max", "sg_window", "sg_polyorder"):
            if not np.isfinite(float(settings[field])):
                raise ValueError("Invalid analysis settings")
        if settings["x_min"] >= settings["x_max"]:
            raise ValueError("Invalid energy range")
        products = result["products"]
        if not isinstance(products, dict) or not products or set(products) - {"raw", "second"}:
            raise ValueError("Unsupported analysis product")
        for product in products.values():
            seen = set()
            for point in product["points"]:
                if any(key not in point for key in _FIELDS):
                    raise ValueError("Incomplete peak point")
                row = point["row_index"]
                if not isinstance(row, int) or not 0 <= row < gate.size or gate[row] != point["y"] or point["y"] not in ys:
                    raise ValueError("Peak row does not match source Y coordinate")
                if not all(np.isfinite(float(point[k])) for k in ("y", "energy", "amplitude", "prominence")):
                    raise ValueError("Peak values must be finite")
                if point["status"] not in ("accepted", "uncertain") or point["polarity"] not in ("peak", "dip"):
                    raise ValueError("Invalid peak quality or polarity")
                if not isinstance(point["track_id"], (int, str)) or isinstance(point["track_id"], bool):
                    raise ValueError("Invalid track identifier")
                key = (row, point["track_id"])
                if key in seen:
                    raise ValueError("Multiple points assigned to one track in a row")
                seen.add(key)
    except (KeyError, TypeError, OverflowError) as exc:
        raise ValueError("Invalid peak analysis result") from exc


def _tracks(points):
    return sorted({p["track_id"] for p in points}, key=lambda value: (isinstance(value, str), value))


def _append(sheet, values):
    # Build only the new row. Looking it up via max_row/max_column scans
    # the entire worksheet and makes large peak tables quadratic to write.
    cells = []
    for value in values:
        cell = Cell(sheet, value=value)
        if isinstance(value, str):
            cell.data_type = "s"
        cells.append(cell)
    sheet.append(cells)


def _write_workbook(path, payload):
    book = Workbook()
    book.remove(book.active)
    try:
        for product, name in (("raw", "DRR_Peaks"), ("second", "D2E_Peaks")):
            sheet = book.create_sheet(name)
            points = payload["products"].get(product, {}).get("points", [])
            tracks = _tracks([p for p in points if p["status"] == "accepted"])
            _append(sheet, [payload["y_label"]] + [f"Track {track} energy (eV)" for track in tracks])
            lookup = {(p["y"], p["track_id"]): p["energy"] for p in points if p["status"] == "accepted"}
            for y in payload["y_values"]:
                _append(sheet, [y] + [lookup.get((y, track)) for track in tracks])
        info = book.create_sheet("Peak_Info")
        _append(info, ["product", *_FIELDS, *_METRIC_FIELDS])
        for product, data in payload["products"].items():
            for point in data["points"]:
                _append(info, [product] + [point[key] for key in _FIELDS] + [point.get(key) for key in _METRIC_FIELDS])
        params = book.create_sheet("Parameters")
        _append(params, ["Parameter", "Value"])
        for key, value in {**payload["settings"], **payload["metadata"]}.items():
            _append(params, [key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
        if "provenance" in payload:
            _append(params, ["provenance", json.dumps(payload["provenance"], ensure_ascii=False)])
        if 'result_filter' in payload:
            _append(params,['result_filter',json.dumps(payload['result_filter'],ensure_ascii=False)])
        for sheet in book:
            sheet.freeze_panes = "B2" if sheet.title.endswith("Peaks") else "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="244A64")
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = min(65, max(16, max(len(str(c.value or "")) for c in column) + 2))
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, float):
                        cell.number_format = "0.##########"
        book.save(path)
    finally:
        book.close()


def _write_comparison(path, result, cube):
    settings = result["settings"]
    fig = Figure(figsize=(12, 5.2), layout="constrained")
    FigureCanvasAgg(fig)
    try:
        for ax, product, title in zip(fig.subplots(1, 2), ("raw", "second"), ("DR/R", "Second energy derivative")):
            ax.set(title=title, xlabel="Energy (eV)", ylabel=result["y_label"])
            if product not in result["products"]:
                ax.text(.5, .5, "Not analyzed", ha="center", va="center", transform=ax.transAxes)
                continue
            source = cube
            if product == "second":
                # Calculate on the entire energy axis before restricting the plot.
                derivative_values = np.full(np.shape(cube.Z), np.nan)
                derivative_settings = PeakAnalysisSettings(**settings)
                for index in np.flatnonzero(np.isin(cube.gate, result["y_values"])):
                    derivative_values[index], _ = compute_second_derivative_row(cube, int(index), derivative_settings)
                source = replace(cube, Z=derivative_values, cbar_label="d2(DR/R)/dE2")
            xmask = (source.energy >= settings["x_min"]) & (source.energy <= settings["x_max"])
            ymask = np.isin(source.gate, result["y_values"])
            xs, ys = np.asarray(source.energy)[xmask], np.asarray(source.gate)[ymask]
            values = np.asarray(source.Z)[np.ix_(ymask, xmask)]
            xi, yi = np.argsort(xs), np.argsort(ys)
            mesh = ax.pcolormesh(xs[xi], ys[yi], values[np.ix_(yi, xi)], shading="nearest", cmap="RdBu_r", rasterized=True)
            fig.colorbar(mesh, ax=ax, label=source.cbar_label)
            points = result["products"][product]["points"]
            accepted_points = [p for p in points if p["status"] == "accepted"]
            tracks = {}
            colors = {"peak": "#18733E", "dip": "#7A238B"}
            for point in accepted_points:
                tracks.setdefault(point["track_id"], []).append(point)
            for track_points in tracks.values():
                if len(track_points) < 2:
                    continue
                accepted = {p["y"]: p["energy"] for p in track_points}
                # NaN breaks the line across missing/uncertain rows.
                ax.plot([accepted.get(y, np.nan) for y in ys[yi]], ys[yi], "-", lw=1.2,
                        color=colors[track_points[0]["polarity"]], label="_nolegend_")
            for polarity, marker in (("peak", "o"), ("dip", "v")):
                selected = [p for p in accepted_points if p["polarity"] == polarity]
                if selected:
                    ax.scatter([p["energy"] for p in selected], [p["y"] for p in selected],
                               marker=marker, color=colors[polarity], s=18,
                               edgecolors="white", linewidths=.35, label=f"Accepted {polarity}s")
            uncertain = [p for p in points if p["status"] == "uncertain"]
            if uncertain:
                ax.scatter([p["energy"] for p in uncertain], [p["y"] for p in uncertain], marker="x", color="black", label="Uncertain")
            ax.set_xlim(settings["x_min"], settings["x_max"])
            if points:
                ax.legend(fontsize=7, loc="best")
        fig.suptitle("Peak analysis — accepted tracks; crosses mark uncertainty")
        fig.savefig(path, dpi=160, format="png")
    finally:
        fig.clear()


def export_drr_peak_analysis(folder, base, result, raw_cube) -> dict[str, Path]:
    """Export a unique XLSX/JSON/PNG set, without touching original map exports.

    Wide sheets contain eV energies only for accepted detections. All selected
    Y rows and stable track columns remain present; Peak_Info retains uncertainty.
    """
    _validate(result, raw_cube)
    payload = json.loads(json.dumps(result, allow_nan=False))
    payload["metadata"] = {
        **payload.get("metadata", {}),
        "exported_at_utc": datetime.now(timezone.utc).isoformat(),
        "source_title": str(raw_cube.title),
        "source_shape": list(np.shape(raw_cube.Z)),
        "energy_unit": "eV",
        "wide_table_policy": "Accepted energies only; missing and uncertain detections are blank.",
        "derivative_processing": "Savitzky-Golay derivative on full energy axis before range selection",
    }
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r'[^\w. -]+', "_", str(base).replace("\\", "/").split("/")[-1]).strip(" .")[:100] or "DRR"
    stem = f"{safe}_{uuid4().hex[:12]}"
    outputs = {ext: folder / f"{stem}_peak_analysis.{ext}" for ext in ("xlsx", "json")}
    outputs["png"] = folder / f"{stem}_peaks_comparison.png"
    created = []
    try:
        with tempfile.TemporaryDirectory(prefix=".peak-analysis-", dir=folder) as staging:
            stage = {ext: Path(staging) / f"analysis.{ext}" for ext in outputs}
            _write_workbook(stage["xlsx"], payload)
            stage["json"].write_text(json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
            _write_comparison(stage["png"], payload, raw_cube)
            for ext, destination in outputs.items():
                # Exclusive creation protects even against a concurrent name collision.
                with destination.open("xb") as target:
                    created.append(destination)
                    target.write(stage[ext].read_bytes())
    except Exception:
        for path in created:
            path.unlink(missing_ok=True)
        raise
    return outputs
