"""Origin-friendly P2P exports; every dataset retains its own Y grid."""
import json
import numpy as np
from core.drr_comparison_groups import parse_condition
from core.drr_p2p_metrics import metric_values


def field_value(dataset):
    files=dataset.provenance.get('measurement_files',[])
    condition=parse_condition(files[0]) if files else None
    return condition['field'] if condition else None


def write_workbook(path,records):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    from openpyxl.utils import get_column_letter
    wb=Workbook();wb.remove(wb.active)
    ordered=sorted(records,key=lambda r:(field_value(r['dataset']) if field_value(r['dataset']) is not None else float('inf'),r['dataset'].name))
    for name,metric,unit in [('Amplitude','amplitude','DRR'),('Emax','maximum','eV'),('Emin','minimum','eV'),('Separation','separation','meV')]:
        ws=wb.create_sheet(name)
        for index,record in enumerate(ordered):
            field=field_value(record['dataset'])
            col=index*2+1;label=f'D{index+1}'+(f'_B{field:g}T' if field is not None else '')
            ws.cell(1,col,f'{label}_Y');ws.cell(1,col+1,f'{label}_{name}_{unit}')
            result=record['result']
            if result is not None:
                values=metric_values(result,metric)
                for row,(y,value) in enumerate(zip(result[:,0],values),2):
                    ws.cell(row,col,float(y));ws.cell(row,col+1,float(value) if np.isfinite(value) else None)
                    ws.cell(row,col).number_format='0.000000';ws.cell(row,col+1).number_format='0.000000'
            ws.column_dimensions[get_column_letter(col)].width=19;ws.column_dimensions[get_column_letter(col+1)].width=24
        ws.freeze_panes='A2'
    meta=wb.create_sheet('Metadata')
    meta.append(['ID','Dataset','B_T','Y_label','Energy_min_eV','Energy_max_eV','Y_min','Y_max','Smoothing_parameters','Status'])
    quality=wb.create_sheet('Quality');quality.append(['ID','Y','Edge_warning','Finite_samples','Status'])
    for index,record in enumerate(ordered):
        d=record['dataset'];label=f'D{index+1}'
        meta.append([label,d.name,field_value(d),d.cube.gate_label,*map(float,record['bounds']),json.dumps(record.get('smoothing',{'method':'Off'}),ensure_ascii=False),record.get('error') or 'ok'])
        if record['result'] is not None:
            for row,edge in zip(record['result'],record['edge']):
                quality.append([label,float(row[0]),bool(edge),int(row[6]),'ok' if np.isfinite(row[1]) else 'insufficient data'])
    for ws in wb:
        ws.freeze_panes='A2'
        for row in ws:
            for cell in row:
                if isinstance(cell.value,str):cell.data_type='s'
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='264A70')
    meta.column_dimensions['B'].width=65;meta.column_dimensions['I'].width=60
    wb.save(path);wb.close()
