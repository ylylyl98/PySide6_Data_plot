"""Origin-friendly SHG analysis and transposed spectra in one workbook.

Each measurement keeps its own angle and wavelength axes. The JSON sidecar is
the history/reproduction record; CSV copies are optional views of the same rows.
"""
from __future__ import annotations

from contextlib import ExitStack
import csv
from dataclasses import dataclass
from datetime import datetime, timezone
import json
import math
from pathlib import Path
from typing import Callable, Iterable

import numpy as np

from app_version import __version__
from core.export import (EXPORT_METADATA_SCHEMA_VERSION, SHG_RESULT_FIELDS,
                         _metadata_jsonable, _source_descriptor, iter_shg_result_rows)
from core.load_cancellation import check_load_cancelled
from core.shg_export_paths import validate_shg_output_paths
from core.shg_fit import evaluate_shg_angular_model
from core.shg_settings import ShgFitSettings


_MAX_ROWS = 1_048_576
_MAX_COLUMNS = 16_384
_FIT_POINTS = 721


def _value(value):
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def _json_value(value):
    value = _metadata_jsonable(value)
    if isinstance(value, dict):
        return {key: _json_value(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_value(item) for item in value]
    return _value(value)


def _text(value: str) -> str:
    # Spreadsheet cells cannot contain XML 1.0 control characters. Keep the
    # source identity readable and reversible without creating formula cells.
    return "".join(f"\\u{ord(char):04x}" if ord(char) < 32 and char not in "\t\n\r" else char
                   for char in value)


def _order(result):
    angle = np.asarray(result.measured_angle_deg, float)
    return np.argsort(np.where(np.isfinite(angle), angle, np.inf), kind="stable")


def _source(folder, name, role):
    path = Path(name)
    path = (Path(folder) / path).resolve() if not path.is_absolute() else path.resolve()
    descriptor = _source_descriptor(str(folder), str(path), role=role)
    descriptor["name"] = path.name
    root = Path(folder).resolve()
    if path.is_relative_to(root):
        descriptor["source_relative_path"] = path.relative_to(root).as_posix()
    return descriptor


@dataclass
class _Table:
    name: str
    headers: list[str]
    row_count: int
    rows: Callable[[], Iterable]
    metadata: dict


def _column(header, role, units, dataset=None):
    return dict(header=header, designation=role, units=units,
                **({"dataset": dataset} if dataset else {}))


def _overview_table(folder, results, roles, inputs, created_utc, tables):
    """Keep source identities readable even when the workbook travels alone."""
    compare = len(results) == 2
    sources = {item["role"]: item for item in inputs}
    labels = {"single": "Measurement", "reference": "Reference A", "sample": "Sample B"}
    rows = [["Analysis", "Mode", "Compare / Twist angle" if compare else "Single file",
             "Reference A is compared with Sample B." if compare else "One measured angular sweep."]]
    for role in roles:
        source = sources["measurement" if role == "single" else role]
        rows.append(["Source files", labels[role], source["name"], source["source_path"]])
    for role, result in zip(roles, results):
        prefix = role.title() if compare else "Spectra"
        rows.append(["Data mapping", labels[role], f"{prefix}_Raw / {prefix}_Corrected",
                     f"{role.title()}_* columns in Plot / Fit; dataset = {role} in Summary / Diagnostics."
                     if compare else "Plot / Fit / Summary / Diagnostics describe this measurement."])
        rows.append(["Processing", labels[role] + " background", result.settings.background_method,
                     "Background model used to calculate corrected spectra and integrated intensities."])
        background = sources.get("background" if role == "single" else role + "_background")
        if background is not None:
            rows.append(["Background files", labels[role], background["name"], background["source_path"]])
        rows.append(["Processing", labels[role] + " peak (nm)", result.settings.peak_center_nm,
                     f"Integration window: {result.settings.gate_min_nm:g} to {result.settings.gate_max_nm:g} nm."])
    rows.extend([
        ["Analysis", "Experiment folder", str(Path(folder).resolve()), "Location at the time of export."],
        ["Analysis", "Exported (UTC)", created_utc, "Same export timestamp as analysis_settings.json."],
        ["Worksheet guide", "Overview", "Source files and workbook contents", "Start here to identify the measurement or reference/sample pair."],
    ])
    descriptions = {
        "Plot": "Angle, integrated intensity and error; each dataset has its own X axis.",
        "Fit": "Smooth fitted curves; blank when fitting is unavailable or disabled.",
        "Summary": "Source names, fit parameters, uncertainties and twist-angle results.",
        "Diagnostics": "All acquisition rows, source-row identities, quality flags and residuals.",
    }
    for table in tables:
        description = descriptions.get(table.name, table.metadata.get("description", ""))
        role = table.metadata.get("dataset")
        detail = (labels[role] + ": wavelength (nm) in the first column; remaining columns are intensity (counts)."
                  if role else "See the first row for column names and units.")
        rows.append(["Worksheet guide", table.name, description, detail])
    return _Table("Overview", ["Section", "Role / sheet", "File / value", "Full path / description"],
                  len(rows), lambda: iter(rows),
                  dict(description="Source roles, complete paths, processing context and worksheet guide."))


def _plot_table(results, roles):
    columns, blocks = [], []
    for role, result in zip(roles, results):
        angle = np.asarray(result.measured_angle_deg, float)
        area = np.asarray(result.integrated_area, float)
        indices = _order(result)
        indices = indices[np.asarray(result.included, bool)[indices]
                          & np.isfinite(angle[indices]) & np.isfinite(area[indices])]
        blocks.append((angle[indices], area[indices], np.asarray(result.area_uncertainty)[indices]))
        prefix = "" if len(results) == 1 else role.title() + "_"
        columns.extend((_column(prefix + "Angle_deg", "X", "deg", role),
                        _column(prefix + "Intensity_counts_nm", "Y", "counts*nm", role),
                        _column(prefix + "Error_counts_nm", "Y Error", "counts*nm", role)))
    count = max((len(block[0]) for block in blocks), default=0)

    def rows():
        for index in range(count):
            yield [values[index] if index < len(values) else None for block in blocks for values in block]

    return _Table("Plot", [column["header"] for column in columns], count, rows,
                  dict(columns=columns, inclusion="Finite, processing-included points; independent angle axes."))


def _fit_table(results, roles, fits):
    columns, blocks = [], []
    for role, result, fit in zip(roles, results, fits):
        if fit is None:
            continue
        angle = np.asarray(result.measured_angle_deg, float)
        selected = angle[np.asarray(fit.fit_mask, bool) & np.isfinite(angle)]
        if not selected.size:
            raise ValueError(f"The {role} SHG fit has no finite included angles.")
        x = np.linspace(float(selected.min()), float(selected.max()), _FIT_POINTS)
        y = evaluate_shg_angular_model(x, fit.i0, fit.amplitude, fit.x_center_deg)
        blocks.append((x, y))
        prefix = "" if len(results) == 1 else role.title() + "_"
        columns.extend((_column(prefix + "Fit_angle_deg", "X", "deg", role),
                        _column(prefix + "Fit_intensity_counts_nm", "Y", "counts*nm", role)))
    count = _FIT_POINTS if blocks else 0
    if not columns:
        columns = [_column("Fit_angle_deg", "X", "deg"),
                   _column("Fit_intensity_counts_nm", "Y", "counts*nm")]
    return _Table("Fit", [column["header"] for column in columns], count,
                  lambda: ([values[index] for block in blocks for values in block] for index in range(count)),
                  dict(columns=columns, available=bool(blocks),
                       sampling="721 model points within each fit's actual included angle range; no extrapolation."))


def _summary_table(results, roles, fits, twist, fit_settings, fit_message):
    rows = []
    parameters = (("i0", "i0_uncertainty", "counts*nm"),
                  ("amplitude", "amplitude_uncertainty", "counts*nm"),
                  ("x_center_deg", "x_center_uncertainty_deg", "deg"),
                  ("r_squared", None, ""), ("rmse", None, "counts*nm"),
                  ("reduced_chi_squared", None, ""), ("point_count", None, "points"),
                  ("angle_span_deg", None, "deg"))
    for role, result, fit in zip(roles, results, fits):
        status = "success" if fit is not None else "disabled" if not fit_settings.get("enabled", True) else "unavailable"
        rows.append([role, "fit_status", status, None, "", "" if fit is not None else fit_message])
        rows.append([role, "source_csv", result.data.source_file, None, "", ""])
        rows.append([role, "processing_included_rows", int(np.count_nonzero(result.included)), None, "points", ""])
        for name, error_name, units in parameters:
            rows.append([role, name, getattr(fit, name) if fit is not None else None,
                         getattr(fit, error_name) if fit is not None and error_name else None, units, status])
    if len(results) == 2:
        rows.append(["comparison", "twist_status", "success" if twist is not None else "unavailable",
                     None, "", "Both individual fits are required." if twist is None else ""])
        for name, error_name, units in (
            ("nearest_delta_x_center_deg", None, "deg"), ("phase_branch", None, ""),
            ("delta_x_center_deg", "delta_x_center_uncertainty_deg", "deg"),
            ("signed_twist_angle_deg", "twist_uncertainty_deg", "deg"),
            ("absolute_twist_angle_deg", "twist_uncertainty_deg", "deg"),
        ):
            rows.append(["comparison", name, getattr(twist, name) if twist is not None else None,
                         getattr(twist, error_name) if twist is not None and error_name else None,
                         units, "success" if twist is not None else "unavailable"])
    rows.append(["workbook", "Origin_import", "Use the first row as column names; set Plot columns to repeating X/Y/Y Error.",
                 None, "", "Fit uses independent X/Y pairs. Spectra use first column X, all remaining columns Y."])
    return _Table("Summary", ["dataset", "parameter", "value", "uncertainty", "units", "status_or_note"],
                  len(rows), lambda: iter(rows), dict(description="Fit statistics and twist results; unavailable values remain blank."))


def _spectra_tables(result, role, compare, include_cleaned):
    from openpyxl.utils import get_column_letter
    wave = np.asarray(result.data.wavelength_nm, float)
    order = _order(result)
    prefix = role.title() if compare else "Spectra"
    headers, mapping = ["Wavelength_nm"], []
    for offset, index in enumerate(order, start=2):
        angle = float(result.measured_angle_deg[index])
        source_row = int(result.data.source_rows[index])
        label = f"{angle:.12g} deg" if math.isfinite(angle) else "Unknown angle"
        header = f"{label} [row {source_row}]"
        headers.append(header)
        mapping.append(dict(column=get_column_letter(offset), header=header, source_index=int(index),
                            source_row=source_row, measured_angle_deg=_value(angle),
                            raw_measured_angle_deg=_value(result.data.measured_angle_deg[index]),
                            target_angle_deg=_value(result.data.target_angle_deg[index]),
                            included=bool(result.included[index]), quality_flag=result.quality_flags[index]))
    full = np.arange(wave.size)
    cropped = (np.flatnonzero((wave >= result.settings.left_min_nm) & (wave <= result.settings.right_max_nm))
               if result.settings.background_method != "none" else full)
    arrays = [("Raw", result.data.spectra, full, "Loaded measurement counts; full wavelength range."),
              ("Corrected", result.corrected, cropped,
               "Cosmic-cleaned counts minus background; restricted to the background analysis window when background subtraction is enabled.")]
    if include_cleaned:
        arrays.append(("Cleaned", result.cleaned_spectra, full, "Cosmic-cleaned counts before background subtraction; full wavelength range."))
    tables = []
    for kind, matrix, wave_indices, description in arrays:
        def rows(matrix=matrix, wave_indices=wave_indices):
            for wavelength_index in wave_indices:
                yield [wave[wavelength_index], *np.asarray(matrix)[order, wavelength_index]]
        tables.append(_Table(prefix + "_" + kind, headers, len(wave_indices), rows,
            dict(dataset=role, spectrum_type=kind.casefold(), description=description,
                 wavelength_column="A", wavelength_units="nm", intensity_units="counts",
                 angle_units="deg", columns=mapping,
                 wavelength_range_nm=([float(wave[wave_indices].min()), float(wave[wave_indices].max())]
                                      if len(wave_indices) else []),
                 background_analysis_window_nm=[result.settings.left_min_nm, result.settings.right_max_nm],
                 layout="One wavelength per row, one acquisition per column; repeated angles are never averaged.")))
    return tables, mapping


def _diagnostics_table(results, roles, fits, mappings):
    def rows():
        for role, result, fit, mapping in zip(roles, results, fits, mappings):
            for row, spectrum in zip(iter_shg_result_rows(result.data, result, result.settings, fit), mapping):
                yield [role, spectrum["column"], spectrum["header"], *[row[key] for key in SHG_RESULT_FIELDS]]
    return _Table("Diagnostics", ["dataset", "spectrum_column", "spectrum_header", *SHG_RESULT_FIELDS],
                  sum(len(result.data.source_rows) for result in results), rows,
                  dict(description="All acquisition rows, including excluded and failed rows; source-row mapping matches the spectra columns."))


def _validate(results, fits):
    if len(results) not in (1, 2) or len(fits) != len(results):
        raise ValueError("SHG workbook export needs one result, or a reference and sample with matching fits.")
    for result, fit in zip(results, fits):
        count = len(result.data.source_rows)
        shape = (count, len(result.data.wavelength_nm))
        if any(np.asarray(values).shape != shape for values in
               (result.data.spectra, result.cleaned_spectra, result.corrected)):
            raise ValueError("SHG spectrum arrays do not match their angle and wavelength axes.")
        if len(result.measured_angle_deg) != count:
            raise ValueError("SHG angles do not match acquisition rows.")
        if fit is not None and any(np.asarray(values).shape != (count,) for values in
                                   (fit.fit_mask, fit.fitted_intensity, fit.residual)):
            raise ValueError("SHG fit arrays do not match acquisition rows.")


def export_shg_workbook(folder, *, results, fits=(), twist=None, fit_settings=None,
                        fit_message="", processed_name="Processed Data", include_csv=False,
                        include_cleaned=False):
    """Write one workbook and one complete settings JSON without overwriting.

    Callers reserve a unique analysis directory (and may stage it atomically).
    A failed export removes only files created by this call. Historical CSV APIs
    remain available; optional CSVs here are exact worksheet copies.
    """
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    check_load_cancelled()
    results, fits = tuple(results), tuple(fits)
    if not fits:
        fits = (twist.reference_fit, twist.sample_fit) if twist is not None else (None,) * len(results)
    _validate(results, fits)
    roles = ("reference", "sample") if len(results) == 2 else ("single",)
    if twist is not None and len(results) != 2:
        raise ValueError("A twist result requires reference and sample SHG data.")
    if fit_settings is None:
        fit_settings = next((fit.settings for fit in fits if fit is not None), ShgFitSettings(enabled=False))
    fit_settings = fit_settings.to_dict() if hasattr(fit_settings, "to_dict") else dict(fit_settings)
    spectra, mappings = [], []
    for role, result in zip(roles, results):
        tables, mapping = _spectra_tables(result, role, len(results) == 2, include_cleaned)
        spectra.extend(tables)
        mappings.append(mapping)
    tables = [_plot_table(results, roles), _fit_table(results, roles, fits),
              _summary_table(results, roles, fits, twist, fit_settings, fit_message),
              _diagnostics_table(results, roles, fits, mappings), *spectra]
    created_utc = datetime.now(timezone.utc).isoformat()
    inputs, datasets = [], []
    for role, result, fit in zip(roles, results, fits):
        check_load_cancelled()
        inputs.append(_source(folder, result.data.source_file, "measurement" if role == "single" else role))
        if result.background_file:
            inputs.append(_source(folder, result.background_file, "background" if role == "single" else role + "_background"))
        datasets.append(dict(role=role, source_csv=result.data.source_file, background_csv=result.background_file,
            settings=result.settings.to_dict(), detected_columns=result.data.detected_columns,
            acquisition_rows=len(result.data.source_rows), included_rows=int(np.count_nonzero(result.included)),
            cosmic_pixels_removed=int(np.sum(result.cosmic_pixels_removed)),
            cosmic_affected_rows=int(np.count_nonzero(result.cosmic_pixels_removed)),
            angular_fit=fit.to_dict() if fit is not None else None))
    tables.insert(0, _overview_table(folder, results, roles, inputs, created_utc, tables))
    for table in tables:
        if len(table.headers) > _MAX_COLUMNS or table.row_count + 1 > _MAX_ROWS:
            raise ValueError(f"SHG sheet {table.name} exceeds the XLSX limit of {_MAX_COLUMNS} columns or {_MAX_ROWS} rows. "
                             "Export a smaller acquisition; no data has been truncated.")
    out_dir = Path(folder) / processed_name
    paths = {"workbook": out_dir / "shg.xlsx", "settings": out_dir / "analysis_settings.json"}
    if include_csv:
        paths.update({"csv_" + table.name.casefold(): out_dir / (table.name.casefold() + ".csv") for table in tables})
    validate_shg_output_paths(paths.values())
    if any(path.exists() for path in paths.values()):
        raise FileExistsError("An SHG workbook export already exists here. Use a new analysis directory.")
    payload = dict(mode="SHG Compare / Twist Angle" if len(results) == 2 else "SHG Processing",
        schema_version=EXPORT_METADATA_SCHEMA_VERSION, workbook_schema_version=1, workflow="SHG",
        dataset_type="derived_analysis_with_spectra", package=out_dir.name, app_version=__version__,
        created_utc=created_utc, inputs=inputs, datasets=datasets,
        settings=results[0].settings.to_dict(), processing_settings=results[0].settings.to_dict(),
        fit_settings=fit_settings, fit_message=fit_message,
        angular_fit=fits[0].to_dict() if len(results) == 1 and fits[0] is not None else None,
        angular_fits={role: fit.to_dict() if fit is not None else None for role, fit in zip(roles, fits)},
        twist_fit=twist.to_dict() if twist is not None else None,
        export_options=dict(include_csv=bool(include_csv), include_cleaned=bool(include_cleaned)),
        sheets={table.name: dict(headers=table.headers, data_rows=table.row_count, **table.metadata) for table in tables},
        outputs=[path.name for path in paths.values()],
        output_manifest=[dict(role=key, filename=path.name) for key, path in paths.items()])
    if len(results) == 1:
        payload.update(source_csv=results[0].data.source_file, background_csv=results[0].background_file)
    else:
        payload.update(reference_csv=results[0].data.source_file, sample_csv=results[1].data.source_file)

    out_dir.mkdir(parents=True, exist_ok=True)
    book = Workbook(write_only=True)
    book.properties.title = "SHG analysis and angle-resolved spectra"
    book.properties.subject = "Origin: Plot X/Y/Y Error; Fit X/Y pairs; Spectra first column X, remaining columns Y."
    created = []
    saved = False
    try:
        for table in tables:
            check_load_cancelled()
            sheet = book.create_sheet(table.name)
            sheet.freeze_panes = "B2" if "spectrum_type" in table.metadata else "A2"
            sheet.sheet_view.showGridLines = True
            overview = table.name == "Overview"
            widths = (16, 26, 44, 76) if overview else [24] + [22] * (len(table.headers) - 1)
            for index in range(1, len(table.headers) + 1):
                sheet.column_dimensions[get_column_letter(index)].width = widths[index - 1]
            row_number = 0
            with ExitStack() as stack:
                writer = None
                if include_csv:
                    path = paths["csv_" + table.name.casefold()]
                    stream = stack.enter_context(path.open("x", encoding="utf-8", newline=""))
                    created.append(path)
                    writer = csv.writer(stream)

                def append(values, header=False):
                    nonlocal row_number
                    row_number += 1
                    check_load_cancelled()
                    values = [_value(value) for value in values]
                    if writer is not None:
                        writer.writerow(values)
                    if overview:
                        # Write-only worksheets need row heights before cells.
                        # Reserve extra width for CJK characters in source paths.
                        lines = max(sum(max(1, math.ceil(sum(2 if ord(char) > 255 else 1 for char in line)
                                                        / (width - 2)))
                                        for line in str(value or "").split("\n"))
                                    for value, width in zip(values, widths))
                        sheet.row_dimensions[row_number].height = 24 if header else max(30, lines * 16 + 8)
                    cells = []
                    for value in values:
                        cell = WriteOnlyCell(sheet, value=_text(value) if isinstance(value, str) else value)
                        if isinstance(value, str):
                            cell.data_type = "s"
                        if overview:
                            cell.alignment = Alignment(vertical="top", wrap_text=True)
                        if header:
                            cell.font = Font(bold=True, color="FFFFFF")
                            cell.fill = PatternFill("solid", fgColor="245B85")
                        cells.append(cell)
                    sheet.append(cells)

                append(table.headers, True)
                for row in table.rows():
                    append(row)
            sheet.close()
        book.active = 0
        check_load_cancelled()
        with paths["workbook"].open("xb") as stream:
            created.append(paths["workbook"])
            book.save(stream)
        saved = True
        check_load_cancelled()
        with paths["settings"].open("x", encoding="utf-8") as stream:
            created.append(paths["settings"])
            json.dump(_json_value(payload), stream, indent=2, ensure_ascii=False, allow_nan=False)
            stream.write("\n")
        check_load_cancelled()
        return paths
    except BaseException:
        for path in reversed(created):
            path.unlink(missing_ok=True)
        raise
    finally:
        # Write-only worksheets own temporary XML files until the archive is
        # saved. Close and remove those scratch files on cancellation/failure.
        if not saved:
            for sheet in book.worksheets:
                if not sheet.closed:
                    sheet.close()
                writer = getattr(sheet, "_writer", None)
                if writer is not None and Path(writer.out).exists():
                    writer.cleanup()
        book.close()


__all__ = ["export_shg_workbook"]
